"""Atualiza o deck MP'27 do Circuito Panamericano (Bridge CP_v6) sobre a versão salva pela Juliana.

Uso:
    python atualizar_ppt_cp_mp27_v6.py <Bridge CP_v6.pptx> <Base_Book_CP_.xlsx> <MENS CP 2027 - MP'27.xlsx>
                                       <Faturamento CP - MP'27.xlsx> <Base de Custo MP'27_v2.xlsx> <saida.pptx>

Ajustes de 20/10 (pedido dela, nova rentabilidade: custo fixo 10.692 / EBIT 2.845):
- slide 1: mantém a tabela colada por ela (imagem); refaz o bridge com a abertura dela e o painel "Resumo executivo";
- slide 2: painel menor, gráficos maiores, mensal de eventos com a sazonalidade da MENS;
- slide 3: bridge e composição do custo fixo pela abertura dela; comentários revistos;
- slide 4: postergações 611 (lista 2026) -> 411 (MP'27): o que foi cortado e o que foi revisto;
- slide 5: perguntas-chave em cartões com indicador.
Fonte dos números do MP'27 e do R9: aba LOCAL da MENS (a tabela colada vem dela). Histórico (FY24, FY25, MP26): Base Book.
"""
import sys
from lxml import etree
import openpyxl
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_AXIS_CROSSES

SRC, BB, MENS, FAT, BC, OUT = sys.argv[1:7]

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MIDBLUE = RGBColor(0x2E, 0x5A, 0xA8)
LIGHT = RGBColor(0xDC, 0xE3, 0xF0)
PANEL = RGBColor(0xF3, 0xF5, 0xFA)
EDGE = RGBColor(0x3C, 0x4B, 0x66)
SKY = RGBColor(0x9D, 0xC3, 0xE6)
GREY = RGBColor(0x80, 0x80, 0x80)
DGREY = RGBColor(0x40, 0x40, 0x40)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)
HIST = [RGBColor(0xBF, 0xBF, 0xBF), RGBColor(0xA6, 0xA6, 0xA6), RGBColor(0xC9, 0xD6, 0xEA),
        RGBColor(0x8E, 0xA9, 0xCF), NAVY]                       # FY24, FY25, MP26, R9, MP27
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def r0(v):
    """arredondamento comercial (meio para cima); 9.619,49998 do base book vira 9.620."""
    v = round(v, 2)
    return int(v + 0.5) if v >= 0 else -int(-v + 0.5)


# ------------------------------------------------------------------ dados: histórico (Base Book) + MP'27/R9 (MENS)
m = openpyxl.load_workbook(BB, data_only=True)["Month_FY_LC_"]
loc = openpyxl.load_workbook(MENS, data_only=True)["LOCAL"]
BB_ROW = {"net": 11, "ev": 12, "oem": 13, "rent": 14, "rd": 15, "evg": 16, "fix": 21, "gm": 22, "adm": 23,
          "ops": 24, "mkt": 25, "csa": 26, "da": 27, "ebit": 31}
MENS_ROW = {"net": 21, "ev": 22, "oem": 23, "rent": 24, "rd": 25, "evg": 26, "fix": 30, "gm": 31, "adm": 32,
            "ops": 33, "mkt": 34, "csa": 35, "da": 36, "ebit": 40}


def serie(k):
    """[Actual 2024, Actual 2025, MP26, R9, MP'27] em BRL mil."""
    hist = [m[c + str(BB_ROW[k])].value or 0 for c in "DEF"]
    r9_ = (loc["S" + str(MENS_ROW[k])].value or 0) / 1e3
    mp_ = (loc["Q" + str(MENS_ROW[k])].value or 0) / 1e3
    assert abs(r9_ - (m["G" + str(BB_ROW[k])].value or 0)) < 0.5, k        # R9 igual nas duas fontes
    return hist + [r9_, mp_]


net, ev, oem, rent, rd, evg = (serie(k) for k in ("net", "ev", "oem", "rent", "rd", "evg"))
fix, gm, adm, ops, mkt, csa, da = (serie(k) for k in ("fix", "gm", "adm", "ops", "mkt", "csa", "da"))
ebit = serie("ebit")
r9, mp, mp26, fy25 = 3, 4, 2, 1
assert abs(sum(x[mp] for x in (ev, oem, rent, rd, evg)) - net[mp]) < 0.5
assert abs(sum(x[mp] for x in (gm, adm, ops, mkt, csa, da)) - fix[mp]) < 0.5
assert abs(net[mp] + fix[mp] - ebit[mp]) < 0.5

# ------------------------------------------------------------------ dados: eventos brutos mensais
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
ev_mp_m = [(loc.cell(8, c).value or 0) / 1e3 for c in range(5, 17)]           # MENS LOCAL E8:P8
fat = openpyxl.load_workbook(FAT, data_only=True)["Volume MP27"]
ev_r9_m = [(fat.cell(13, c).value or 0) / 1e3 for c in range(2, 14)]
rd_mp_m = (loc.cell(11, 5).value or 0) / 1e3
rd_r9_m = (fat.cell(16, 2).value or 0) / 1e3

# ------------------------------------------------------------------ postergações (Base de Custo, Sheet1)
sh = openpyxl.load_workbook(BC, data_only=True)["Sheet1"]
just = {}
for r in range(13, 26):
    k, v, t = sh.cell(r, 3).value, sh.cell(r, 4).value, sh.cell(r, 5).value
    if k and isinstance(v, (int, float)):
        just[str(k).strip()] = (v / 1e3, (t or "").strip().lower())
# lista original de 2026 (611, deck v5) x valor no MP'27; itens cortados marcados por ela (rascunho de 19/10)
POST = [("Operacional", "Destinação do resíduo (Sansão)", 170, just["Cercamento | Manutenções - Sansão"][0]),
        ("Manut. Geral", "Restauração APP", 72, just["Restauração APP"][0]),
        ("Operacional", "Manutenção corretiva elétrica", 23, just["Elétrica"][0]),
        ("Marketing", "Ações de marketing", 24, just["Ações de marketing"][0]),
        ("Operacional", "Análise qualidade do ar (ar-condicionado)", 5,
         just["Análise qualidade do ar (ar-condicionado)"][0]),
        ("Operacional", "Reajustes de contratos", 74, 0),
        ("Operacional", "Pintura box / sinalizações", 70, 0),
        ("Manut. Geral", "Plantio bambuzal + área restauração APP", 50, 0),
        ("Manut. Geral", "Retirada de bambu", 45, 0),
        ("Marketing", "Viagens | reclassificações", 35, 0),
        ("Operacional", "Melhorias – sanitários/vestiários", 21, 0),
        ("Operacional", "Custos I.T.", 20, 0),
        ("Operacional", "Pluvial | controle água da chuva", 2, 0)]
post_orig = sum(o for _, _, o, _ in POST)
post = sum(v for _, _, _, v in POST)
cortado = sum(o for _, _, o, v in POST if not v)
revisto = sum(v - o for _, _, o, v in POST if v)
assert r0(post_orig) == 611 and abs(post - sum(v for v, t in just.values() if "posterga" in t)) < 0.01
assert abs(post_orig - cortado + revisto - post) < 0.01

# ------------------------------------------------------------------ abertura do custo fixo (Juliana, 20/10)
HSE, LABOUR, INFL, REDUCAO, PORTARIA = 300, 138, 213, -150, 40
INFL_PCT = 4.4
d_fix = -(fix[mp] - fix[r9])
DRIVERS = [("Postergações\n26 → 27", post), ("Itens obrig.\nHSE", HSE), ("Labour", LABOUR),
           (f"Inflação\n({INFL_PCT:.1f}%)".replace(".", ","), INFL),
           ("Redução de\ndespesas", REDUCAO), ("Projeto\nportaria", PORTARIA)]
assert abs(sum(v for _, v in DRIVERS) - d_fix) < 1, (sum(v for _, v in DRIVERS), d_fix)
HL = ("Postergações\n26 → 27", "Itens obrig.\nHSE")

# ------------------------------------------------------------------ números-chave
terc = [a + b + c + d for a, b, c, d in zip(ev, oem, rent, evg)]      # terceiros (tudo menos R&D)
ev_tot = [a + b for a, b in zip(ev, evg)]
oem_tot = [a + b for a, b in zip(oem, rent)]
rd_chart = list(rd)
rd_chart[r9] = rd[mp26]                    # R9 = MP26 (9.620); base book 9.619,5 (arredondamento)
d_net = net[mp] - net[r9]
d_rd, d_terc = rd[mp] - rd[r9], terc[mp] - terc[r9]
d_ebit = ebit[mp] - ebit[r9]
ros_r9, ros_mp = ebit[r9] / net[r9] * 100, ebit[mp] / net[mp] * 100
ebit_sem_post = ebit[mp] + post
ros_sem = ebit_sem_post / net[mp] * 100
ebitda = ebit[mp] - da[mp]
pontual = post + HSE
recorrente = d_fix - pontual
sens_ev = ev_tot[mp] * 0.10                # eventos: margem de contribuição de 100%
ebit_sem_reaj = ebit[mp] - d_rd
pico = max(range(12), key=lambda i: ev_mp_m[i])
vale = min(range(12), key=lambda i: ev_mp_m[i])


def fmt(v, sign=False, paren=False):
    s = f"{abs(r0(v)):,}".replace(",", ".")
    if r0(v) < 0:
        return f"({s})" if paren else "-" + s
    return ("+" if sign and r0(v) > 0 else "") + s


def pct(a, b):
    return (a / b - 1) * 100


def p1(x, sign=True):
    return (f"{x:+.1f}" if sign else f"{x:.1f}").replace(".", ",")


# ------------------------------------------------------------------ helpers de desenho (mesmos do gerador v6)
def line(sl, x1, y1, x2, y2, color=GREY, w=0.75, dash=None, arrow=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    if arrow:
        etree.SubElement(c.line._get_or_add_ln(), A + "tailEnd").set("type", "triangle")
    return c


def text(sl, x, y, w, h, runs, size=9, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bold=False, italic=False, spacing=1.0, after=0):
    """runs: str (\\n = novo parágrafo) ou lista de parágrafos; parágrafo = lista de (texto, {fmt})."""
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    paras = [[(ln, {})] for ln in runs.split("\n")] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        for t, f in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(f.get("size", size))
            r.font.bold = f.get("bold", bold)
            r.font.italic = f.get("italic", italic)
            r.font.color.rgb = f.get("color", color)
    return tb


def rect(sl, x, y, w, h, fill, edge=None, ew=0.75, shape=MSO_SHAPE.RECTANGLE):
    r = sl.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = fill
    if edge is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = edge
        r.line.width = Pt(ew)
    r.shadow.inherit = False
    return r


def headline(sl, main, accent, y=62, size=16):
    text(sl, 36, y, 900, 40, [[(main, {"bold": True, "color": NAVY, "size": size}),
                               (accent, {"bold": True, "color": MIDBLUE, "size": size})]])


def break_mark(sl, c, bw, axis):
    top = axis - 11
    f = sl.shapes.build_freeform(Pt(c - bw / 2 - 3), Pt(top + 5))
    f.add_line_segments([(Pt(c + bw / 2 + 3), Pt(top - 3)), (Pt(c + bw / 2 + 3), Pt(top + 1)),
                         (Pt(c - bw / 2 - 3), Pt(top + 9))], close=True)
    b = f.convert_to_shape()
    b.fill.solid()
    b.fill.fore_color.rgb = WHITE
    b.line.fill.background()
    b.shadow.inherit = False


def oval_label(sl, cx_, cy, txt, w=64):
    ov = rect(sl, cx_ - w / 2, cy - 10, w, 20, WHITE, GREY, shape=MSO_SHAPE.OVAL)
    tf = ov.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BLACK


def waterfall(sl, x0, x1, axis, room, start, steps, end_label, start_label, highlight=(), tot_h=18,
              lab_size=9, bracket=None, cat_size=8.5):
    """Bridge em formas nativas. steps = [(rótulo, delta)]; totais com eixo truncado."""
    n = len(steps) + 2
    dx = (x1 - x0) / n
    bw = min(46.0, dx * 0.66)
    cx = [x0 + dx * (i + 0.5) for i in range(n)]
    cum, lo, hi = 0, 0, 0
    for _, v in steps:
        cum += v
        lo, hi = min(lo, cum), max(hi, cum)
    sc = (room - tot_h - 24) / max(hi - lo, 1)
    base = axis - tot_h + lo * sc if lo < 0 else axis - tot_h            # nível do total inicial
    line(sl, x0, axis, x1, axis, color=BLACK, w=0.75)
    rect(sl, cx[0] - bw / 2, base, bw, axis - base, NAVY)
    text(sl, cx[0] - 30, base - 19, 60, 16, fmt(start), size=13, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.BOTTOM)
    level = base
    min_top = base
    for i, (lab, v) in enumerate(steps, start=1):
        new = level - v * sc
        top, bot = min(level, new), max(level, new)
        min_top = min(min_top, top)
        hl = lab in highlight
        rect(sl, cx[i] - bw / 2, top, bw, max(bot - top, 1.2), LIGHT, MIDBLUE if hl else EDGE, 1.75 if hl else 0.75)
        line(sl, cx[i - 1] + bw / 2, level, cx[i] - bw / 2, level, color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
        text(sl, cx[i] - 30, top - 14, 60, 12, fmt(v, sign=True), size=lab_size, bold=True,
             color=BLACK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        level = new
    end = start + sum(v for _, v in steps)
    rect(sl, cx[-1] - bw / 2, level, bw, axis - level, NAVY)
    line(sl, cx[-2] + bw / 2, level, cx[-1] - bw / 2, level, color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
    text(sl, cx[-1] - 30, level - 19, 60, 16, fmt(end), size=13, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.BOTTOM)
    break_mark(sl, cx[0], bw, axis)
    break_mark(sl, cx[-1], bw, axis)
    labels = [start_label] + [l for l, _ in steps] + [end_label]
    for c, lab in zip(cx, labels):
        hl = lab in highlight
        text(sl, c - dx / 2, axis + 4, dx, 26, lab, size=cat_size, bold=hl or lab in (start_label, end_label),
             color=MIDBLUE if hl else BLACK, align=PP_ALIGN.CENTER)
    if bracket:
        ytop = min(min_top, level) - 32
        line(sl, cx[0], ytop, cx[0], base - 22)
        line(sl, cx[0], ytop, cx[-1], ytop)
        line(sl, cx[-1], ytop, cx[-1], level - 22, arrow=True)
        oval_label(sl, (cx[0] + cx[-1]) / 2, ytop, bracket)
    return end


def style_chart(ch, size=8):
    ch.font.size = Pt(size)
    ch.font.color.rgb = DGREY
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(size)
    ca.format.line.color.rgb = GREY
    ca.has_major_gridlines = False


def badge(sl, x, y, d, label="vs R9"):
    rect(sl, x, y, 64, 18, LIGHT, None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(sl, x, y, 64, 18, f"{p1(d)}% {label}", size=8.5, bold=True,
         color=GREEN if d >= 0 else RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def small_multiple(sl, x, y, w, h, title, vals, sub, ymax):
    text(sl, x, y, w - 66, 16, title, size=11, bold=True, color=NAVY)
    text(sl, x, y + 15, w, 12, sub, size=8, italic=True, color=GREY)
    cd = CategoryChartData()
    cd.categories = ["FY24", "FY25", "MP26", "R9", "MP’27"]
    cd.add_series(title, [r0(v) for v in vals])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(x), Pt(y + 28), Pt(w), Pt(h - 28), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    style_chart(ch, size=8.5)
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.series[0].data_labels
    dl.show_value = True
    dl.number_format = '#,##0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(8.5)
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = HIST[i]
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = ymax
    badge(sl, x + w - 64, y, pct(vals[4], vals[3]))


def table(sl, x, y, col_w, rows, row_h=19, header=None, size=9, total_last=False, col_align=None):
    nr = len(rows) + (1 if header else 0)
    gt = sl.shapes.add_table(nr, len(col_w), Pt(x), Pt(y), Pt(sum(col_w)), Pt(row_h * nr))
    tbl = gt.table
    pr = gt._element.graphic.graphicData.tbl.tblPr
    pr.set("firstRow", "0")
    pr.set("bandRow", "0")
    sid = pr.find(A + "tableStyleId")
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Pt(w)
    col_align = col_align or [PP_ALIGN.LEFT] * (len(col_w) - 1) + [PP_ALIGN.RIGHT]
    allrows = ([header] if header else []) + rows
    for r_, vals in enumerate(allrows):
        tbl.rows[r_].height = Pt(row_h)
        is_h = header and r_ == 0
        is_t = total_last and r_ == nr - 1
        for c_, v in enumerate(vals):
            ce = tbl.cell(r_, c_)
            ce.margin_top = ce.margin_bottom = Pt(1)
            ce.margin_left = ce.margin_right = Pt(5)
            ce.vertical_anchor = MSO_ANCHOR.MIDDLE
            ce.text = ""
            p = ce.text_frame.paragraphs[0]
            p.alignment = col_align[c_]
            t, f = (v if isinstance(v, tuple) else (v, {}))
            rr = p.add_run()
            rr.text = str(t)
            rr.font.size = Pt(f.get("size", size))
            rr.font.bold = f.get("bold", bool(is_h or is_t))
            rr.font.italic = f.get("italic", False)
            rr.font.color.rgb = f.get("color", WHITE if is_h else BLACK)
            # sem isto a fonte padrão (18pt) do parágrafo vazio infla a altura da linha
            etree.SubElement(p._p, A + "endParaRPr").set("sz", str(int(f.get("size", size) * 100)))
            if is_h:
                ce.fill.solid()
                ce.fill.fore_color.rgb = NAVY
            elif is_t:
                ce.fill.solid()
                ce.fill.fore_color.rgb = LIGHT
            else:
                ce.fill.background()
        if not (is_h or is_t):
            for c_ in range(len(col_w)):
                tcPr = tbl.cell(r_, c_)._tc.get_or_add_tcPr()
                ln = etree.SubElement(tcPr, A + "lnB", w="3175")
                sf = etree.SubElement(ln, A + "solidFill")
                etree.SubElement(sf, A + "srgbClr", val="D9D9D9")
    return gt


def gain_label(sl):
    text(sl, 10, 52, 200, 14, "(+) gain  |  BRL ’000", size=9, italic=True)


def panel_blocks(sl, x, y, w, h, title, items, size=10):
    """Painel com blocos (título em negrito + texto) distribuídos igualmente na altura – sem sobra."""
    rect(sl, x, y, w, h, PANEL, LIGHT)
    rect(sl, x, y, 4, h, MIDBLUE)
    text(sl, x + 16, y + 10, w - 28, 20, title, size=13, bold=True, color=NAVY)
    top = y + 38
    bh = (y + h - 8 - top) / len(items)
    for i, (head, body) in enumerate(items):
        by = top + i * bh
        if i:
            line(sl, x + 16, by, x + w - 14, by, color=RGBColor(0xC9, 0xD3, 0xE6), w=0.5)
        text(sl, x + 16, by + 4, w - 30, bh - 8,
             [[(head, {"bold": True, "color": NAVY, "size": size + 0.5})],
              [(body, {"color": DGREY, "size": size})]],
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.08, after=3)


def limpar(sl, manter):
    """Remove as formas do slide (menos as de `manter`) e solta os gráficos removidos."""
    for shp in list(sl.shapes):
        if manter(shp):
            continue
        if shp.has_chart:
            rid = shp._element.xpath(".//c:chart/@r:id")[0]
            sl.part.drop_rel(rid)
        shp._element.getparent().remove(shp._element)


def nota(sl, txt):
    sl.notes_slide.notes_text_frame.text = txt


prs = Presentation(SRC)
s1, s2, s3, s4, s5 = prs.slides

# ================================================================== slide 1 – resumo
# mantém a tabela colada por ela (imagem), o objeto think-cell, o título, "(+) gain" e o texto vertical
limpar(s1, lambda shp: shp.name in ("Object 4", "Title 2") or shp.shape_type == 13
       or (shp.has_text_frame and shp.text_frame.text.startswith(("(+) gain", "Result Variant"))))
steps = [("R&D\n(intercomp.)", d_rd), ("Terceiros", d_terc)] + [(l, -v) for l, v in DRIVERS]
assert abs(ebit[r9] + sum(v for _, v in steps) - ebit[mp]) < 1
waterfall(s1, 66, 700, 492, 150, ebit[r9], steps, "MP’27", "R9", highlight=HL,
          bracket=f"{p1(pct(ebit[mp], ebit[r9]))}%", cat_size=8)
text(s1, 66, 528, 640, 10, "Terceiros = eventos (externos e grupo), testes OEM e aluguéis de longo prazo. "
     f"Postergações: lista original de 2026 de {fmt(post_orig)}, revisada para {fmt(post)}.",
     size=7, italic=True, color=GREY)

panel_blocks(s1, 722, 70, 226, 450, "Resumo executivo", [
    (f"EBIT {fmt(ebit[mp])}  |  {fmt(d_ebit, sign=True)} vs R9 ({p1(pct(ebit[mp], ebit[r9]))}%)",
     f"ROS de {p1(ros_r9, False)}% para {p1(ros_mp, False)}%: a receita cresce +{fmt(d_net)} e absorve "
     f"o aumento de +{fmt(d_fix)} do custo fixo."),
    (f"Receita {fmt(net[mp])}  |  {p1(pct(net[mp], net[r9]))}%",
     f"Contrato de R&D +{p1(pct(rd_mp_m, rd_r9_m), False)}% (+{fmt(d_rd)}) e terceiros +{fmt(d_terc)} "
     f"com a retomada de eventos (bruto {fmt(sum(ev_mp_m))} vs {fmt(sum(ev_r9_m))} no R9)."),
    (f"Custo fixo {fmt(-fix[mp])}  |  {p1(pct(fix[mp], fix[r9]))}%",
     f"{pontual / d_fix * 100:.0f}% do aumento é pontual ou obrigatório: postergações ({fmt(post)}) e "
     f"novos itens HSE ({fmt(HSE)}). Inflação, labour e portaria parcialmente compensados por "
     f"redução de despesas ({fmt(REDUCAO)})."),
    (f"Postergações: de {fmt(post_orig)} para {fmt(post)}",
     f"Lista de 2026 revisada, com corte de {fmt(post_orig - post)}. Sem elas, EBIT de "
     f"{fmt(ebit_sem_post)} (ROS {p1(ros_sem, False)}%)."),
    (f"EBITDA {fmt(ebitda)}  |  {p1(ebitda / net[mp] * 100, False)}% da receita",
     f"Forte geração de caixa; D&A estável em {fmt(-da[mp])}."),
], size=9.3)
nota(s1, "Tabela: colada pela Juliana (MENS CP 2027 - MP'27, aba LOCAL).\n"
     "Bridge EBIT (R$ mil): " + " | ".join(f"{l.replace(chr(10), ' ')} {v:,.1f}" for l, v in
                                         [("R9", ebit[r9])] + steps + [("MP'27", ebit[mp])]) +
     "\nAbertura do custo fixo informada pela Juliana em 20/10: postergações 411, HSE 300, labour 138, "
     "inflação 4,4% 213, redução de despesas -150, portaria 40 (= 952).")

# ================================================================== slide 2 – receita por origem
limpar(s2, lambda shp: shp.is_placeholder)
gain_label(s2)
headline(s2, f"Receita líquida de BRL {p1(net[mp] / 1e3, False)} mi ({p1(pct(net[mp], net[r9]))}% vs R9)",
         " — R&D garante a base e terceiros retomam após 2026 atípico")

PX, PW = 760, 188                         # painel de destaques (mais estreito que na v6)
W, Y0, H, GAP = 230, 100, 205, 14
ymax = max(rd_chart) * 1.18
small_multiple(s2, 36, Y0, W, H, "R&D Tests", rd_chart, "intercompany – contrato mensal", ymax)
small_multiple(s2, 36 + W + GAP, Y0, W, H, "Terceiros", terc, "eventos + OEM + aluguéis", ymax * 0.5)

xt = 36 + 2 * (W + GAP)
text(s2, xt, Y0, W - 66, 16, "Total", size=11, bold=True, color=NAVY)
text(s2, xt, Y0 + 15, W, 12, "receita líquida = R&D + terceiros", size=8, italic=True, color=GREY)
cd = CategoryChartData()
cd.categories = ["FY24", "FY25", "MP26", "R9", "MP’27"]
tot_r = [r0(v) for v in net]
cd.add_series("R&D", [r0(v) for v in rd_chart])
cd.add_series("Terceiros", [r0(v) for v in terc])
cd.add_series("Total", [0] * 5)
gf = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Pt(xt), Pt(Y0 + 28), Pt(W), Pt(H - 28), cd)
ch = gf.chart
style_chart(ch, size=8.5)
ch.has_title = False
ch.has_legend = False
plot = ch.plots[0]
plot.gap_width = 45
plot.overlap = 100
for se, col in zip(plot.series, (NAVY, SKY, None)):
    if col is None:
        se.format.fill.background()
        for j, pt in enumerate(se.points):
            dl = pt.data_label
            dl.position = XL_LABEL_POSITION.INSIDE_BASE
            dl.text_frame.text = fmt(tot_r[j])
            dl.text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
            dl.text_frame.paragraphs[0].runs[0].font.bold = True
            dl.text_frame.paragraphs[0].runs[0].font.color.rgb = BLACK
        continue
    se.format.fill.solid()
    se.format.fill.fore_color.rgb = col
    dl = se.data_labels
    dl.show_value = True
    dl.number_format = '#,##0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.CENTER
    dl.font.size = Pt(8)
    dl.font.color.rgb = WHITE if col == NAVY else NAVY
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = max(net) * 1.13
badge(s2, xt + W - 64, Y0, pct(net[mp], net[r9]))
for k, (lab, col) in enumerate((("R&D", NAVY), ("Terceiros", SKY))):
    rect(s2, xt + 76 + k * 62, Y0 + H + 2, 8, 8, col)
    text(s2, xt + 87 + k * 62, Y0 + H, 50, 12, lab, size=8, color=DGREY)

# eventos mês a mês (bruto) – sazonalidade da MENS
yb = Y0 + H + 20
text(s2, 36, yb, 470, 16, "Eventos externos – faturamento bruto mensal (BRL ’000)", size=11, bold=True, color=NAVY)
text(s2, 36, yb + 15, 470, 12, f"R9 {fmt(sum(ev_r9_m))}  →  MP’27 {fmt(sum(ev_mp_m))} no ano "
     f"({p1(pct(sum(ev_mp_m), sum(ev_r9_m)))}%)", size=8, italic=True, color=GREY)
cd = CategoryChartData()
cd.categories = MESES
cd.add_series("2026 (R9)", [r0(v) for v in ev_r9_m])
cd.add_series("MP’27", [r0(v) for v in ev_mp_m])
gf = s2.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Pt(30), Pt(yb + 26), Pt(490), Pt(522 - yb - 26), cd)
ch = gf.chart
style_chart(ch, size=8)
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(8)
ch.value_axis.minimum_scale = -90
ch.value_axis.maximum_scale = 560
ch.value_axis.crosses = XL_AXIS_CROSSES.MINIMUM
R9COL = RGBColor(0x8E, 0xA9, 0xCF)
for k, (se, col) in enumerate(zip(ch.plots[0].series, (R9COL, NAVY))):
    se.smooth = False
    se.format.line.color.rgb = col
    se.format.line.width = Pt(1.75 if k == 0 else 2.5)
    if k == 0:
        se.format.line.dash_style = MSO_LINE.DASH
    se.marker.style = 8
    se.marker.size = 5 if k == 0 else 6
    se.marker.format.fill.solid()
    se.marker.format.fill.fore_color.rgb = col if k else WHITE
    se.marker.format.line.color.rgb = col
    lab_col = NAVY if k else RGBColor(0x5B, 0x7D, 0xB1)
    for j in range(12):                        # rótulo da linha mais alta acima, da mais baixa abaixo
        mine, other = (ev_r9_m[j], ev_mp_m[j]) if k == 0 else (ev_mp_m[j], ev_r9_m[j])
        pl = se.points[j].data_label
        pl.has_text_frame = False
        pl.show_value = True
        pl.number_format = '#,##0'
        pl.number_format_is_linked = False
        pl.font.size = Pt(8)
        pl.font.bold = bool(k)
        pl.font.color.rgb = lab_col
        pl.position = XL_LABEL_POSITION.ABOVE if mine > other or (mine == other and k == 1) else XL_LABEL_POSITION.BELOW

# mix da receita MP'27 (rosca)
xm = 544
text(s2, xm, yb, 200, 16, "Mix da receita MP’27", size=11, bold=True, color=NAVY)
mix = [("R&D", rd[mp], NAVY), ("Eventos", ev_tot[mp], MIDBLUE), ("OEM + aluguéis", oem_tot[mp], SKY)]
cd = CategoryChartData()
cd.categories = [l for l, _, _ in mix]
cd.add_series("Mix", [v for _, v, _ in mix])
DS = 118
gf = s2.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Pt(xm - 6), Pt(yb + 24), Pt(DS), Pt(DS), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
plot = ch.plots[0]
for (_, _, col), pt in zip(mix, plot.series[0].points):
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
    pt.format.line.color.rgb = WHITE
    pt.format.line.width = Pt(1.5)
dn = plot._element
hs = dn.find(C + "holeSize")
if hs is None:
    hs = etree.SubElement(dn, C + "holeSize")
hs.set("val", "58")
fsa = dn.find(C + "firstSliceAng")
if fsa is not None:
    fsa.set("val", "0")
text(s2, xm - 6 + DS / 2 - 27, yb + 24 + DS / 2 - 14, 54, 28,
     [[(f"{fmt(net[mp])}", {"bold": True, "size": 10, "color": NAVY})], [("BRL ’000", {"size": 7, "color": GREY})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for k, (lab, v, col) in enumerate(mix):
    yy = yb + 38 + k * 32
    rect(s2, xm + DS + 6, yy + 4, 9, 9, col)
    text(s2, xm + DS + 20, yy, 80, 30, [[(f"{v / net[mp] * 100:.0f}%", {"bold": True, "size": 11, "color": NAVY}),
                                         (f"  {fmt(v)}", {"size": 8, "color": GREY})],
                                        [(lab, {"size": 8, "color": DGREY})]], spacing=0.95)
text(s2, xm, yb + 152, 200, 40, [[("Terceiros: ", {"bold": True, "size": 8.5, "color": NAVY}),
                                  (f"{terc[mp] / net[mp] * 100:.0f}% da receita ({fmt(terc[mp])}); eventos "
                                   "com margem de contribuição de 100%.", {"size": 8.5, "color": DGREY})]],
     spacing=1.05)

panel_blocks(s2, PX, Y0, PW, 420, "Destaques", [
    (f"R&D  |  {rd[mp] / net[mp] * 100:.0f}% da receita",
     f"Contrato de {fmt(rd_r9_m)} para {fmt(rd_mp_m)} mil/mês brutos (+{p1(pct(rd_mp_m, rd_r9_m), False)}%): "
     f"+{fmt(d_rd)}."),
    (f"Terceiros  |  {p1(pct(terc[mp], terc[r9]))}%",
     f"Eventos +{fmt(ev_tot[mp] - ev_tot[r9])} após 2026 atípico (Copa do Mundo). Ainda "
     f"{p1(pct(terc[mp], terc[fy25]))}% vs FY25: premissa conservadora."),
    ("Sazonalidade",
     f"Pico em {MESES[pico].lower()} ({fmt(ev_mp_m[pico])}) e menor mês em {MESES[vale].lower()} "
     f"({fmt(ev_mp_m[vale])}); {p1(sum(ev_mp_m[8:]) / sum(ev_mp_m) * 100, False)}% do ano em set–dez "
     f"(R9: {p1(sum(ev_r9_m[8:]) / sum(ev_r9_m) * 100, False)}%)."),
    ("Sensibilidade",
     f"Cada 10% de eventos = ±{fmt(sens_ev)} de EBIT (sem custo variável)."),
], size=8.8)
nota(s2, "Fonte: MENS CP 2027 - MP'27 (aba LOCAL: receita líquida e eventos brutos mensais), Base Book (histórico) "
     "e Faturamento CP - MP'27 (R9 mensal).\n"
     f"Terceiros inclui Events sales group ({fmt(evg[mp])} no MP'27), eventos do grupo Pirelli.\n"
     f"R&D R9 mostrado como {fmt(rd_chart[r9])} (igual MP26); base book {rd[r9]:,.1f}.")

# ================================================================== slide 3 – custo fixo
limpar(s3, lambda shp: shp.is_placeholder)
gain_label(s3)
headline(s3, f"Custo fixo de BRL {p1(-fix[mp] / 1e3, False)} mi ({p1(pct(fix[mp], fix[r9]))}% vs R9)",
         f" — {pontual / d_fix * 100:.0f}% do aumento vem de postergações e de novos itens obrigatórios de HSE")
text(s3, 36, 100, 520, 16, "Custo fixo R9 → MP’27 (BRL ’000, custo = positivo)", size=11, bold=True, color=NAVY)
waterfall(s3, 36, 560, 440, 290, -fix[r9], DRIVERS, "MP’27", "R9", highlight=HL,
          bracket=f"{p1(pct(fix[mp], fix[r9]))}%")
text(s3, 36, 500, 530, 22, "Por linha do P&L: estrutura operacional "
     f"{fmt(-(ops[mp] - ops[r9]), sign=True)}, administrativo {fmt(-(adm[mp] - adm[r9]), sign=True)}, "
     f"manutenção geral {fmt(-(gm[mp] - gm[r9]), sign=True)}, CSA/CSC {fmt(-(csa[mp] - csa[r9]), sign=True)}, "
     f"marketing {fmt(-(mkt[mp] - mkt[r9]), sign=True)} e D&A {fmt(-(da[mp] - da[r9]), sign=True)}.",
     size=7, italic=True, color=GREY)

text(s3, 590, 100, 358, 16, "Composição do aumento (BRL ’000)", size=11, bold=True, color=NAVY)
G = {"color": GREY, "size": 8}
rows = [[("Pontual", G), f"Postergações de 2026 → 2027 (lista de {fmt(post_orig)} revisada)", fmt(post)],
        [("Obrigatório", G), "Novos itens HSE: licenças ambientais + outorga e biodigestor", fmt(HSE)],
        [("Recorrente", G), "Labour: dissídio e provisões", fmt(LABOUR)],
        [("Recorrente", G), f"Inflação de contratos ({p1(INFL_PCT, False)}%)", fmt(INFL)],
        [("Recorrente", G), "Redução de despesas", (fmt(REDUCAO), {"color": GREEN})],
        [("Projeto", G), "Projeto da portaria", fmt(PORTARIA)],
        ["", ("Total", {"bold": True}), (fmt(d_fix), {"bold": True, "color": MIDBLUE})]]
table(s3, 590, 120, [64, 244, 50], rows, row_h=19, header=["Natureza", "Driver", "BRL ’000"], size=8.5,
      total_last=True, col_align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT])

panel_blocks(s3, 590, 290, 358, 230, "O que está por trás do aumento", [
    (f"{pontual / d_fix * 100:.0f}% pontual ou obrigatório  |  {fmt(pontual)}",
     f"Postergações de 2026 ({fmt(post)}) e novos itens de HSE ({fmt(HSE)}): renovação das licenças "
     "ambientais + outorga e coleta de lixo orgânico (biodigestor)."),
    (f"Custo recorrente: {p1(recorrente / -fix[r9] * 100)}%  |  {fmt(recorrente, sign=True)}",
     f"Inflação de {p1(INFL_PCT, False)}% ({fmt(INFL)}), labour ({fmt(LABOUR)}) e portaria ({fmt(PORTARIA)}); "
     f"a redução de despesas de {fmt(-REDUCAO)} absorve {-REDUCAO / INFL * 100:.0f}% da inflação."),
    (f"D&A {fmt(-da[mp])}  |  {-da[mp] / -fix[mp] * 100:.0f}% do custo fixo",
     "Estável vs R9 e sem efeito caixa."),
], size=9)
nota(s3, "Abertura do custo fixo informada pela Juliana em 20/10 (= variação da MENS, aba LOCAL).\n"
     + " | ".join(f"{l.replace(chr(10), ' ')} {v:,.1f}" for l, v in DRIVERS) + f" | Total MENS {d_fix:,.1f}")

# ================================================================== slide 4 – postergações
limpar(s4, lambda shp: shp.is_placeholder)
gain_label(s4)
headline(s4, f"Postergações de 2026: de BRL {fmt(post_orig)} mil para {fmt(post)} mil",
         f" — revisamos a lista e cortamos {fmt(post_orig - post)}")

text(s4, 36, 100, 520, 16, "Lista de postergações 2026 → MP’27 (BRL ’000)", size=11, bold=True, color=NAVY)
kept = sorted([p for p in POST if p[3]], key=lambda p: -p[3])
cut = sorted([p for p in POST if not p[3]], key=lambda p: -p[2])
rows = []
for g, it, o, v in kept:
    st = ("Mantido", {"color": NAVY, "bold": True}) if abs(v - o) < 0.5 else \
        (f"Revisado ({fmt(v - o, sign=True)})", {"color": MIDBLUE, "bold": True})
    rows.append([(g, G), it, fmt(o), (fmt(v), {"bold": True}), st])
for g, it, o, v in cut:
    gc = {"color": GREY}
    rows.append([(g, G), (it, gc), (fmt(o), gc), ("–", gc), ("Cortado", {"color": RED, "bold": True})])
rows.append(["", ("TOTAL", {"bold": True}), (fmt(post_orig), {"bold": True}), (fmt(post), {"bold": True,
             "color": MIDBLUE}), (f"{fmt(post - post_orig)}", {"bold": True, "color": RED})])
table(s4, 36, 120, [76, 236, 58, 58, 92], rows, row_h=24.5,
      header=["Área", "Item", "Lista 2026", "MP’27", "Status"], size=9, total_last=True,
      col_align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER])

# mini-ponte 611 -> 411 (barras em escala real, sem corte de eixo)
bx0, bx1, axis, top = 590, 948, 262, 142
text(s4, bx0, 100, 358, 16, f"De {fmt(post_orig)} para {fmt(post)} (BRL ’000)", size=11, bold=True, color=NAVY)
sc = (axis - top) / post_orig
cols = [("Lista 2026", 0, post_orig, NAVY), ("Itens\ncortados", post_orig - cortado, post_orig, None),
        ("Revisão de\nvalores", post_orig - cortado, post_orig - cortado + revisto, None),
        ("MP’27", 0, post, NAVY)]
dx = (bx1 - bx0) / len(cols)
bw = dx * 0.55
line(s4, bx0, axis, bx1, axis, color=BLACK, w=0.75)
prev = None
for i, (lab, a, b, col) in enumerate(cols):
    c = bx0 + dx * (i + 0.5)
    y_top, y_bot = axis - b * sc, axis - a * sc
    if col:
        rect(s4, c - bw / 2, y_top, bw, y_bot - y_top, col)
        val = fmt(b)
    else:
        neg = i == 1
        rect(s4, c - bw / 2, y_top, bw, y_bot - y_top, LIGHT, RED if neg else MIDBLUE, 1.5)
        val = fmt(-cortado) if neg else fmt(revisto, sign=True)
    text(s4, c - 40, y_top - 17, 80, 14, val, size=11 if col else 10, bold=True,
         color=BLACK if col else (RED if i == 1 else MIDBLUE), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    text(s4, c - dx / 2, axis + 4, dx, 26, lab, size=8.5, bold=bool(col), align=PP_ALIGN.CENTER)
    if prev is not None:
        line(s4, prev, axis - [post_orig, post_orig - cortado, post_orig - cortado + revisto][i - 1] * sc,
             c - bw / 2, axis - [post_orig, post_orig - cortado, post_orig - cortado + revisto][i - 1] * sc,
             color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
    prev = c + bw / 2

top_kept = kept[0][3] + kept[1][3]
panel_blocks(s4, 590, 300, 358, 220, "Mensagens-chave", [
    ("2026: alavanca para proteger o resultado",
     "Com a queda de eventos (ano de Copa do Mundo), adiamos estes gastos sem comprometer a operação "
     "e a segurança do Circuito."),
    (f"Cortamos {fmt(post_orig - post)} da lista original",
     f"{len(cut)} itens saíram ({fmt(cortado)}), como pintura de box/sinalizações, bambu e viagens; "
     f"Sansão, APP e elétrica foram reorçados (+{fmt(revisto)})."),
    (f"Ficam {fmt(post)}, {top_kept / post * 100:.0f}% em Sansão e APP",
     f"Gasto pontual: sem ele, EBIT {fmt(ebit_sem_post)} e ROS {p1(ros_sem, False)}% "
     f"(vs {p1(ros_mp, False)}%)."),
], size=9)
nota(s4, "Valores MP'27: Base de Custo MP'27_v2.xlsx / Sheet1 (justificativa 'postergação de 2026'). "
     "Lista original de 2026 (611): deck Bridge CP_v5, slide 4. Itens cortados: marcados pela Juliana no rascunho "
     "'CLAUDE OLHAR AQUI' (19/10).\n"
     f"Cortados {cortado:,.0f} | revisões de valor +{revisto:,.0f} (Sansão 170→180, restauração 72→167, "
     f"elétrica 23→35) | líquido {post - post_orig:,.0f}.")

# ================================================================== slide 5 – perguntas-chave
limpar(s5, lambda shp: shp.is_placeholder)
gain_label(s5)
headline(s5, "Perguntas-chave", " — o que sustenta o MP’27 do Circuito Panamericano")
QA = [
    (fmt(d_ebit, sign=True), "EBIT vs R9", "O que explica a melhora do EBIT?",
     f"Receita +{fmt(d_net)} (R&D +{fmt(d_rd)}; terceiros +{fmt(d_terc)}) supera o aumento do custo fixo "
     f"(+{fmt(d_fix)}). EBIT {fmt(ebit[mp])}, ROS {p1(ros_mp, False)}%."),
    (f"{rd[mp] / net[mp] * 100:.0f}%", "da receita é R&D", "Qual a dependência do contrato de R&D?",
     f"Receita intercompany. Sem o reajuste do contrato (+{p1(pct(rd_mp_m, rd_r9_m), False)}%), o EBIT seria "
     f"{fmt(ebit_sem_reaj)}, abaixo do R9 ({fmt(ebit[r9])})."),
    (fmt(sum(ev_mp_m)), "eventos brutos", "A retomada de eventos é realista?",
     f"vs {fmt(sum(ev_r9_m))} no R9 e ainda abaixo de FY25 e do MP26. Cada 10% de eventos = "
     f"±{fmt(sens_ev)} de EBIT."),
    (f"{pontual / d_fix * 100:.0f}%", "do aumento é pontual", "O aumento do custo fixo é recorrente?",
     f"Não na maior parte: postergações ({fmt(post)}) e HSE ({fmt(HSE)}). O recorrente cresce "
     f"{p1(recorrente / -fix[r9] * 100, False)}%, já com redução de despesas ({fmt(REDUCAO)})."),
    (fmt(ebitda), f"EBITDA ({p1(ebitda / net[mp] * 100, False)}%)", "E a geração de caixa?",
     f"Quase metade da receita vira caixa operacional. D&A estável em {fmt(-da[mp])}."),
    (fmt(ebit[mp] - ebit[mp26], sign=True), "EBIT vs MP26", "Por que abaixo do MP26?",
     f"Eventos {fmt(ev_tot[mp] - ev_tot[mp26], sign=True)} e custo fixo {fmt(fix[mp] - fix[mp26], sign=True)}, "
     f"compensados em parte pelo R&D {fmt(rd[mp] - rd[mp26], sign=True)}."),
]
CWID, CH, GX, GY, KW = 450, 98, 12, 9, 104
for i, (kpi, cap, q, a) in enumerate(QA):
    x = 36 + (i % 2) * (CWID + GX)
    y = 100 + (i // 2) * (CH + GY)
    rect(s5, x, y, CWID, CH, WHITE, LIGHT)
    rect(s5, x, y, KW, CH, PANEL)
    rect(s5, x, y, 4, CH, MIDBLUE)
    neg = kpi.startswith("-")
    text(s5, x + 6, y + 18, KW - 10, 30, kpi, size=22, bold=True, color=RED if neg else MIDBLUE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s5, x + 6, y + 52, KW - 10, 28, cap, size=8.5, color=GREY, align=PP_ALIGN.CENTER)
    text(s5, x + KW + 14, y + 12, CWID - KW - 26, 18, q, size=11.5, bold=True, color=NAVY)
    text(s5, x + KW + 14, y + 36, CWID - KW - 26, CH - 42, a, size=9.8, color=DGREY, spacing=1.1)
yr = 100 + 3 * (CH + GY) + 2
for k, (head, col, body) in enumerate([
        ("▲  Riscos", RED,
         f"Eventos concentrados em set–out ({fmt(ev_mp_m[8])}–{fmt(ev_mp_m[9])} mil/mês brutos) e OEM sob "
         "demanda das montadoras."),
        ("▼  Alavancas", GREEN,
         "Mesmo racional de 2026: escalonar itens não obrigatórios se eventos não se confirmarem; cada 10% "
         f"acima da premissa soma {fmt(sens_ev)} de EBIT.")]):
    x = 36 + k * (CWID + GX)
    rect(s5, x, yr, CWID, 520 - yr, PANEL, LIGHT)
    rect(s5, x, yr, 4, 520 - yr, col)
    text(s5, x + 16, yr + 10, CWID - 30, 70, [[(head, {"bold": True, "size": 11.5, "color": col})],
                                              [(body, {"size": 9.8, "color": DGREY})]], spacing=1.1, after=4)
nota(s5, "Slide de apoio para perguntas do presidente. Números da MENS (aba LOCAL), Base Book (histórico) "
     "e abertura do custo fixo informada pela Juliana em 20/10.")

prs.save(OUT)
print("ok", OUT)
print("EBIT R9/MP27", round(ebit[r9], 1), round(ebit[mp], 1), "net", round(net[mp], 1), "fix", round(fix[mp], 1))
print("bridge", [(l.replace(chr(10), ' '), round(v, 1)) for l, v in steps])
print(f"post {post:.1f} (orig {post_orig}, cortado {cortado}, revisto {revisto:.1f}) pontual {pontual / d_fix * 100:.1f}%"
      f" recorrente {recorrente:.1f} ({recorrente / -fix[r9] * 100:.2f}%) ebitda {ebitda:.1f} sem_post {ebit_sem_post:.1f}"
      f" ros_sem {ros_sem:.2f} sem_reaj {ebit_sem_reaj:.1f} sens {sens_ev:.1f}")
print("eventos MP27", [r0(v) for v in ev_mp_m], "total", round(sum(ev_mp_m), 1))
