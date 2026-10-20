"""Refaz o slide 4 (postergações 2026 -> MP'27) da Bridge CP_v7 a partir da tabela que a Juliana ajustou à mão.

Uso:
    python refazer_slide4_ppt_cp_mp27_v7.py <Bridge CP_v7.pptx> <MENS CP 2027 - MP'27.xlsx> <saida.pptx>

Só o slide 4 é tocado. A tabela dela é a fonte: lista de 2026 (611) = itens mantidos no MP'27 (411) + itens cortados (200),
sem "revisão de valores" (os valores da lista já estão atualizados por ela). Linhas em branco da tabela (itens que ela tirou) são ignoradas.
"""
import sys
from lxml import etree
import openpyxl
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE

SRC, MENS, OUT = sys.argv[1:4]

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MIDBLUE = RGBColor(0x2E, 0x5A, 0xA8)
LIGHT = RGBColor(0xDC, 0xE3, 0xF0)
PANEL = RGBColor(0xF3, 0xF5, 0xFA)
GREY = RGBColor(0x80, 0x80, 0x80)
DGREY = RGBColor(0x40, 0x40, 0x40)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(0xC6, 0x28, 0x28)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def fmt(v, sign=False):
    r = int(abs(v) + 0.5)
    s = f"{r:,}".replace(",", ".")
    return ("-" + s) if v < 0 and r else (("+" if sign and r else "") + s)


def p1(x):
    return f"{x:.1f}".replace(".", ",")


# ------------------------------------------------------------------ dados: tabela dela (fonte) + MENS
prs = Presentation(SRC)
s4 = prs.slides[3]
tbl = [sh for sh in s4.shapes if sh.has_table][0].table
ITENS = []                                        # (área, item, lista 2026, MP'27, status)
for r in list(tbl.rows)[1:]:
    c = [ce.text.strip() for ce in r.cells]
    if c[1] and c[1].upper() != "TOTAL" and c[2]:
        ITENS.append((c[0], c[1], float(c[2]), 0.0 if c[3] in ("", "–", "-") else float(c[3]), c[4]))
lista = sum(i[2] for i in ITENS)
mp27 = sum(i[3] for i in ITENS)
mantidos = sorted([i for i in ITENS if i[3]], key=lambda i: -i[3])
cortados = sorted([i for i in ITENS if not i[3]], key=lambda i: -i[2])
cort = sum(i[2] for i in cortados)
assert round(lista) == 611 and round(mp27) == 411 and all(i[2] == i[3] for i in mantidos), (lista, mp27)
assert abs(lista - cort - mp27) < 0.01

loc = openpyxl.load_workbook(MENS, data_only=True)["LOCAL"]
net, ebit = loc["Q21"].value / 1e3, loc["Q40"].value / 1e3
ebit_sem, ros_mp, ros_sem = ebit + mp27, ebit / net * 100, (ebit + mp27) / net * 100
top2 = mantidos[0][3] + mantidos[1][3]


# ------------------------------------------------------------------ helpers
def line(sl, x1, y1, x2, y2, color=GREY, w=0.75, dash=None):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    return c


def text(sl, x, y, w, h, runs, size=9, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bold=False, italic=False, spacing=1.0, after=0):
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    paras = [[(ln, {})] for ln in runs.split("\n")] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        for t, f in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(f.get("size", size))
            r.font.bold = f.get("bold", bold)
            r.font.italic = f.get("italic", italic)
            r.font.color.rgb = f.get("color", color)
    return tb


def rect(sl, x, y, w, h, fill, edge=None, ew=0.75):
    r = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if edge is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = edge
        r.line.width = Pt(ew)
    r.shadow.inherit = False
    return r


def table(sl, x, y, col_w, rows, row_h, header, size=9, col_align=None):
    nr = len(rows) + 1
    gt = sl.shapes.add_table(nr, len(col_w), Pt(x), Pt(y), Pt(sum(col_w)), Pt(row_h * nr))
    t = gt.table
    pr = gt._element.graphic.graphicData.tbl.tblPr
    pr.set("firstRow", "0")
    pr.set("bandRow", "0")
    sid = pr.find(A + "tableStyleId")
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(col_w):
        t.columns[i].width = Pt(w)
    for r_, vals in enumerate([header] + rows):
        t.rows[r_].height = Pt(row_h)
        is_h, is_t = r_ == 0, r_ == nr - 1
        for c_, v in enumerate(vals):
            ce = t.cell(r_, c_)
            ce.margin_top = ce.margin_bottom = Pt(1)
            ce.margin_left = ce.margin_right = Pt(5)
            ce.vertical_anchor = MSO_ANCHOR.MIDDLE
            ce.text = ""
            p = ce.text_frame.paragraphs[0]
            p.alignment = col_align[c_]
            txt, f = (v if isinstance(v, tuple) else (v, {}))
            rr = p.add_run()
            rr.text = str(txt)
            rr.font.size = Pt(f.get("size", size))
            rr.font.bold = f.get("bold", bool(is_h or is_t))
            rr.font.color.rgb = f.get("color", WHITE if is_h else BLACK)
            etree.SubElement(p._p, A + "endParaRPr").set("sz", str(int(f.get("size", size) * 100)))
            if is_h:
                ce.fill.solid()
                ce.fill.fore_color.rgb = NAVY
            elif is_t:
                ce.fill.solid()
                ce.fill.fore_color.rgb = LIGHT
            else:
                ce.fill.background()
        if not (is_h or is_t):
            for c_ in range(len(col_w)):
                tcPr = t.cell(r_, c_)._tc.get_or_add_tcPr()
                ln = etree.Element(A + "lnB", w="3175")
                tcPr.insert(0, ln)                       # lnB tem que vir antes do preenchimento da célula
                sf = etree.SubElement(ln, A + "solidFill")
                etree.SubElement(sf, A + "srgbClr", val="D9D9D9")


def panel_blocks(sl, x, y, w, h, title, items, size=10):
    rect(sl, x, y, w, h, PANEL, LIGHT)
    rect(sl, x, y, 4, h, MIDBLUE)
    text(sl, x + 16, y + 10, w - 28, 20, title, size=13, bold=True, color=NAVY)
    top = y + 38
    bh = (y + h - 8 - top) / len(items)
    for i, (head, body) in enumerate(items):
        by = top + i * bh
        if i:
            line(sl, x + 16, by, x + w - 14, by, color=RGBColor(0xC9, 0xD3, 0xE6), w=0.5)
        text(sl, x + 16, by + 4, w - 30, bh - 8,
             [[(head, {"bold": True, "color": NAVY, "size": size + 0.5})], [(body, {"color": DGREY, "size": size})]],
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.08, after=3)


# ------------------------------------------------------------------ refaz o slide 4 (mantém título e "(+) gain")
for shp in list(s4.shapes):
    if not (shp.is_placeholder or (shp.has_text_frame and shp.text_frame.text.startswith("(+) gain"))):
        shp._element.getparent().remove(shp._element)

text(s4, 36, 62, 900, 40, [[(f"Postergações de 2026: de BRL {fmt(lista)} mil para {fmt(mp27)} mil",
                              {"bold": True, "color": NAVY, "size": 16}),
                             (f" — revisamos a lista e cortamos {fmt(cort)}",
                              {"bold": True, "color": MIDBLUE, "size": 16})]])
text(s4, 36, 100, 520, 16, "Lista de postergações 2026 → MP’27 (BRL ’000)", size=11, bold=True, color=NAVY)
G = {"color": GREY, "size": 8.5}
rows = []
for g, it, o, v, st in mantidos:
    rows.append([(g, G), it, fmt(o), (fmt(v), {"bold": True}), ("Mantido", {"color": NAVY, "bold": True})])
for g, it, o, v, st in cortados:
    gc = {"color": GREY}
    rows.append([(g, G), (it, gc), (fmt(o), gc), ("–", gc), (st, {"color": RED, "bold": True, "size": 8.5})])
rows.append(["", "TOTAL", fmt(lista), (fmt(mp27), {"bold": True, "color": MIDBLUE}),
             (fmt(-cort), {"bold": True, "color": RED})])
table(s4, 36, 120, [74, 200, 56, 56, 126], rows, 30, ["Área", "Item", "Lista 2026", "MP’27", "Status"], size=9,
      col_align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER])

# mini-ponte 611 -> -200 -> 411 (barras em escala real)
bx0, bx1, axis, top = 590, 948, 262, 142
text(s4, bx0, 100, 358, 16, f"De {fmt(lista)} para {fmt(mp27)} (BRL ’000)", size=11, bold=True, color=NAVY)
sc = (axis - top) / lista
cols = [("Lista 2026", 0, lista, NAVY), ("Itens\ncortados", mp27, lista, None), ("MP’27", 0, mp27, NAVY)]
dx = (bx1 - bx0) / len(cols)
bw = dx * 0.5
line(s4, bx0, axis, bx1, axis, color=BLACK, w=0.75)
nivel = None
for i, (lab, a, b, col) in enumerate(cols):
    c = bx0 + dx * (i + 0.5)
    y_top, y_bot = axis - b * sc, axis - a * sc
    if col:
        rect(s4, c - bw / 2, y_top, bw, y_bot - y_top, col)
    else:
        rect(s4, c - bw / 2, y_top, bw, y_bot - y_top, LIGHT, RED, 1.5)
    val = fmt(b) if col else fmt(-cort)
    text(s4, c - 40, y_top - 17, 80, 14, val, size=11 if col else 10, bold=True,
         color=BLACK if col else RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    text(s4, c - dx / 2, axis + 4, dx, 26, lab, size=8.5, bold=bool(col), align=PP_ALIGN.CENTER)
    if i:
        yl = axis - (lista if i == 1 else mp27) * sc
        line(s4, nivel, yl, c - bw / 2, yl, color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
    nivel = c + bw / 2

panel_blocks(s4, 590, 300, 358, 220, "Mensagens-chave", [
    ("2026: alavanca para proteger o resultado",
     "Com a queda de eventos (ano de Copa do Mundo), adiamos estes gastos sem comprometer a operação "
     "e a segurança do Circuito."),
    (f"Cortamos {fmt(cort)} da lista original",
     f"{len(cortados)} itens saíram: pintura do box (já realizada em 2026), reajuste de contratos, viagens, "
     "sanitários/vestiários, TI e pluvial."),
    (f"Ficam {fmt(mp27)}, {top2 / mp27 * 100:.0f}% em Sansão e APP",
     f"Gasto pontual: sem ele, EBIT {fmt(ebit_sem)} e ROS {p1(ros_sem)}% (vs {p1(ros_mp)}%)."),
], size=9)
s4.notes_slide.notes_text_frame.text = (
    "Slide refeito em 20/10 a partir da tabela ajustada pela Juliana no v7: lista de 2026 = 611 = 411 mantidos no MP'27 + "
    f"200 cortados ({', '.join(f'{i[1]} {i[2]:.0f}' for i in cortados)}). Sem revisão de valores: os valores da lista já estão atualizados.\n"
    f"Mantidos: {', '.join(f'{i[1]} {i[3]:.0f}' for i in mantidos)}. EBIT sem postergações = EBIT MENS {ebit:,.1f} + {mp27:.0f}.")
prs.save(OUT)
print("ok", OUT)
print("lista", lista, "mantidos", mp27, "cortados", cort, [i[1] for i in cortados])
