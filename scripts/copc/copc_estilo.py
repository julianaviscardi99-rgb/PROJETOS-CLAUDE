# -*- coding: utf-8 -*-
"""Sistema visual dos cadernos do COPC — o mesmo do MP_Preparação, para os outros blocos.

FONTE DA VERDADE: `bloco_mp.py`. As constantes e o `tabela()` daqui sao copia fiel do que ele ja
usa. Ele NAO importa este modulo de proposito — e codigo validado em producao e nao se reescreve
logica aprovada so para reaproveitar (regra do CLAUDE.md). Quem muda o padrao muda os dois.

Elementos do padrao:
  · sem linhas de grade em nenhuma aba
  · A1 titulo em negrito 15 · A2 subtitulo italico 10 cinza (quando/o que e)
  · faixa de veredito com a cor do farol
  · cabecalho de tabela azul 1F4E79 com texto branco e borda fina BFBFBF
  · farol pintado na celula, com o rotulo em palavra (VERDE / ANALISE / CORRIGIR)
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

AZUL = "1F4E79"
CINZA_BORDA = "BFBFBF"
FAROL_FILL = {"verde": "C6EFCE", "amarelo": "FFEB9C", "vermelho": "FFC7CE",
              "indisponivel": "D9D9D9", "?": "D9D9D9"}
FAROL_TXT = {"verde": "VERDE", "amarelo": "ANALISE", "vermelho": "CORRIGIR",
             "indisponivel": "INDISPON.", "?": "FALTA"}
THIN = Border(*[Side(style="thin", color=CINZA_BORDA)] * 4)
MOEDA = "#,##0.00"


def farol_de(texto):
    """Normaliza os rotulos que os checks do SP usam para as chaves do padrao do MP."""
    t = str(texto or "").strip().upper()
    if t in ("VERDE", "OK", "PRONTO"):
        return "verde"
    if t in ("INFO", "ANALISE", "ANÁLISE", "REVISAR", "AMARELO"):
        return "amarelo"
    if t in ("S/ DADOS", "SEM DADOS", "INDISPONIVEL", "INDISPONÍVEL", ""):
        return "indisponivel"
    return "vermelho"


def nova_aba(wb, nome, primeira=False):
    ws = wb.active if primeira else wb.create_sheet(nome[:31])
    ws.title = nome[:31]
    ws.sheet_view.showGridLines = False
    return ws


def cabecalho(ws, titulo, subtitulo=None):
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=15)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True, size=10, color="595959")


def veredito(ws, celula, texto, farol):
    c = ws[celula]
    c.value = texto
    c.font = Font(bold=True, size=12)
    c.fill = PatternFill("solid", fgColor=FAROL_FILL.get(farol, "FFFFFF"))


def tabela(ws, r0, headers, rows, title=None, subtitle=None, widths=None, numfmt=None,
           farol_col=None):
    """Desenha uma tabela no padrao. Devolve a proxima linha livre.

    farol_col: indice 1-based da coluna cujo valor e um rotulo de farol, para pintar a celula.
    """
    r = r0
    if title:
        c = ws.cell(row=r, column=1, value=title)
        c.font = Font(bold=True, size=12)
        r += 1
    if subtitle:
        c = ws.cell(row=r, column=1, value=subtitle)
        c.font = Font(italic=True, size=9, color="595959")
        r += 1
    if title or subtitle:
        r += 1
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = THIN
    r += 1
    for row in rows:
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = THIN
            if numfmt and j in numfmt and isinstance(v, (int, float)):
                c.number_format = numfmt[j]
            # celula vazia na coluna de farol nao e "sem dado" — e linha que nao tem farol
            # (o resumo do MDO tem duas linhas de composicao antes da linha do delta).
            if farol_col and j == farol_col and v not in (None, ""):
                f = farol_de(v)
                c.value = FAROL_TXT.get(f, str(v))
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = PatternFill("solid", fgColor=FAROL_FILL.get(f, "FFFFFF"))
        r += 1
    if widths:
        from openpyxl.utils import get_column_letter
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r + 1


def nota(ws, linha, texto):
    c = ws.cell(row=linha, column=1, value=texto)
    c.font = Font(italic=True, size=9, color="595959")
    return linha + 1
