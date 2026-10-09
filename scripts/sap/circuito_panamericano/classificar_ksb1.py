"""Sugere a classificacao dos custos da KSB1 do Circuito Panamericano.

Le o arquivo 'KSB1 - Jan-Dez_.xlsx' (somente leitura, trabalha numa copia temporaria),
aprende com os meses ja classificados (Jan-Jul: coluna 'N documento'; Ago em diante:
coluna 'Classificacao') e com a aba 'Parametros', e sugere a classificacao das linhas
do mes escolhido. NUNCA altera o arquivo original: gera um xlsx novo em data/processed/.

Uso:
    python classificar_ksb1.py --arquivo "<caminho>\\KSB1 - Jan-Dez_.xlsx" --mes Set
    python classificar_ksb1.py --arquivo ... --mes Set --avaliar   # teste: compara com o que ja foi classificado
"""
import argparse
import os
import re
import shutil
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
SAIDA_PADRAO = os.path.join(RAIZ, "data", "processed", "circuito_panamericano")


def carregar_regras_manuais():
    import json
    caminho = os.path.join(os.path.dirname(__file__), "regras_manuais.json")
    if not os.path.exists(caminho):
        return []
    with open(caminho, encoding="utf-8") as f:
        return [{"contem": norm(r["contem"]), "classificacao": r["classificacao"]}
                for r in json.load(f).get("regras", [])]


def norm(txt):
    """minusculo, sem acento, sem espacos extras."""
    if txt is None:
        return ""
    s = unicodedata.normalize("NFKD", str(txt))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def val(txt):
    return "" if txt is None else str(txt).strip()


def nome_versionado(pasta, base, ext):
    """Nunca sobrescreve: base.ext, base_v2.ext, base_v3.ext ..."""
    caminho = os.path.join(pasta, base + ext)
    n = 2
    while os.path.exists(caminho):
        caminho = os.path.join(pasta, f"{base}_v{n}{ext}")
        n += 1
    return caminho


def e_numero(txt):
    return bool(re.fullmatch(r"0*\d+", val(txt)))


REGRAS_MANUAIS = carregar_regras_manuais()


class Linha:
    def __init__(self, mes, linha_excel, d, classe_real):
        self.mes = mes
        self.linha = linha_excel
        self.d = d
        self.real = classe_real  # classificacao ja feita pela usuaria (ou "")
        self.classe = val(d.get("classe de custo"))
        self.forn = val(d.get("fornecedor"))
        self.texto = norm(d.get("texto do pedido"))
        self.cab = norm(d.get("texto de cabecalho de documento"))
        self.centro = val(d.get("centro custo"))
        self.valor = d.get("valor/mr") or 0


def ler_mes(wb, mes):
    if mes not in wb.sheetnames:
        return []
    it = wb[mes].iter_rows(values_only=True)
    cab = next(it, None)
    if not cab:
        return []
    nomes = [norm(c) for c in cab]
    # coluna da classificacao: 'classificacao' (Ago em diante) ou 'n documento' (Jan-Jul)
    col_cls = None
    for alvo in ("classificacao", "n documento"):
        for i, n in enumerate(nomes):
            if n == alvo or (alvo == "n documento" and re.fullmatch(r"n.{0,2} documento", n)):
                col_cls = i
                break
        if col_cls is not None:
            break
    out = []
    for n_lin, r in enumerate(it, start=2):
        if all(v is None for v in r):
            continue
        d = {nomes[i]: r[i] for i in range(len(nomes)) if nomes[i]}
        real = ""
        if col_cls is not None and col_cls < len(r):
            v = val(r[col_cls])
            if v and not e_numero(v):
                real = v
        out.append(Linha(mes, n_lin, d, real))
    return out


def ler_parametros(wb):
    """aba Parametros: Descricao (texto do pedido) -> classificacao."""
    for nome in wb.sheetnames:
        if norm(nome).startswith("parametros"):
            m = {}
            for i, r in enumerate(wb[nome].iter_rows(values_only=True)):
                if i == 0 or not r or r[0] is None or len(r) < 2 or r[1] is None:
                    continue
                m[norm(r[0])] = val(r[1])
            return m
    return {}


def construir_historico(linhas):
    h = {k: defaultdict(Counter) for k in ("cft", "cf", "ft", "t", "c")}
    for l in linhas:
        if not l.real:
            continue
        if l.classe or l.forn:
            h["cft"][(l.classe, l.forn, l.texto)][l.real] += 1
            h["cf"][(l.classe, l.forn)][l.real] += 1
        if l.forn:
            h["ft"][(l.forn, l.texto)][l.real] += 1
        if l.texto:
            h["t"][l.texto][l.real] += 1
        if l.classe:
            h["c"][l.classe][l.real] += 1
    return h


def e_rh(l):
    return l.centro.upper() == "HR_DUMMY" or l.classe.upper().endswith("_HR") or "rateio hr" in l.cab


def sugerir(l, h, params):
    """retorna (sugestao, regra, confianca, alternativas)."""
    if e_rh(l):
        return "RH", "Rateio de RH (centro HR_DUMMY / classe _HR)", "alta", ""
    for reg in REGRAS_MANUAIS:
        if reg["contem"] in l.texto:
            return reg["classificacao"], f"Regra manual ({reg['contem']})", "alta", ""
    if not l.classe and not l.forn:
        return "", "Sem classe/fornecedor (DAC / reclassificacao MS) - decisao manual", "revisar", ""

    def unanime(chave, tabela, minimo):
        c = tabela.get(chave)
        if not c:
            return None
        if len(c) == 1:
            cls, n = next(iter(c.items()))
            return cls, n
        return ("amb", c)

    # historico exato tem prioridade sobre a aba Parametros (que pode estar desatualizada)
    r = unanime((l.classe, l.forn, l.texto), h["cft"], 1)
    if r and r[0] != "amb":
        return r[0], f"Classe+Fornecedor+Texto ({r[1]}x no historico)", "alta", ""
    if l.texto and l.texto in params:
        return params[l.texto], "Aba Parametros (texto do pedido)", "alta", ""
    if l.classe or l.forn:
        r = unanime((l.classe, l.forn), h["cf"], 1)
        if r and r[0] != "amb":
            conf = "alta" if r[1] >= 2 else "media"
            return r[0], f"Classe+Fornecedor ({r[1]}x no historico)", conf, ""
    if l.forn:
        r = unanime((l.forn, l.texto), h["ft"], 1)
        if r and r[0] != "amb":
            return r[0], f"Fornecedor+Texto ({r[1]}x)", "media", ""
    if l.texto:
        r = unanime(l.texto, h["t"], 1)
        if r and r[0] != "amb":
            return r[0], f"Texto do pedido ({r[1]}x)", "media", ""
    # ambiguo: lista as opcoes do historico (mais forte primeiro)
    for tab, chave in (("cf", (l.classe, l.forn)), ("c", l.classe)):
        c = h[tab].get(chave) if (l.classe or l.forn) else None
        if c:
            alts = "; ".join(f"{k} ({n}x)" for k, n in c.most_common(4))
            return "", "Ambiguo no historico", "revisar", alts
    return "", "Sem historico", "revisar", ""


def avaliar(linhas_mes, h, params):
    tot = Counter()
    for l in linhas_mes:
        if not l.real:
            continue
        s, _, conf, _ = sugerir(l, h, params)
        tot["linhas"] += 1
        if conf == "revisar":
            tot["revisar"] += 1
        else:
            tot[f"{conf}_total"] += 1
            if s == l.real:
                tot[f"{conf}_ok"] += 1
    return tot


def gravar(saida, mes, linhas_mes, h, params):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classificacao"
    pend = wb.create_sheet("Pendentes")
    res = wb.create_sheet("Resumo")
    cab = ["Linha na aba", "Data lcto", "Centro custo", "Classe de custo", "Descr. classe", "Fornecedor",
           "Nome", "Texto do pedido", "Valor/MR", "Classificacao ja existente", "SUGESTAO", "Confianca",
           "Regra", "Alternativas do historico"]
    for w in (ws, pend):
        w.append(cab)
        for c in w[1]:
            c.font = Font(bold=True)
    cores = {"alta": "C6EFCE", "media": "FFEB9C", "revisar": "FFC7CE"}
    resumo = defaultdict(lambda: [0, 0.0])
    for l in linhas_mes:
        s, regra, conf, alts = sugerir(l, h, params)
        d = l.d
        linha = [l.linha, d.get("data de lancamento"), l.centro, l.classe,
                 val(d.get("descr.classe custo") or d.get("denom.classe custo")), l.forn,
                 val(d.get("nome 1")), val(d.get("texto do pedido")), l.valor, l.real, s, conf, regra, alts]
        for w in ([ws, pend] if conf != "alta" else [ws]):
            w.append(linha)
            w.cell(w.max_row, 12).fill = PatternFill("solid", fgColor=cores[conf])
            if isinstance(linha[1], datetime):
                w.cell(w.max_row, 2).number_format = "dd/mm/yyyy"
            w.cell(w.max_row, 9).number_format = "#,##0.00"
        resumo[conf][0] += 1
        resumo[conf][1] += float(l.valor or 0)
    for w in (ws, pend):
        for i, larg in enumerate([10, 11, 11, 13, 28, 13, 30, 38, 13, 30, 38, 10, 40, 60], start=1):
            w.column_dimensions[get_column_letter(i)].width = larg
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions
    res.append(["Confianca", "Linhas", "Valor liquido (R$)"])
    for conf in ("alta", "media", "revisar"):
        res.append([conf, resumo[conf][0], round(resumo[conf][1], 2)])
    res.append([])
    res.append(["alta = pode confiar, so conferir por amostragem | media = conferir | revisar = decisao sua"])
    wb.save(saida)
    return resumo


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--arquivo", required=True, help="caminho do KSB1 - Jan-Dez_.xlsx")
    ap.add_argument("--mes", required=True, choices=MESES, help="aba do mes a classificar")
    ap.add_argument("--saida", default=SAIDA_PADRAO, help="pasta de saida (padrao: data/processed/circuito_panamericano)")
    ap.add_argument("--avaliar", action="store_true",
                    help="teste: esconde a classificacao do mes e compara a sugestao com o que ja foi feito")
    a = ap.parse_args()

    tmp = tempfile.mkdtemp()
    copia = os.path.join(tmp, "ksb1.xlsx")
    shutil.copyfile(a.arquivo, copia)  # nunca le/escreve direto no original
    try:
        wb = openpyxl.load_workbook(copia, read_only=True, data_only=True)
        todas = {m: ler_mes(wb, m) for m in MESES}
        params = ler_parametros(wb)
        wb.close()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    alvo = todas[a.mes]
    historico_linhas = [l for m, ls in todas.items() if m != a.mes for l in ls]
    h = construir_historico(historico_linhas)
    print(f"Historico: {sum(1 for l in historico_linhas if l.real)} linhas classificadas | "
          f"Parametros: {len(params)} regras | Linhas de {a.mes}: {len(alvo)}")

    if a.avaliar:
        t = avaliar(alvo, h, params)
        n = t["linhas"]
        print(f"\nTESTE em {a.mes}: {n} linhas ja classificadas por voce")
        for conf in ("alta", "media"):
            tt, ok = t[f"{conf}_total"], t[f"{conf}_ok"]
            if tt:
                print(f"  confianca {conf:5s}: {tt:4d} linhas ({tt / n:.0%}) | acerto {ok}/{tt} ({ok / tt:.0%})")
        print(f"  para revisar     : {t['revisar']:4d} linhas ({t['revisar'] / n:.0%})")
        return

    os.makedirs(a.saida, exist_ok=True)
    saida = nome_versionado(a.saida, f"Sugestao_classificacao_{a.mes}_{datetime.now():%Y-%m-%d}", ".xlsx")
    resumo = gravar(saida, a.mes, alvo, h, params)
    print("\nResumo:")
    for conf in ("alta", "media", "revisar"):
        print(f"  {conf:8s}: {resumo[conf][0]:4d} linhas | R$ {resumo[conf][1]:,.2f}")
    print(f"\nArquivo gerado: {saida}")


if __name__ == "__main__":
    main()
