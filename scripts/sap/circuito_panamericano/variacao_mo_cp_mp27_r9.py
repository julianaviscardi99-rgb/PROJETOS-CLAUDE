"""Variacao da MO do Circuito Panamericano: MP27 v1 e v2 x R9 FY26 x run-rate R9 (set-dez x3), por grupo de custo.

Uso: python variacao_mo_cp_mp27_r9.py <labour_mp27_v1.xlsx> <labour_mp27_v2.xlsx> <labour_r9.xlsx> <saida.xlsx>
Somente leitura nas fontes. Valores em R$ mil. CP tem um unico centro de custo (4000): a quebra e por grupo e por categoria (Gerente/Horista).
Grupos iguais aos da Fitted (variacao_mo_r9_runrate_mp27.py) para as apresentacoes ficarem comparaveis.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

f1, f2, f9, out = sys.argv[1:5]
MS = [f"M{i:02d}" for i in range(1, 13)]
RR = ["M09", "M10", "M11", "M12"]

GRUPO = {  # conta -> (grupo apresentacao, grupo detalhe)
    **{c: ("Salarios e PPR", "Salarios + horas extras") for c in ("4200000P", "4200100P", "4201000P", "4201100P")},
    **{c: ("Salarios e PPR", "Provisoes 13o e ferias") for c in ("4200110P", "4200120P", "4201110P", "4201120P")},
    **{c: ("Salarios e PPR", "PPR / PLR") for c in ("4200130P", "4201130P")},
    **{c: ("Encargos", "Encargos") for c in ("4200150P", "4200160P", "4200170P", "4200180P", "4201150P", "4201160P", "4201170P", "4201180P")},
    **{c: ("Assistencia medica", "Assistencia medica") for c in ("4245000P", "4245100P")},
    **{c: ("Outros beneficios", "Outros beneficios") for c in ("4247100P", "4254100P", "4254110P")},
}
G1 = ["Salarios e PPR", "Encargos", "Assistencia medica", "Outros beneficios"]
G2 = ["Salarios + horas extras", "Provisoes 13o e ferias", "PPR / PLR", "Encargos", "Assistencia medica", "Outros beneficios"]


def carregar(path):
    df = pd.read_excel(path, sheet_name="Labour Cost", dtype={"SAP Code": str, "Cost Center": str, "SAP Account": str})
    df = df[df["SAP Code"].notna()].copy()
    cols = [c for c in df.columns if str(c)[:2] == "20" and "." in str(c)]
    df = df.rename(columns={c: f"M{str(c)[-2:]}" for c in cols})
    df["SAP Account"] = df["SAP Account"].str.strip()
    assert df["SAP Account"].isin(GRUPO).all(), df.loc[~df["SAP Account"].isin(GRUPO), "SAP Account"].unique()
    df["G1"] = df["SAP Account"].map(lambda a: GRUPO[a][0])
    df["G2"] = df["SAP Account"].map(lambda a: GRUPO[a][1])
    return df


v1, v2, r9 = carregar(f1), carregar(f2), carregar(f9)
FONTES = [("R9 FY26", r9, MS, 1), ("R9 run-rate", r9, RR, 3), ("MP27 v1", v1, MS, 1), ("MP27 v2", v2, MS, 1)]


def valor(d, gcol, g, meses, mult=1, cat=None):
    s = d[d[gcol] == g]
    if cat:
        s = s[s["Category"] == cat]
    return round(float(s[meses].sum().sum()) * mult / 1e3, 1)


HDR = PatternFill("solid", fgColor="1F3864"); SUB = PatternFill("solid", fgColor="D9E1F2")
NUM = '#,##0;[Red]-#,##0;"-"'; NUM1 = '#,##0.0;[Red]-#,##0.0;"-"'; PCT = '0.0%;[Red]-0.0%;"-"'
BLUE = Font(color="0000FF")
COLS = ["Grupo", "R9 ano", "Run-rate R9", "MP27 v1", "MP27 v2", "Δ v2 vs R9 ano", "Δ % v2 vs R9 ano",
        "Δ v2 vs run-rate", "Δ % v2 vs run-rate", "Δ v2 vs v1", "Δ % v2 vs v1", "Mix v2", "Médio/mês v2"]


def head(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row, i, t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def tabela(ws, r, titulo, gcol, grupos, cat=None):
    ws.cell(r, 1, titulo).font = Font(bold=True, size=12); r += 1
    head(ws, r, COLS); r += 1
    ini = r
    for g in grupos:
        ws.cell(r, 1, g)
        for j, (_, d, mm, mult) in enumerate(FONTES):
            c = ws.cell(r, 2 + j, valor(d, gcol, g, mm, mult, cat)); c.font = BLUE
        r += 1
    tot = r
    ws.cell(tot, 1, "Total MO")
    for col in "BCDE":
        ws[f"{col}{tot}"] = f"=SUM({col}{ini}:{col}{tot-1})"
    for i in range(ini, tot + 1):
        ws[f"F{i}"] = f"=E{i}-B{i}"; ws[f"G{i}"] = f'=IF(B{i}=0,"n/a",F{i}/B{i})'
        ws[f"H{i}"] = f"=E{i}-C{i}"; ws[f"I{i}"] = f'=IF(C{i}=0,"n/a",H{i}/C{i})'
        ws[f"J{i}"] = f"=E{i}-D{i}"; ws[f"K{i}"] = f'=IF(D{i}=0,"n/a",J{i}/D{i})'
        ws[f"L{i}"] = f'=IF(E${tot}=0,"n/a",E{i}/E${tot})'; ws[f"M{i}"] = f"=E{i}/12"
        for col in "BCDEFHJ":
            ws[f"{col}{i}"].number_format = NUM
        ws[f"M{i}"].number_format = NUM1
        for col in "GIKL":
            ws[f"{col}{i}"].number_format = PCT; ws[f"{col}{i}"].alignment = Alignment(horizontal="right")
    for j in range(1, 14):
        ws.cell(tot, j).font = Font(bold=True); ws.cell(tot, j).fill = SUB
    return tot + 3


def larguras(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[L(i)].width = w


wb = Workbook()
NOTA = ("Circuito Panamericano, centro de custo 4000 (único). Run-rate = média set-dez/26 x 12 (R9 tem valores diferentes de 1S26 por entrada tardia de horistas e estornos). "
        "Azul = valor das fontes; preto = fórmula. R$ mil.")

ws = wb.active; ws.title = "Resumo (4 grupos)"
ws["A1"] = "MO Circuito Panamericano - MP27 v1 e v2 x R9 FY26 x run-rate R9 (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = NOTA; ws["A2"].font = Font(italic=True, color="666666")
tabela(ws, 4, "Total CP", "G1", G1)
larguras(ws, [26, 10, 12, 10, 10, 14, 14, 14, 14, 12, 12, 9, 13])
ws.freeze_panes = "B4"

ws = wb.create_sheet("Detalhe (6 grupos)")
ws["A1"] = "MO CP - abertura: salário / provisões / PPR separados (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = NOTA; ws["A2"].font = Font(italic=True, color="666666")
tabela(ws, 4, "Total CP", "G2", G2)
larguras(ws, [26, 10, 12, 10, 10, 14, 14, 14, 14, 12, 12, 9, 13])

ws = wb.create_sheet("Por categoria")
ws["A1"] = "MO CP por categoria: Gerente (mensalista) x Horista (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = NOTA; ws["A2"].font = Font(italic=True, color="666666")
r = 4
for cat in ("GERENTE", "HORISTA"):
    r = tabela(ws, r, cat.capitalize(), "G1", G1, cat)
larguras(ws, [26, 10, 12, 10, 10, 14, 14, 14, 14, 12, 12, 9, 13])

ws = wb.create_sheet("Mensal")
ws["A1"] = "MO CP total por mês (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
head(ws, 3, ["Mês", "R9 2026", "MP27 v1", "MP27 v2", "Δ v2 vs R9", "Δ % v2 vs R9", "Δ v2 vs v1"])
for i, m in enumerate(MS):
    r = 4 + i
    ws.cell(r, 1, f"{i+1:02d}")
    for j, d in enumerate((r9, v1, v2)):
        c = ws.cell(r, 2 + j, round(float(d[m].sum()) / 1e3, 1)); c.font = BLUE
    ws[f"E{r}"] = f"=D{r}-B{r}"; ws[f"F{r}"] = f'=IF(B{r}=0,"n/a",E{r}/B{r})'; ws[f"G{r}"] = f"=D{r}-C{r}"
r = 16
ws.cell(r, 1, "Total")
for col in "BCDEG":
    ws[f"{col}{r}"] = f"=SUM({col}4:{col}15)"
ws[f"F{r}"] = f"=E{r}/B{r}"
for rr in range(4, 17):
    for col in "BCDEG": ws[f"{col}{rr}"].number_format = NUM1
    ws[f"F{rr}"].number_format = PCT; ws[f"F{rr}"].alignment = Alignment(horizontal="right")
for j in range(1, 8): ws.cell(16, j).font = Font(bold=True); ws.cell(16, j).fill = SUB
larguras(ws, [8, 12, 12, 12, 13, 13, 13])

ws = wb.create_sheet("Contas")
head(ws, 1, ["Categoria", "Conta SAP", "Descrição", "Grupo (resumo)", "Grupo (detalhe)", "R9 ano", "Run-rate R9", "MP27 v1", "MP27 v2", "Δ v2 vs R9 ano", "Δ v2 vs run-rate", "Δ v2 vs v1"])
base = pd.concat([v2, v1, r9]).drop_duplicates(["Category", "SAP Account"])[["Category", "SAP Account", "SAP Account Description"]]
base = base.sort_values(["Category", "SAP Account"])
r = 2
for _, b in base.iterrows():
    vals = []
    for _, d, mm, mult in FONTES:
        s = d[(d.Category == b.Category) & (d["SAP Account"] == b["SAP Account"])]
        vals.append(round(float(s[mm].sum().sum()) * mult / 1e3, 1))
    for j, x in enumerate([b.Category, b["SAP Account"], b["SAP Account Description"], GRUPO[b["SAP Account"]][0], GRUPO[b["SAP Account"]][1]] + vals, 1):
        c = ws.cell(r, j, x)
        if j >= 6: c.number_format = NUM1; c.font = BLUE
    ws[f"J{r}"] = f"=I{r}-F{r}"; ws[f"K{r}"] = f"=I{r}-G{r}"; ws[f"L{r}"] = f"=I{r}-H{r}"
    for col in "JKL": ws[f"{col}{r}"].number_format = NUM1
    r += 1
ws.cell(r, 3, "TOTAL").font = Font(bold=True)
for col in "FGHIJKL":
    ws[f"{col}{r}"] = f"=SUM({col}2:{col}{r-1})"; ws[f"{col}{r}"].number_format = NUM1; ws[f"{col}{r}"].font = Font(bold=True)
larguras(ws, [11, 11, 34, 18, 24, 10, 12, 10, 10, 14, 14, 12])
ws.freeze_panes = "D2"

wb.save(out)
print("ok", out)
