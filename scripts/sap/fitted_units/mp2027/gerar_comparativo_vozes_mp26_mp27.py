"""MP2027 - comparativo MP2026 x R9 2026 x MP2027 por VOZ (conta gestorial).

Fonte: aba DataBase_Detail dos arquivos de Detalhe de Despesas:
  - MP26: Detalhe_Despesas_Fitted Units_MP2026_v3.xlsx            (Total C/ Curva, CQ:DB)
  - R9:   Detalhe_Despesas_Fitted Units_Forecast Setembro_final.xlsx (Total C/ Curva, CQ:DB)
  - MP27: Detalhe_Despesas_Fitted Units_Budget'27_V3.xlsx         (Total C/ Curva, BC:BN)
Os totais por unidade fecham com a aba "Resumo Custos" de cada arquivo.

"Voz" = Gestorial + Descrição Gestorial (ex.: 4257000 Aluguéis). Valores em R$ mil,
custo positivo; Δ positivo = custo maior em 2027. O ranking das abas é pela
variação MP27 x R9 (o R9 é a base mais recente do ano corrente).

Uso: python gerar_comparativo_vozes_mp26_mp27.py [MP26.xlsx] [MP27.xlsx] [R9.xlsx] [saida_dir]
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter

BASE = Path(r"\\FSS024-01BR.group.pirelli.com\GFU_DAC\Management Plan")
ARQ26 = BASE / "MP 2026" / "Detalhe_Despesas_Fitted Units_MP2026_v3.xlsx"
ARQ27 = BASE / "MP 2027" / "Detalhe_Despesas_Fitted Units_Budget'27_V3.xlsx"
ARQ9 = Path(r"\\FSS024-01BR.group.pirelli.com\EO_FITTED\BU FITTED\Forecast\Fcst\Fcst 2026\R9 2026"
            r"\Detalhe_Despesas_Fitted Units_Forecast Setembro_final.xlsx")
SAIDA = Path(__file__).resolve().parents[4] / "data" / "processed" / "mp2027"

COLS = {  # (primeira, última) coluna dos 12 meses "Total C/ Curva"
    26: ("CQ", "DB"),
    9: ("CQ", "DB"),
    27: ("BC", "BN"),
}
# índices 0-based na linha
I_CM, I_GEST, I_DESC, I_DETALHE, I_FORN, I_CAT, I_VAR = 0, 2, 3, 7, 9, 13, 23
I_PREMISSA27, I_INDICE27 = CI("AB") - 1, CI("Z") - 1

AZUL = PatternFill("solid", fgColor="1F3864")
CINZA = PatternFill("solid", fgColor="D9D9D9")
VERDE = PatternFill("solid", fgColor="E2EFDA")
VERMELHO = PatternFill("solid", fgColor="FCE4D6")
FMT = '#,##0.0;[Red]-#,##0.0;"-"'
FMT_PCT = '0.0%;[Red]-0.0%;"-"'
UNIDADES = ["SJP", "IBI", "GOI", "RES", "GER"]
ORDEM_CM = {cm: i for i, cm in enumerate(UNIDADES)}


def ler(arq, ano):
    a, b = COLS[ano]
    ws = load_workbook(arq, data_only=True, read_only=True)["DataBase_Detail"]
    linhas = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[I_CM] is None:
            continue
        valor = sum(v for v in r[CI(a) - 1:CI(b)] if isinstance(v, (int, float))) / 1000
        linhas.append({
            "cm": str(r[I_CM]).strip(),
            "gest": str(r[I_GEST]).split(".")[0].strip() if r[I_GEST] is not None else "",
            "desc": str(r[I_DESC] or "").strip(),
            "detalhe": str(r[I_DETALHE] or "").strip(),
            "forn": str(r[I_FORN] or "").strip(),
            "cat": str(r[I_CAT] or "").strip(),
            "var": str(r[I_VAR] or "").strip(),
            "valor": valor,
            "premissa": r[I_PREMISSA27] if ano == 27 and len(r) > I_PREMISSA27 else None,
            "indice": r[I_INDICE27] if ano == 27 and len(r) > I_INDICE27 else None,
        })
    return linhas


def agrega(linhas, chave):
    out = defaultdict(float)
    for l in linhas:
        out[chave(l)] += l["valor"]
    return out


def cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(linha, i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def status(v9, v27):
    """Situação da voz de MP27 em relação ao R9."""
    if abs(v9) < 0.05 and abs(v27) >= 0.05:
        return "Voz nova vs R9"
    if abs(v27) < 0.05 and abs(v9) >= 0.05:
        return "Voz eliminada vs R9"
    return "Aumento" if v27 > v9 else ("Redução" if v27 < v9 else "Igual")


TIT_VALORES = ["MP26 (R$ mil)", "R9 (R$ mil)", "MP27 (R$ mil)",
               "Δ MP27 vs R9", "Δ % vs R9", "Δ MP27 vs MP26", "Δ % vs MP26", "Situação vs R9"]


def formulas_valores(ws, i, n, v26, v9, v27):
    """Colunas n+1..n+8: MP26, R9, MP27, Δ vs R9, %, Δ vs MP26, %, situação."""
    c26, c9, c27, cd9, cd26 = (get_column_letter(n + k) for k in (1, 2, 3, 4, 6))
    ws.cell(i, n + 1, round(v26, 3)).number_format = FMT
    ws.cell(i, n + 2, round(v9, 3)).number_format = FMT
    ws.cell(i, n + 3, round(v27, 3)).number_format = FMT
    ws.cell(i, n + 4, f"={c27}{i}-{c9}{i}").number_format = FMT
    ws.cell(i, n + 5, f'=IF({c9}{i}=0,"n/a",{cd9}{i}/ABS({c9}{i}))').number_format = FMT_PCT
    ws.cell(i, n + 6, f"={c27}{i}-{c26}{i}").number_format = FMT
    ws.cell(i, n + 7, f'=IF({c26}{i}=0,"n/a",{cd26}{i}/ABS({c26}{i}))').number_format = FMT_PCT
    ws.cell(i, n + 8, status(v9, v27))


def aba_comparativa(wb, nome, linhas_cmp, titulos_chave, larguras):
    """linhas_cmp: lista de (chaves..., v26, v9, v27). Δ e Δ% são fórmulas."""
    ws = wb.create_sheet(nome)
    n = len(titulos_chave)
    cabecalho(ws, 1, [*titulos_chave, *TIT_VALORES])
    for i, (*chaves, v26, v9, v27) in enumerate(linhas_cmp, 2):
        for j, k in enumerate(chaves, 1):
            ws.cell(i, j, k)
        formulas_valores(ws, i, n, v26, v9, v27)
    ult = len(linhas_cmp) + 1
    tot = ult + 1
    ws.cell(tot, 1, "TOTAL").font = Font(bold=True)
    for col in (n + 1, n + 2, n + 3, n + 4, n + 6):
        L = get_column_letter(col)
        c = ws.cell(tot, col, f"=SUM({L}2:{L}{ult})")
        c.font = Font(bold=True)
        c.number_format = FMT
        c.fill = CINZA
    for col_pct, col_base, col_d in ((n + 5, n + 2, n + 4), (n + 7, n + 1, n + 6)):
        B, D = get_column_letter(col_base), get_column_letter(col_d)
        c = ws.cell(tot, col_pct, f'=IF({B}{tot}=0,"n/a",{D}{tot}/ABS({B}{tot}))')
        c.number_format = FMT_PCT
        c.font = Font(bold=True)
    ws.freeze_panes = ws.cell(2, n + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(n + 8)}{ult}"
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    return ws


def main():
    a26 = Path(sys.argv[1]) if len(sys.argv) > 1 else ARQ26
    a27 = Path(sys.argv[2]) if len(sys.argv) > 2 else ARQ27
    a9 = Path(sys.argv[3]) if len(sys.argv) > 3 else ARQ9
    saida = Path(sys.argv[4]) if len(sys.argv) > 4 else SAIDA
    l26, l27, l9 = ler(a26, 26), ler(a27, 27), ler(a9, 9)
    todas = l26 + l9 + l27

    desc = {}  # prefere a descrição do MP27, depois R9, depois MP26
    for l in todas:
        desc[l["gest"]] = l["desc"]
    cats = defaultdict(set)
    for l in todas:
        cats[l["gest"]].add(l["cat"])

    def tripla(chave):
        k26, k9, k27 = agrega(l26, chave), agrega(l9, chave), agrega(l27, chave)
        return k26, k9, k27, set(k26) | set(k9) | set(k27)

    def dv9(x):  # x = (..., v26, v9, v27)
        return -abs(x[-1] - x[-2])

    # 1) voz consolidada
    k26, k9, k27, chaves = tripla(lambda l: l["gest"])
    vozes = sorted(((g, desc[g], " / ".join(sorted(cats[g])), k26.get(g, 0), k9.get(g, 0), k27.get(g, 0))
                    for g in chaves), key=dv9)

    # 2) voz por unidade
    u26, u9, u27, chaves = tripla(lambda l: (l["cm"], l["gest"]))
    por_un = sorted(((k[0], k[1], desc[k[1]], u26.get(k, 0), u9.get(k, 0), u27.get(k, 0)) for k in chaves),
                    key=lambda x: (ORDEM_CM.get(x[0], 9), dv9(x)))
    por_un_rank = sorted(por_un, key=dv9)

    # 3) categoria por unidade
    c26, c9, c27, chaves = tripla(lambda l: (l["cm"], l["cat"], l["var"]))
    por_cat = sorted(((k[0], k[1], "Variável" if k[2] == "V" else "Fixo",
                       c26.get(k, 0), c9.get(k, 0), c27.get(k, 0)) for k in chaves),
                     key=lambda x: (ORDEM_CM.get(x[0], 9), dv9(x)))

    # 4) linhas de detalhe
    chave_d = lambda l: (l["cm"], l["gest"], l["detalhe"], l["forn"])  # noqa: E731
    d26, d9, d27, chaves = tripla(chave_d)
    prem = {chave_d(l): (l["premissa"], l["indice"]) for l in l27}
    det = sorted(chaves, key=lambda k: -abs(d27.get(k, 0) - d9.get(k, 0)))[:80]

    wb = Workbook()
    wb.remove(wb.active)

    # ---- Resumo
    ws = wb.create_sheet("Resumo")
    ws["A1"] = "Fitted Units - Comparativo de despesas por voz: MP2026 x R9 2026 x MP2027"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("R$ mil · custo positivo · Δ positivo = custo maior no MP27 · valores finais (Total C/ Curva) · "
                "ranking pela variação MP27 vs R9")
    ws["A2"].font = Font(italic=True)
    cabecalho(ws, 4, ["Unidade", "MP26", "R9", "MP27", "Δ vs R9", "Δ % vs R9", "Δ vs MP26", "Δ % vs MP26"])
    for i, cm in enumerate(UNIDADES, 5):
        ws.cell(i, 1, cm)
        for col, d in ((2, c26), (3, c9), (4, c27)):
            ws.cell(i, col, round(sum(v for k, v in d.items() if k[0] == cm), 3))
        ws.cell(i, 5, f"=D{i}-C{i}")
        ws.cell(i, 6, f'=IF(C{i}=0,"n/a",E{i}/ABS(C{i}))')
        ws.cell(i, 7, f"=D{i}-B{i}")
        ws.cell(i, 8, f'=IF(B{i}=0,"n/a",G{i}/ABS(B{i}))')
    t = 5 + len(UNIDADES)
    ws.cell(t, 1, "TOTAL FITTED UNITS").font = Font(bold=True)
    for col in "BCDEG":
        ws[f"{col}{t}"] = f"=SUM({col}5:{col}{t - 1})"
        ws[f"{col}{t}"].font = Font(bold=True)
        ws[f"{col}{t}"].fill = CINZA
    ws[f"F{t}"] = f"=E{t}/ABS(C{t})"
    ws[f"H{t}"] = f"=G{t}/ABS(B{t})"
    for r in range(5, t + 1):
        for col in "BCDEG":
            ws[f"{col}{r}"].number_format = FMT
        for col in "FH":
            ws[f"{col}{r}"].number_format = FMT_PCT

    def top(titulo, itens, linha, cor):
        ws.cell(linha, 1, titulo).font = Font(bold=True, size=12)
        cabecalho(ws, linha + 1, ["Unidade", "Gestorial", "Voz", "MP26", "R9", "MP27", "Δ vs R9", "Δ % vs R9"])
        for i, (cm, g, d, a, r9, b) in enumerate(itens, linha + 2):
            for j, v in enumerate((cm, g, d), 1):
                ws.cell(i, j, v)
            for j, v in ((4, a), (5, r9), (6, b)):
                ws.cell(i, j, round(v, 3)).number_format = FMT
            ws.cell(i, 7, f"=F{i}-E{i}").number_format = FMT
            ws.cell(i, 8, f'=IF(E{i}=0,"n/a",G{i}/ABS(E{i}))').number_format = FMT_PCT
            for c in range(1, 9):
                ws.cell(i, c).fill = cor
        return linha + 2 + len(itens)

    aum = [x for x in por_un_rank if x[5] > x[4]][:10]
    red = [x for x in por_un_rank if x[5] < x[4]][:10]
    prox = top("10 maiores AUMENTOS vs R9 (unidade x voz)", aum, t + 2, VERMELHO)
    top("10 maiores REDUÇÕES vs R9 (unidade x voz)", red, prox + 1, VERDE)
    for col, w in zip("ABCDEFGH", (24, 12, 38, 13, 13, 13, 13, 12)):
        ws.column_dimensions[col].width = w

    # ---- comparativas
    aba_comparativa(wb, "Vozes (total)", vozes, ["Gestorial", "Voz", "Categoria"], (12, 40, 30, *[14] * 8))
    aba_comparativa(wb, "Vozes por unidade", por_un, ["Unidade", "Gestorial", "Voz"], (10, 12, 40, *[14] * 8))
    aba_comparativa(wb, "Categorias", por_cat, ["Unidade", "Categoria", "Tipo"], (10, 24, 12, *[14] * 8))

    # ---- Detalhe Serviço/Produto (por centro de montagem e consolidado)
    def norm(s):
        return " ".join(s.split()).upper()

    nome_serv = {}  # prefere o texto do MP27, depois R9, depois MP26
    for l in todas:
        nome_serv[norm(l["detalhe"])] = l["detalhe"] or "(sem detalhe)"
    s26, s9, s27, chaves = tripla(lambda l: (l["cm"], norm(l["detalhe"])))
    serv_cm = sorted(((k[0], nome_serv[k[1]], s26.get(k, 0), s9.get(k, 0), s27.get(k, 0)) for k in chaves),
                     key=lambda x: (ORDEM_CM.get(x[0], 9), dv9(x)))
    t26, t9, t27, chaves = tripla(lambda l: norm(l["detalhe"]))
    serv_tot = sorted(((nome_serv[k], t26.get(k, 0), t9.get(k, 0), t27.get(k, 0)) for k in chaves), key=dv9)
    aba_comparativa(wb, "Serviço por CM", serv_cm, ["Centro de montagem", "Detalhe serviço/produto"], (12, 55, *[14] * 8))
    aba_comparativa(wb, "Serviço (total)", serv_tot, ["Detalhe serviço/produto"], (60, *[14] * 8))

    # ---- linhas de detalhe
    ws = wb.create_sheet("Maiores linhas")
    cabecalho(ws, 1, ["Unidade", "Gestorial", "Voz", "Serviço / produto", "Fornecedor", *TIT_VALORES[:7],
                      "Premissa MP27", "Índice de reajuste"])
    for i, k in enumerate(det, 2):
        cm, g, dt, fo = k
        for j, v in enumerate([cm, g, desc[g], dt, fo], 1):
            ws.cell(i, j, v)
        formulas_valores(ws, i, 5, d26.get(k, 0), d9.get(k, 0), d27.get(k, 0))
        ws.cell(i, 13).value = None  # situação não se aplica a esta aba
        p, ix = prem.get(k, (None, None))
        ws.cell(i, 13, p)
        ws.cell(i, 14, ix)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:N{len(det) + 1}"
    for i, w in enumerate((9, 11, 32, 40, 34, 12, 12, 12, 12, 10, 12, 10, 22, 16), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32

    # ---- notas
    ws = wb.create_sheet("Notas")
    t26, t9, t27 = (sum(l["valor"] for l in x) for x in (l26, l9, l27))
    notas = [
        "Fontes",
        f"  MP26: {a26.name} (aba DataBase_Detail, 'Total C/ Curva' Jan-Dez)",
        f"  R9:   {a9.name} (aba DataBase_Detail, 'Total C/ Curva' Jan-Dez; Jan-Ago realizado/estimado + forecast do R9)",
        f"  MP27: {a27.name} (aba DataBase_Detail, 'Total C/ Curva' Jan-Dez)",
        f"Conferência: o total fecha com a aba 'Resumo Custos' de cada arquivo (MP26 {t26:,.1f} / R9 {t9:,.1f} / MP27 {t27:,.1f} R$ mil).",
        "",
        "Definições",
        "  Voz = Gestorial + Descrição Gestorial. Δ = MP27 − referência (positivo = custo maior). Δ % sobre o valor absoluto da referência.",
        "  'Voz nova / eliminada vs R9' = existe só em um dos dois (inclui mudança de classificação entre vozes).",
        "  Abas ordenadas pela variação MP27 vs R9. Aluguéis negativos na IBI vêm da própria base - não alterado.",
        "",
        "Atenção",
        "  RES (Resende) não existe no MP26: todo o valor aparece como aumento/voz nova nessa comparação.",
        "  GER = rateio da Gerência.",
        "  Os arquivos de MP26 e MP27 foram alterados em 06/10/2026 (16:37-16:40); o comparativo reflete essa leitura.",
        "  Valores do MP27 ainda em revisão (mão de obra e depreciação provisórias; premissas por linha na aba 'Maiores linhas').",
    ]
    for i, n in enumerate(notas, 1):
        ws.cell(i, 1, n)
        if n and not n.startswith(" ") and not n.startswith("Conferência"):
            ws.cell(i, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 140

    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / "Comparativo_Vozes_MP26_R9_MP27.xlsx"
    wb.save(destino)
    print("Arquivo gerado:", destino)
    print(f"Total MP26 {t26:,.1f} | R9 {t9:,.1f} | MP27 {t27:,.1f} (R$ mil)")
    for cm in UNIDADES:
        print(f"  {cm}: MP26 {sum(v for k, v in c26.items() if k[0] == cm):9,.1f} "
              f"R9 {sum(v for k, v in c9.items() if k[0] == cm):9,.1f} "
              f"MP27 {sum(v for k, v in c27.items() if k[0] == cm):9,.1f}")
    print("Top 12 vozes por variacao MP27 vs R9:")
    for g, d, c, a, r9, b in vozes[:12]:
        print(f"  {g} {d[:34]:34} MP26 {a:9,.1f} R9 {r9:9,.1f} MP27 {b:9,.1f}  var {b - r9:+9,.1f}")


if __name__ == "__main__":
    main()
