"""Gera o bridge de EBIT do MP'27 (vs R9 e vs MP'26) com vinculo externo real para o MENS FITTED.

Uso: python gerar_bridge_ebit_linkado.py <saida.xlsx> <copia_local_do_MENS.xlsx> [nome_do_MENS_na_pasta]
O vinculo e relativo: a saida deve ficar na mesma pasta do MENS. A copia local so serve para o cache.
"""
import os, openpyxl, re, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.external_link.external import (ExternalLink, ExternalBook, ExternalSheetNames,
    ExternalSheetDataSet, ExternalSheetData, ExternalRow, ExternalCell)
from openpyxl.packaging.relationship import Relationship

OUT = sys.argv[1]
SRC_LOCAL = sys.argv[2]
SRC_NAME = sys.argv[3] if len(sys.argv) > 3 else os.path.basename(SRC_LOCAL)
src = openpyxl.load_workbook(SRC_LOCAL, data_only=True)
sheets = src.sheetnames
refs = {}


def X(sheet, addr):
    a = addr.replace('$', '')
    refs[(sheet, a)] = src[sheet][a].value
    sh = f"'[1]{sheet}'" if ' ' in sheet else f"[1]{sheet}"
    col = re.match(r'[A-Z]+', a).group()
    row = a[len(col):]
    return f"{sh}!${col}${row}"


wb = openpyxl.Workbook()
F = "Calibri"
hdr = PatternFill("solid", fgColor="1F3864")
sub = PatternFill("solid", fgColor="D9E1F2")
inp = PatternFill("solid", fgColor="FFF2CC")
tot = PatternFill("solid", fgColor="E2EFDA")
NUM = '#,##0.0;[Red]-#,##0.0;"-"'
EBIT27 = "EBIT MP'27"


def put(ws, a, v, bold=False, fill=None, color=None, fmt=None, align=None, italic=False, size=10):
    c = ws[a]
    c.value = v
    c.font = Font(name=F, bold=bold, color=color, italic=italic, size=size)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    return c


# ---------------- Premissas ----------------
pr = wb.active
pr.title = "Premissas"
pr.sheet_view.showGridLines = False
put(pr, "B2", "Premissas do bridge de EBIT — MP'27", True, size=14)
put(pr, "B3", f"Arquivo-fonte (vinculado): {SRC_NAME}, na mesma pasta. Células amarelas = digitar; o resto é fórmula.", italic=True)


def band(r, text):
    put(pr, f"B{r}", text, True, fill=sub)
    for c in "CDE":
        pr[f"{c}{r}"].fill = sub


band(5, "1) Transporte — nova regra ANTT (IBI)")
put(pr, "B6", "Frete unitário ANTES da regra (R$/pç Fiat, líquido de 9,25%)")
put(pr, "D6", -1.57073125, fill=inp, fmt='0.0000')
put(pr, "E6", "Valor da v2 (-4.158,3 mil / 2.647.379 pç). Ajustar se a base (R9 / MP'26) tiver outro frete.", italic=True)
put(pr, "B7", "Frete unitário ATUAL (R$/pç Fiat) — vem da aba IBI")
put(pr, "D7", f"={X('IBI', 'Q79')}/{X('IBI', 'Q53')}", fmt='0.0000')
put(pr, "B8", "Variação do frete unitário (%)")
put(pr, "D8", "=D7/D6-1", fmt='0.0%')
put(pr, "B9", "Volume Fiat MP'27 (pç)")
put(pr, "D9", f"={X('IBI', 'Q53')}", fmt='#,##0')
put(pr, "B10", "Efeito ANTT no EBIT (R$ mil) = vol. Fiat × (frete atual − anterior)", True)
put(pr, "D10", "=D9*(D7-D6)/1000", True, fill=tot, fmt=NUM)

band(12, "2) Jaguar (JLR) — rateio do custo fixo de Ibirité")
put(pr, "D13", "vs R9", True, align="center")
put(pr, "E13", "vs MP'26", True, align="center")
put(pr, "B14", "Custo fixo da base atribuído à Jaguar que SAI com ela (R$ mil; positivo = melhora o EBIT)")
put(pr, "D14", "=IF(ISNUMBER(D17),D17+D15,0)", fmt=NUM)
put(pr, "E14", "=IF(ISNUMBER(E17),E17+E15,0)", fmt=NUM)
put(pr, "B15", "Margem direta da Jaguar na base (R$ mil) — aba Impacto JLR")
put(pr, "D15", f"={X('Impacto JLR', 'D15')}", fmt=NUM)
put(pr, "E15", f"={X('Impacto JLR', 'E15')}", fmt=NUM)
put(pr, "B16", "Efeito total da saída da Jaguar no EBIT (R$ mil) = −margem direta + rateio fixo", True)
put(pr, "D16", "=-D15+D14", True, fill=tot, fmt=NUM)
put(pr, "E16", "=-E15+E14", True, fill=tot, fmt=NUM)
put(pr, "B17", "Efeito da perda de volume da Jaguar informado (R$ mil, negativo = piora o EBIT)")
put(pr, "D17", -550, fill=inp, fmt=NUM)
put(pr, "E17", None, fill=inp, fmt=NUM)
put(pr, "B18", "Informado vs R9 = -550 (Fcst R9). D14 = informado + margem direta (a diferença vs a margem calculada cai na linha de rateio fixo). Quando tiver o rateio de Ibirité, digitar direto em D14/E14 (E17 vazio = 0 vs MP'26).", italic=True)
pr["B18"].alignment = Alignment(wrap_text=True, vertical="top")
pr.row_dimensions[18].height = 42

band(19, "3) Como ler o bridge")
notes = [
    "Volume = efeito de volume à margem de contribuição unitária da base (coluna T/W de cada aba). Na IBI, é o da aba menos o volume da Jaguar.",
    "Mix e Preço: só a IBI tem mix (Fiat/Iveco/CNH), vindo da aba Impacto JLR. SJP, GO e RES têm preço único, então Mix = 0.",
    "Preço = variação do preço médio × volume (cada unidade identifica seus ganhos/reajustes; a base é o preço de lista).",
    "Jaguar = −margem direta (faturamento − rodas − materiais − frete) + rateio fixo (premissa acima). É retirada das linhas de volume, mix e variável, sem dupla contagem.",
    "ANTT = efeito do novo frete por peça sobre o volume Fiat da IBI; é retirado do custo variável.",
    "Demais variáveis = variação de MO variável, handling, materiais, transporte (ex-ANTT) e outros variáveis, a volume constante (coluna T/W).",
    "Demais fixos = variação dos custos fixos (MO, depreciação, IFRS16, aluguel, condomínio, outros), exceto o rateio da Jaguar.",
    "Outros = preço adicional + outros + other income. Conferências: base + efeitos = EBIT MP'27 e efeitos = variante da aba (Q49/Q50).",
    "Atualização: abrir este arquivo e aceitar 'Atualizar vínculos'. Se a v3 mudar de nome: Dados > Editar Vínculos > Alterar origem.",
]
for i, t in enumerate(notes):
    put(pr, f"B{20 + i}", "• " + t)
    pr[f"B{20 + i}"].alignment = Alignment(wrap_text=True, vertical="top")
    pr.row_dimensions[20 + i].height = 28
pr.column_dimensions["A"].width = 2
pr.column_dimensions["B"].width = 80
pr.column_dimensions["C"].width = 2
pr.column_dimensions["D"].width = 14
pr.column_dimensions["E"].width = 60

# ---------------- Bridge ----------------
bw = wb.create_sheet("Bridge", 0)
bw.sheet_view.showGridLines = False
put(bw, "B2", "Bridge de EBIT — MP'27 (R$ mil)", True, size=14)
put(bw, "B3", "Todos os valores são fórmulas ligadas à v3 teste daniel. Positivo = melhora o EBIT.", italic=True)
COLS = {"SJP": "C", "IBI": "D", "GO": "E", "RES": "F"}
L_BASE, L_VOL, L_MIX, L_PRC = "EBIT base", "Volume", "Mix", "Preço"
L_JMD, L_JFX, L_ANTT = "Jaguar — margem direta", "Jaguar — rateio custo fixo Ibirité", "Transporte ANTT (IBI)"
L_VAR, L_FIX, L_OUT = "Demais custos variáveis", "Demais custos fixos", "Outros"
L_DELTA = "Δ EBIT total"
L_CK1 = "Conferência: base + efeitos − EBIT MP'27"
L_CK2 = "Conferência: efeitos − variante da aba"
LINES = [L_BASE, L_VOL, L_MIX, L_PRC, L_JMD, L_JFX, L_ANTT, L_VAR, L_FIX, L_OUT, EBIT27, L_DELTA, L_CK1, L_CK2]
STRONG = (L_BASE, EBIT27, L_DELTA)
ORIGEM = {L_VOL: "aba (T/W8) / Impacto JLR", L_MIX: "Impacto JLR (só IBI)", L_PRC: "aba (T/W16) / Impacto JLR",
          L_JMD: "Impacto JLR", L_JFX: "Premissas (digitar)", L_ANTT: "Premissas",
          L_VAR: "aba, linha 18", L_FIX: "aba, linha 31", L_OUT: "aba, linhas 12+13+41"}


def block(r0, base, b, t, jr, er, mdc, pcol, qrow):
    put(bw, f"B{r0}", f"MP'27 vs {base}", True, fill=hdr, color="FFFFFF")
    for u, c in COLS.items():
        put(bw, f"{c}{r0}", u, True, fill=hdr, color="FFFFFF", align="center")
    put(bw, f"G{r0}", "TOTAL", True, fill=hdr, color="FFFFFF", align="center")
    put(bw, f"H{r0}", "Origem / status", True, fill=hdr, color="FFFFFF")
    R = {n: r0 + 1 + i for i, n in enumerate(LINES)}
    for n in LINES:
        put(bw, f"B{R[n]}", n, bold=n in STRONG)
    for u, c in COLS.items():
        f = {}
        f[L_BASE] = f"={X(u, b + '43')}"
        if u == "IBI":
            f[L_VOL] = f"={X(u, t + '8')}-{X('Impacto JLR', 'K' + str(jr))}"
            f[L_MIX] = f"={X('Impacto JLR', 'L' + str(er))}"
            f[L_PRC] = f"={X('Impacto JLR', 'M' + str(er))}"
            f[L_JMD] = f"=-{X('Impacto JLR', mdc + '15')}"
            f[L_JFX] = f"=Premissas!{pcol}14"
            f[L_ANTT] = "=Premissas!$D$10"
            econ = f"({X('Impacto JLR', mdc + '8')}-{X('Impacto JLR', mdc + '15')})"
            f[L_VAR] = f"={X(u, t + '18')}-{econ}-{c}{R[L_ANTT]}"
            f[L_FIX] = f"={X(u, t + '31')}-{c}{R[L_JFX]}"
        else:
            f[L_VOL] = f"={X(u, t + '8')}"
            f[L_MIX] = 0
            f[L_PRC] = f"={X(u, t + '16')}"
            f[L_JMD] = 0
            f[L_JFX] = 0
            f[L_ANTT] = 0
            f[L_VAR] = f"={X(u, t + '18')}"
            f[L_FIX] = f"={X(u, t + '31')}"
        f[L_OUT] = f"={X(u, t + '12')}+{X(u, t + '13')}+{X(u, t + '41')}"
        f[EBIT27] = f"={X(u, 'Q43')}"
        f[L_DELTA] = f"={c}{R[EBIT27]}-{c}{R[L_BASE]}"
        eff = f"SUM({c}{R[L_VOL]}:{c}{R[L_OUT]})"
        f[L_CK1] = f"=ROUND({c}{R[L_BASE]}+{eff}-{c}{R[EBIT27]},3)"
        f[L_CK2] = f"=ROUND({eff}-{X(u, 'Q' + qrow)},3)"
        for n, v in f.items():
            put(bw, f"{c}{R[n]}", v, fmt=NUM)
    for n in LINES:
        r = R[n]
        if n.startswith("Conferência"):
            put(bw, f"G{r}", f"=ROUND(SUM(C{r}:F{r}),3)", fmt=NUM)
        else:
            put(bw, f"G{r}", f"=SUM(C{r}:F{r})", bold=n in STRONG, fmt=NUM)
    for n in STRONG:
        for c in "BCDEFGH":
            bw[f"{c}{R[n]}"].fill = tot
            bw[f"{c}{R[n]}"].font = Font(name=F, bold=True, size=10)
    for n in (L_CK1, L_CK2):
        r = R[n]
        put(bw, f"H{r}", f'=IF(AND(C{r}=0,D{r}=0,E{r}=0,F{r}=0,G{r}=0),"OK","VERIFICAR")', True)
    dif = f'ROUND(G{R[L_BASE]}-{X("TOTAL", b + "43")},3)+ROUND(G{R[EBIT27]}-{X("TOTAL", "Q43")},3)'
    put(bw, f"H{R[L_DELTA]}",
        f'=IF({dif}=0,"Soma das unidades = aba TOTAL","ATENÇÃO: aba TOTAL difere em "&TEXT({dif},"#,##0.0")&" (base digitada)")',
        italic=True)
    for n, tx in ORIGEM.items():
        put(bw, f"H{R[n]}", tx, italic=True, color="7F7F7F")


block(5, "R9", "S", "T", 24, 27, "D", "D", "49")
block(22, "MP'26", "V", "W", 36, 39, "E", "E", "50")


put(bw, "B39", "Alertas", True, fill=sub)
for c in "CDEFGH":
    bw[f"{c}39"].fill = sub
alerts = [
    "R9: o EBIT base da aba TOTAL (S43 digitado) é maior que a soma das 4 unidades; a diferença está em agosto (TOTAL!L47). Este bridge usa a soma das unidades.",
    "MP'26 da RES: a aba só tem EBIT (337,5), sem volume, preço ou custos na coluna V. Por isso, vs MP'26, o preço da RES mostra a receita inteira e Outros mostra -337,5.",
    "Mix: só calculado para a IBI (clientes Fiat/Iveco/CNH). Nas demais unidades o preço é único.",
]
for i, t in enumerate(alerts):
    put(bw, f"B{40 + i}", "• " + t)

# ---------------- Detalhe custos ----------------
dw = wb.create_sheet("Detalhe custos")
dw.sheet_view.showGridLines = False
put(dw, "B2", "Variação de custo por linha (R$ mil) — direto das abas; na IBI inclui ANTT e Jaguar", True, size=13)
VAR = [(19, "Labour (MO variável)"), (20, "Handling"), (21, "Direct Materials"), (22, "Transportation"), (23, "Other Variable")]
FIX = [(32, "Labour (MO fixa)"), (33, "Depreciation"), (34, "IFRS16"), (35, "Rents"), (36, "Condominio"), (37, "Other Fixed")]
r = 4
for base, t in (("R9", "T"), ("MP'26", "W")):
    put(dw, f"B{r}", f"MP'27 vs {base}", True, fill=hdr, color="FFFFFF")
    for u, c in COLS.items():
        put(dw, f"{c}{r}", u, True, fill=hdr, color="FFFFFF", align="center")
    put(dw, f"G{r}", "TOTAL", True, fill=hdr, color="FFFFFF", align="center")
    r += 1
    for title, grp in (("Custos variáveis", VAR), ("Custos fixos", FIX)):
        put(dw, f"B{r}", title, True, fill=sub)
        for c in "CDEFG":
            dw[f"{c}{r}"].fill = sub
        r += 1
        first = r
        for row, lab in grp:
            put(dw, f"B{r}", lab)
            for u, c in COLS.items():
                put(dw, f"{c}{r}", f"={X(u, t + str(row))}", fmt=NUM)
            put(dw, f"G{r}", f"=SUM(C{r}:F{r})", fmt=NUM)
            r += 1
        put(dw, f"B{r}", "Subtotal", True, fill=tot)
        for c in "CDEFG":
            put(dw, f"{c}{r}", f"=SUM({c}{first}:{c}{r - 1})", True, fill=tot, fmt=NUM)
        r += 2
    r += 1
dw.column_dimensions["A"].width = 2
dw.column_dimensions["B"].width = 34
for c in "CDEFG":
    dw.column_dimensions[c].width = 12
bw.column_dimensions["A"].width = 2
bw.column_dimensions["B"].width = 44
for c in "CDEFG":
    bw.column_dimensions[c].width = 12
bw.column_dimensions["H"].width = 38
bw.freeze_panes = "C5"

# ---------------- vínculo externo com cache ----------------
by = {}
for (s, a), v in refs.items():
    by.setdefault(s, {})[a] = v
sdata = []
for i, s in enumerate(sheets):
    rowsd = {}
    for a, v in by.get(s, {}).items():
        col = re.match(r'[A-Z]+', a).group()
        rn = int(a[len(col):])
        if isinstance(v, (int, float)):
            cell = ExternalCell(r=a, v=repr(float(v)))
        elif v is None:
            continue
        else:
            cell = ExternalCell(r=a, t="str", v=str(v))
        rowsd.setdefault(rn, []).append(cell)
    sdata.append(ExternalSheetData(sheetId=i, row=[ExternalRow(r=k, cell=v) for k, v in sorted(rowsd.items())]))
book = ExternalBook(sheetNames=ExternalSheetNames(sheetName=sheets),
                    sheetDataSet=ExternalSheetDataSet(sheetData=sdata), id="rId1")
link = ExternalLink(externalBook=book)
link.file_link = Relationship(type="externalLinkPath", Target=SRC_NAME.replace(" ", "%20"),
                              TargetMode="External", Id="rId1")
wb._external_links.append(link)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("ok", len(refs), "refs")
