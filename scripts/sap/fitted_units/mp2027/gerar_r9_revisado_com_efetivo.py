"""MP2027 - R9 REVISADO: efetivo Jan-Set/2026 (KSB1 Flash) + Out-Dez pela premissa do R9.

Pedido da usuária (2026-10-07): o Detalhe de Despesas do R9 (Forecast Setembro) não tem
o efetivo de Jul e Ago. Trocar Jan-Set pelo efetivo do "KSB1 September Flash 2026.xlsx"
(aba BASE_KSB1, Gestorial e Centro de Montagem já resolvidos por fórmula) e manter
Out-Dez como está no R9 ("Total C/ Curva" da DataBase_Detail). Compara com o MP27 (V3)
por unidade e Gestorial, para explicar a variação vs R9.

Escopo: só as vozes (Gestoriais) que existem no Detalhe. O KSB1 traz TODA a despesa
(mão de obra etc.); o que não está no Detalhe fica na aba "Fora do Detalhe".

Uso: gerar_r9_revisado_com_efetivo.py [KSB1] [R9] [MP27] [pasta_saida] [ultimo_mes_efetivo]
Valores em R$ mil, custo positivo.
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as CL

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "_shared"))
from ksb1_core import nome_com_versao  # noqa: E402

REDE_CUSTOS = Path(r"\\FSS024-01BR.group.pirelli.com\GFU_DAC")
ARQ_KSB1 = (REDE_CUSTOS / "Custos Fitted Units" / "Resultados Fitted" / "2026" / "09 - Sep"
            / "09_Sep_Flash" / "KSB1 September Flash 2026.xlsx")
ARQ_R9 = Path(r"\\FSS024-01BR.group.pirelli.com\EO_FITTED\BU FITTED\Forecast\Fcst\Fcst 2026\R9 2026"
              r"\Detalhe_Despesas_Fitted Units_Forecast Setembro_final.xlsx")
ARQ_MP27 = (REDE_CUSTOS / "Management Plan" / "MP 2027"
            / "Detalhe_Despesas_Fitted Units_Budget'27_V3.xlsx")
SAIDA = Path(__file__).resolve().parents[4] / "data" / "processed" / "mp2027"

COL_DETALHE = {"R9": ("CQ", "DB"), "MP27": ("BC", "BN")}
I_CM, I_GEST, I_DESC = 0, 2, 3
# BASE_KSB1 (índices 0-based da linha)
K_VALOR, K_MES, K_GEST, K_DESC, K_CM = 16, 18, 19, 20, 21

UNIDADES = ["SJP", "IBI", "GOI", "RES", "GER"]
MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

AZUL = PatternFill("solid", fgColor="1F3864")
CINZA = PatternFill("solid", fgColor="D9D9D9")
EFET = PatternFill("solid", fgColor="E2EFDA")
PREM = PatternFill("solid", fgColor="FFF2CC")
FMT = '#,##0.0;[Red]-#,##0.0;"-"'


def cod(g):
    """Gestorial como texto sem decimal ('4257000')."""
    if g is None:
        return ""
    try:
        return str(int(float(g)))
    except (TypeError, ValueError):
        return str(g).split(".")[0].strip()


def ler_detalhe(arq, versao):
    """(cm, gestorial) -> (descrição, [12 meses em R$ mil])."""
    a, b = COL_DETALHE[versao]
    ws = load_workbook(arq, data_only=True, read_only=True)["DataBase_Detail"]
    out, desc = defaultdict(lambda: [0.0] * 12), {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[I_CM] is None:
            continue
        chave = (str(r[I_CM]).strip(), cod(r[I_GEST]))
        desc.setdefault(chave, str(r[I_DESC] or "").strip())
        for i, v in enumerate(r[CI(a) - 1:CI(b)]):
            if isinstance(v, (int, float)):
                out[chave][i] += v / 1000
    return out, desc


def ler_ksb1(arq, ultimo_mes):
    """(cm, gestorial) -> [12 meses em R$ mil] (só até ultimo_mes) + descrições."""
    ws = load_workbook(arq, data_only=True, read_only=True)["BASE_KSB1"]
    out, desc = defaultdict(lambda: [0.0] * 12), {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        v, m, cm = r[K_VALOR], r[K_MES], r[K_CM]
        if not isinstance(v, (int, float)) or not isinstance(m, (int, float)) or not cm:
            continue
        if not 1 <= m <= ultimo_mes:
            continue
        chave = (str(cm).strip(), cod(r[K_GEST]))
        desc.setdefault(chave, str(r[K_DESC] or "").strip())
        out[chave][int(m) - 1] += v / 1000
    return out, desc


def cab(ws, linha, titulos, fills=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(linha, i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = (fills or {}).get(i, AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    ksb = Path(sys.argv[1]) if len(sys.argv) > 1 else ARQ_KSB1
    r9a = Path(sys.argv[2]) if len(sys.argv) > 2 else ARQ_R9
    mp27a = Path(sys.argv[3]) if len(sys.argv) > 3 else ARQ_MP27
    saida = Path(sys.argv[4]) if len(sys.argv) > 4 else SAIDA
    ult = int(sys.argv[5]) if len(sys.argv) > 5 else 9

    r9, d9 = ler_detalhe(r9a, "R9")
    m27, d27 = ler_detalhe(mp27a, "MP27")
    efe, de = ler_ksb1(ksb, ult)

    gest_detalhe = {g for (_, g) in r9} | {g for (_, g) in m27}
    desc = {**de, **d27, **d9}

    # --- R9 revisado: efetivo nos meses 1..ult, premissa R9 nos demais
    rev, fora = {}, defaultdict(lambda: [0.0] * 12)
    for chave in set(r9) | set(m27) | {k for k in efe if k[1] in gest_detalhe}:
        base = r9.get(chave, [0.0] * 12)
        e = efe.get(chave, [0.0] * 12)
        rev[chave] = [e[i] if i < ult else base[i] for i in range(12)]
    for chave, v in efe.items():
        if chave[1] not in gest_detalhe:
            for i in range(12):
                fora[chave][i] += v[i]

    # --- validação do método: Jan-Jun (R9 já tem efetivo) e Jul-Ult (R9 sem efetivo)
    print(f"Validação KSB1 x R9 original (R$ mil) - só gestoriais do Detalhe, ult. mês efetivo = {ult}")
    print(f"{'Un':4s} {'R9 Jan-Jun':>11s} {'KSB1 Jan-Jun':>13s} {'dif':>8s} | {'R9 Jul-Set':>11s} {'KSB1 Jul-Set':>13s} {'dif':>8s}")
    for u in UNIDADES:
        ks = [k for k in set(r9) | set(efe) if k[0] == u and k[1] in gest_detalhe]
        a1 = sum(sum(r9.get(k, [0] * 12)[:6]) for k in ks)
        b1 = sum(sum(efe.get(k, [0] * 12)[:6]) for k in ks)
        a2 = sum(sum(r9.get(k, [0] * 12)[6:ult]) for k in ks)
        b2 = sum(sum(efe.get(k, [0] * 12)[6:ult]) for k in ks)
        print(f"{u:4s} {a1:11.1f} {b1:13.1f} {b1 - a1:8.1f} | {a2:11.1f} {b2:13.1f} {b2 - a2:8.1f}")

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Resumo"
    abas = {}
    for u in UNIDADES:
        ws = wb.create_sheet(u)
        abas[u] = ws
        tit = ["Gestorial", "Descrição", "R9 original (ano)", f"R9 revisado: efetivo Jan-{MESES[ult-1].title()}",
               f"R9 revisado: premissa {MESES[ult].title()}-Dez", "R9 REVISADO (ano)", "MP27 (V3)",
               "Δ MP27 vs R9 revisado", "Δ MP27 vs R9 original", "Efeito da revisão (R9 rev. − R9 orig.)",
               *MESES]
        fills = {i: (EFET if 13 <= i < 13 + ult else PREM) for i in range(13, 25)}
        fills = {k: PatternFill("solid", fgColor="548235" if k < 13 + ult else "BF8F00") for k in fills}
        cab(ws, 3, tit, fills)
        ws.cell(1, 1, f"{u} - R9 revisado com efetivo (R$ mil, custo positivo). Verde = efetivo KSB1; amarelo = premissa R9.").font = Font(bold=True)
        chaves = sorted([k for k in rev if k[0] == u],
                        key=lambda k: -abs(sum(m27.get(k, [0] * 12)) - sum(rev[k])))
        for n, k in enumerate(chaves):
            r = 4 + n
            ws.cell(r, 1, k[1]); ws.cell(r, 2, desc.get(k, ""))
            ws.cell(r, 3, round(sum(r9.get(k, [0] * 12)), 3))
            for i in range(12):
                ws.cell(r, 13 + i, round(rev[k][i], 3))
            ef, pr = CL(13), CL(12 + ult)
            ws.cell(r, 4, f"=SUM({ef}{r}:{pr}{r})")
            ws.cell(r, 5, f"=SUM({CL(13 + ult)}{r}:{CL(24)}{r})")
            ws.cell(r, 6, f"=D{r}+E{r}")
            ws.cell(r, 7, round(sum(m27.get(k, [0] * 12)), 3))
            ws.cell(r, 8, f"=G{r}-F{r}")
            ws.cell(r, 9, f"=G{r}-C{r}")
            ws.cell(r, 10, f"=F{r}-C{r}")
        fim = 3 + len(chaves)
        t = fim + 1
        ws.cell(t, 2, "TOTAL").font = Font(bold=True)
        for c in list(range(3, 25)):
            ws.cell(t, c, f"=SUM({CL(c)}4:{CL(c)}{fim})").font = Font(bold=True)
            ws.cell(t, c).fill = CINZA
        for row in ws.iter_rows(min_row=4, max_row=t, min_col=3, max_col=24):
            for c in row:
                c.number_format = FMT
        ws.freeze_panes = "C4"
        ws.column_dimensions["A"].width = 11; ws.column_dimensions["B"].width = 36
        for c in range(3, 25):
            ws.column_dimensions[CL(c)].width = 14 if c < 11 else 9
        ws.row_dimensions[3].height = 48
        abas[u].tot = t

    cab(ws0, 3, ["Unidade", "R9 original", "R9 REVISADO", "MP27 (V3)", "Δ MP27 vs R9 revisado",
                 "Δ MP27 vs R9 original", "Efeito da revisão do R9"])
    ws0.cell(1, 1, "R9 revisado = efetivo KSB1 Jan-Set + premissa R9 Out-Dez. R$ mil, custo positivo, só vozes do Detalhe.").font = Font(bold=True)
    for n, u in enumerate(UNIDADES):
        r = 4 + n
        t = abas[u].tot
        ws0.cell(r, 1, u)
        ws0.cell(r, 2, f"='{u}'!C{t}"); ws0.cell(r, 3, f"='{u}'!F{t}"); ws0.cell(r, 4, f"='{u}'!G{t}")
        ws0.cell(r, 5, f"=D{r}-C{r}"); ws0.cell(r, 6, f"=D{r}-B{r}"); ws0.cell(r, 7, f"=C{r}-B{r}")
    r = 4 + len(UNIDADES)
    ws0.cell(r, 1, "TOTAL").font = Font(bold=True)
    for c in range(2, 8):
        ws0.cell(r, c, f"=SUM({CL(c)}4:{CL(c)}{r - 1})").font = Font(bold=True)
        ws0.cell(r, c).fill = CINZA
    for row in ws0.iter_rows(min_row=4, max_row=r, min_col=2, max_col=7):
        for c in row:
            c.number_format = FMT
    ws0.column_dimensions["A"].width = 12
    for c in range(2, 8):
        ws0.column_dimensions[CL(c)].width = 20
    ws0.row_dimensions[3].height = 36

    wf = wb.create_sheet("Fora do Detalhe")
    cab(wf, 1, ["Unidade", "Gestorial", "Descrição", f"Efetivo Jan-{MESES[ult-1].title()} (R$ mil)", "Obs."])
    for n, (k, v) in enumerate(sorted(fora.items(), key=lambda x: -abs(sum(x[1]))), 2):
        wf.cell(n, 1, k[0]); wf.cell(n, 2, k[1]); wf.cell(n, 3, desc.get(k, ""))
        wf.cell(n, 4, round(sum(v), 3)).number_format = FMT
        wf.cell(n, 5, "Gestorial que não existe no Detalhe (ex.: mão de obra) - fora do R9 revisado")
    for c, w in zip("ABCDE", (10, 12, 40, 22, 60)):
        wf.column_dimensions[c].width = w

    wn = wb.create_sheet("Notas")
    for n, txt in enumerate([
        "Fontes: KSB1 September Flash 2026 (BASE_KSB1, Jan-Set, Flash = ainda não é o Actual fechado);",
        "R9 = Detalhe_Despesas Forecast Setembro_final ('Total C/ Curva'); MP27 = Detalhe Budget'27 V3.",
        "Chave de casamento: Unidade (CM) + Gestorial (97% de aderência, ver ontology/classificacao_gestorial_mp2027.json).",
        "Out-Dez do R9 revisado = valores do R9 sem alteração (inclui linhas manuais, ex.: 'Delta MP'26').",
        "Nos meses de efetivo, o valor do R9 original é SUBSTITUÍDO pelo KSB1 (ajustes manuais do R9 nesses meses saem).",
        "Frete e Materiais Diretos do Detalhe vêm do MENS no MP27; no efetivo vêm do KSB1 (4211000 / 4236300).",
    ], 1):
        wn.cell(n, 1, txt)
    wn.column_dimensions["A"].width = 120

    saida.mkdir(parents=True, exist_ok=True)
    caminho = saida / nome_com_versao(saida, "MP27_vs_R9_Revisado_Efetivo_Jan_Set.xlsx")
    wb.save(caminho)
    print("Arquivo:", caminho)


if __name__ == "__main__":
    main()
