#!/usr/bin/env python3
"""
Explica o ganho (ou perda) de Custos do Actual vs o Flash de um mes: compara a
coluna do mes da aba 'Intermediaria' das duas Bases Intermediarias, linha a
linha (Conta Gestorial + Conta Fiscal + Centro de Custo), e gera um Excel com
a ponte por tema, por gestorial, por unidade e o detalhe de cada linha.

Uso: python comparar_flash_actual.py <flash.xlsx> <actual.xlsx> <mes 1-12> <pasta_saida>

So le os arquivos de entrada; a saida nunca sobrescreve (nome_com_versao).
"""
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "_shared"))
from ksb1_core import nome_com_versao  # noqa: E402

MESES_INGLES = {
    1: "January", 2: "February", 3: "March", 4: "April", 5: "May", 6: "June",
    7: "July", 8: "August", 9: "September", 10: "October", 11: "November", 12: "December",
}

COL_MES_JANEIRO = 8  # indice 0-based de 'January' (coluna I)

TEMAS = [
    ("pc", "Créditos PIS/COFINS (contas _PC) — entraram só no Actual"),
    ("manut_mg", "Manutenção MG: provisão do Flash trocada pelo realizado (sem efeito líquido)"),
    ("alug_arm", "Aluguel armazém MG: provisão × lançamento 'MAPPING WRONG' (reclassificação)"),
    ("deprec_pred", "Depreciação prédios alugados (IFRS16): provisão × realizado"),
    ("conserv", "Materiais para conservação/reparo"),
    ("veiculos", "Veículos: aluguel × depreciação (reclassificação entre contas)"),
    ("demais", "Demais variações (realizado × provisão)"),
]
NOME_TEMA = dict(TEMAS)


def tema(fiscal) -> str:
    f = str(fiscal)
    if f.endswith("_PC"):
        return "pc"
    if f == "N14705S001":
        return "manut_mg"
    if f in ("N14301S001", "N1430100RT"):
        return "alug_arm"
    if f in ("N4200500RT", "N42005S000"):
        return "deprec_pred"
    if f in ("N13020S001", "N13020S000"):
        return "conserv"
    if f in ("N14306S000", "N4221100RT", "N14303S000", "N4221000RT", "N42210S000",
             "N1430300RT", "N1430600RT", "N42211S000"):
        return "veiculos"
    return "demais"


def ler(caminho: Path, mes: int):
    ws = openpyxl.load_workbook(caminho, data_only=True, read_only=True)["Intermediária"]
    idx = COL_MES_JANEIRO + mes - 1
    linhas = []
    for r in ws.iter_rows(min_row=2, max_col=36, values_only=True):
        v = r[idx]
        v = v if isinstance(v, (int, float)) else 0.0
        linhas.append({
            "gest": r[0], "gest_desc": r[1], "fiscal": r[2], "fiscal_desc": r[3],
            "cc": r[4], "unidade": r[6], "valor": v, "modg": r[25], "var": r[26],
        })
    return linhas


def somar(linhas, chave):
    g = defaultdict(float)
    info = {}
    for x in linhas:
        k = chave(x)
        g[k] += x["valor"]
        info.setdefault(k, x)
    return g, info


def main(flash, actual, mes, pasta_saida):
    lf, la = ler(flash, mes), ler(actual, mes)
    tot_f, tot_a = sum(x["valor"] for x in lf), sum(x["valor"] for x in la)
    delta = tot_a - tot_f  # negativo = Actual mais barato que o Flash = ganho

    chave = lambda x: (x["gest"], x["gest_desc"], x["fiscal"], x["fiscal_desc"], x["cc"], x["unidade"])
    gf, inf_f = somar(lf, chave)
    ga, inf_a = somar(la, chave)
    detalhe = []
    for k in set(gf) | set(ga):
        d = ga.get(k, 0.0) - gf.get(k, 0.0)
        if abs(d) < 0.005:
            continue
        info = inf_a.get(k) or inf_f.get(k)
        detalhe.append({"k": k, "flash": gf.get(k, 0.0), "actual": ga.get(k, 0.0), "delta": d,
                        "tema": tema(k[2]), "var": info["var"], "modg": info["modg"]})
    detalhe.sort(key=lambda r: (list(NOME_TEMA).index(r["tema"]), r["delta"]))

    def agrupa(campo):
        out = defaultdict(lambda: [0.0, 0.0])
        for x in lf:
            out[campo(x)][0] += x["valor"]
        for x in la:
            out[campo(x)][1] += x["valor"]
        return sorted(((k, v[0], v[1], v[1] - v[0]) for k, v in out.items() if abs(v[1] - v[0]) > 0.005),
                      key=lambda r: r[3])

    por_gest = agrupa(lambda x: f"{x['gest']} - {x['gest_desc']}")
    por_unid = agrupa(lambda x: str(x["unidade"]).title())
    por_tipo = agrupa(lambda x: f"{x['modg']} / {x['var']}")

    wb = openpyxl.Workbook()
    neg = Font(bold=True)
    cab_fill = PatternFill("solid", fgColor="FFE9A8")
    fmt = '#,##0;[Red]-#,##0'

    def cabecalho(ws, linha, titulos):
        for i, t in enumerate(titulos, 1):
            c = ws.cell(linha, i, t)
            c.font, c.fill = neg, cab_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")

    def larguras(ws, ls):
        for i, w in enumerate(ls, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    mes_nome = MESES_INGLES[mes]

    # --- Resumo (ponte por tema)
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = f"Custos {mes_nome}: Actual vs Flash (R$)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Delta negativo = Actual mais barato que o Flash (ganho). Só a coluna do mês da aba Intermediária."
    cabecalho(ws, 4, ["", "R$"])
    ws.append(["Custos Flash", tot_f])
    ws.append(["Custos Actual", tot_a])
    ws.append(["Ganho do Actual vs Flash (Flash − Actual)", tot_f - tot_a])
    ws.cell(ws.max_row, 1).font = neg
    ws.append([])
    ws.append(["Ponte por tema (delta = Actual − Flash)", "Delta R$", "Flash", "Actual"])
    cabecalho(ws, ws.max_row, ["Ponte por tema (delta = Actual − Flash)", "Delta R$", "Flash", "Actual"])
    soma_t = 0.0
    for cod, nome in TEMAS:
        rs = [r for r in detalhe if r["tema"] == cod]
        if not rs:
            continue
        d = sum(r["delta"] for r in rs)
        soma_t += d
        ws.append([nome, d, sum(r["flash"] for r in rs), sum(r["actual"] for r in rs)])
    ws.append(["Total (confere com Actual − Flash)", soma_t])
    ws.cell(ws.max_row, 1).font = neg
    ws.append(["Diferença para o total (deve ser 0)", round(soma_t - delta, 2)])
    ws.append([])
    ws.append(["Mão de obra x despesas (delta)"])
    ws.cell(ws.max_row, 1).font = neg
    for k, f_, a_, d_ in por_tipo:
        ws.append([k, d_, f_, a_])
    for row in ws.iter_rows(min_row=5):
        for c in row[1:]:
            c.number_format = fmt
    larguras(ws, [78, 16, 16, 16])

    # --- Gestorial / Unidade
    for nome_aba, dados, rot in (("Por Gestorial", por_gest, "Conta gestorial"), ("Por Unidade", por_unid, "Unidade")):
        w = wb.create_sheet(nome_aba)
        cabecalho(w, 1, [rot, "Flash", "Actual", "Delta (Actual − Flash)"])
        for r in dados:
            w.append(list(r))
        w.append(["Total", tot_f, tot_a, delta])
        w.cell(w.max_row, 1).font = neg
        for row in w.iter_rows(min_row=2):
            for c in row[1:]:
                c.number_format = fmt
        larguras(w, [44, 16, 16, 22])
        w.freeze_panes = "A2"

    # --- Detalhe
    w = wb.create_sheet("Detalhe por linha")
    cabecalho(w, 1, ["Tema", "Conta gestorial", "Descrição gestorial", "Conta fiscal", "Descrição conta fiscal",
                     "Centro de custo", "Unidade", "MO/DG", "Var/Fixo", "Flash", "Actual", "Delta (Actual − Flash)"])
    for r in detalhe:
        g, gd, fi, fd, cc, un = r["k"]
        w.append([NOME_TEMA[r["tema"]], g, gd, fi, fd, cc, un, r["modg"], r["var"], r["flash"], r["actual"], r["delta"]])
    for row in w.iter_rows(min_row=2):
        for c in row[9:]:
            c.number_format = fmt
    larguras(w, [60, 14, 28, 14, 44, 12, 28, 8, 9, 14, 14, 18])
    w.freeze_panes = "A2"
    w.auto_filter.ref = w.dimensions

    # --- Notas
    w = wb.create_sheet("Notas")
    for t in [
        "Fonte: coluna do mês da aba 'Intermediária' das Bases Intermediárias Flash e Actual (soma de todas as linhas).",
        f"Conferência: soma Flash = {tot_f:,.2f}; soma Actual = {tot_a:,.2f}.",
        "Chave de comparação: Conta Gestorial + Conta Fiscal + Centro de Custo (valores somados quando a chave repete).",
        "O Flash tem linhas de provisão/reclassificação lançadas à mão (linhas coloridas); no Actual essas linhas ficam em branco e o realizado vem do KSB1.",
        "Os temas são agrupados por regra de conta fiscal (função tema() do script) — são hipóteses de leitura, não confirmação contábil.",
        "Contas '_PC' são créditos de PIS/COFINS, só integrados na KSB1 do Actual (ver DECISOES 2026-08-19).",
        "Temas 'sem efeito líquido' mostram a troca de conta entre Flash e Actual; o total do tema é próximo de zero.",
    ]:
        w.append([t])
    w.column_dimensions["A"].width = 130

    pasta_saida.mkdir(parents=True, exist_ok=True)
    nome = nome_com_versao(pasta_saida, f"Ganho Actual vs Flash - {mes_nome} - Custos.xlsx")
    wb.save(pasta_saida / nome)
    print(f"Salvo: {pasta_saida / nome}")
    print(f"Flash={tot_f:,.2f} Actual={tot_a:,.2f} ganho={tot_f - tot_a:,.2f}")
    for cod, nome_t in TEMAS:
        rs = [r for r in detalhe if r["tema"] == cod]
        if rs:
            print(f"  {sum(r['delta'] for r in rs):>14,.2f}  {nome_t}")


if __name__ == "__main__":
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]), int(sys.argv[3]), Path(sys.argv[4]))
