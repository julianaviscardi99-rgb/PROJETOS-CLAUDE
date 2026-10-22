"""Deck MP'27 da TLM no formato Fitted/Circuito: 1 slide do total, 1 consolidado por unidade e 1 por unidade.

Uso:
    python gerar_ppt_tlm_mp27.py <template.pptx> <pl_unidades.json> <ref_r09_mp26.json> <saida.pptx> [fonte_label]

- template.pptx : deck do Circuito (usa a capa e o layout 'D. Title Only', logo e numeração); slides dele são descartados.
- pl_unidades.json : P&L por aba do Resultado TLM MP27 v2 recalculado (colunas A, N=MP27, AB=R09 em cache, AP) - ver
  mp27_extrair_resumo_slide.py.
- ref_r09_mp26.json : linhas-chave de R09 e MP26 lidas do cache dos próprios arquivos de resultado.
Convenção: valores em BRL mil, (-) = custo/perda; Δ positivo = ganho. PBIT = Net Sales + custo variável + custo fixo
(antes de itens não recorrentes - lay-off).
"""
import sys
import json
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE

TEMPLATE, PL_JSON, REF_JSON, OUT = sys.argv[1:5]
FONTE = sys.argv[5] if len(sys.argv) > 5 else "Resultado TLM - MP27 vs R09_v2"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MIDBLUE = RGBColor(0x2E, 0x5A, 0xA8)
LIGHT = RGBColor(0xDC, 0xE3, 0xF0)
PANEL = RGBColor(0xF3, 0xF5, 0xFA)
EDGE = RGBColor(0x3C, 0x4B, 0x66)
SKY = RGBColor(0x9D, 0xC3, 0xE6)
GREY = RGBColor(0x80, 0x80, 0x80)
DGREY = RGBColor(0x40, 0x40, 0x40)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)
MPFILL = RGBColor(0xEE, 0xF2, 0xF9)
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# R09: net sales do Resultado R09 chega como #VALUE! no cache (HC fechado); total vem do Resumo FY do R09
# (Resumo Resultado TLM - MP27 vs R09, aba Resumo FY, G6) e as fábricas são o resíduo.
R09_NET_TOTAL = 93023.973

# ------------------------------------------------------------------ dados
pl = json.load(open(PL_JSON, encoding="utf-8"))
ref = json.load(open(REF_JSON, encoding="utf-8"))


def num(v):
    return float(v) if isinstance(v, (int, float)) and abs(v) < 1e9 else None


def grab(src, sheet, r, col):
    """src: 'MP27' (pl, col N=1/AB=2) ou 'R09'/'MP26' (ref)."""
    if src == "MP27":
        return num(pl[sheet][r - 1][col])
    d = ref[src][sheet].get(str(r))
    return num(d[1]) if d else None


def build(sheet):
    out = {}
    for sc in ("MP26", "R09", "MP27"):
        g = lambda r: grab(sc, sheet, r, 1) or 0.0
        e = dict(net=grab(sc, sheet, 13, 1), var=g(21), fix=g(50), lab=g(23) + g(51), da=g(42) + g(106),
                 tons=g(4), hc=grab(sc, sheet, 19, 1))
        if sc == "R09":
            e["hc"] = num(pl[sheet][18][2])
        out[sc] = e
    return out


SHEETS = ["P&L - TLM", "P&L Log", "BATotal", "SBLog", "CPLog", "CBTotal", "GVLog", "FSLog",
          "P&L Outras Atividades", "SBTransp", "CPEng", "BAMP", "Hevea", "BAQA", "Campneus", "CPCheck",
          "P&L Fab", "CPFab", "FSFab"]
D = {s: build(s) for s in SHEETS}
# net sales R09 em falta (fábricas) = resíduo do total R09
fab_net = R09_NET_TOTAL - D["P&L Log"]["R09"]["net"] - D["P&L Outras Atividades"]["R09"]["net"]
D["P&L Fab"]["R09"]["net"] = fab_net
D["CPFab"]["R09"]["net"] = fab_net - D["FSFab"]["R09"]["net"]
D["P&L - TLM"]["R09"]["net"] = R09_NET_TOTAL


def pbit(e):
    return e["net"] + e["var"] + e["fix"]


def cost(e):
    return e["var"] + e["fix"]


for s in SHEETS:
    for sc in ("MP26", "R09", "MP27"):
        assert D[s][sc]["net"] is not None, (s, sc)
# checagens de consistência (BRL mil)
t = D["P&L - TLM"]
assert abs(pbit(t["R09"]) - 3437.1) < 0.5, pbit(t["R09"])
assert abs(sum(pbit(D[k]["MP27"]) for k in ("P&L Log", "P&L Outras Atividades", "P&L Fab")) - pbit(t["MP27"])) < 0.5
assert abs(sum(pbit(D[k]["R09"]) for k in ("P&L Log", "P&L Outras Atividades", "P&L Fab")) - pbit(t["R09"])) < 0.5
LAYOFF = {"MP27": num(pl["P&L - TLM"][125][1]), "R09": num(ref["R09"]["P&L - TLM"]["126"][1]),
          "MP26": num(ref["MP26"]["P&L - TLM"]["126"][1])}


# ------------------------------------------------------------------ formatação
def r0(v):
    return int(round(v))


def fmt(v, sign=False, paren=False, dec=0):
    a = abs(round(v, dec))
    s = f"{a:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if round(v, dec) < 0:
        return f"({s})" if paren else "-" + s
    return ("+" if sign and round(v, dec) > 0 else "") + s


def p1(x, sign=True):
    x = 0.0 if abs(x) < 0.05 else x
    return (f"{x:+.1f}" if sign and x else f"{x:.1f}").replace(".", ",")


def pct(a, b):
    return (a / b - 1) * 100 if b else 0.0


def mi(v, sign=False):
    return fmt(v / 1000, sign=sign, dec=1)


# ------------------------------------------------------------------ helpers de desenho (padrão do deck do CP)
def line(sl, x1, y1, x2, y2, color=GREY, w=0.75, dash=None, arrow=False):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    if arrow:
        etree.SubElement(c.line._get_or_add_ln(), A + "tailEnd").set("type", "triangle")
    return c


def text(sl, x, y, w, h, runs, size=9, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bold=False, italic=False, spacing=1.0, after=0):
    tb = sl.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    paras = [[(ln, {})] for ln in runs.split("\n")] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        for t_, f in para:
            r = p.add_run()
            r.text = t_
            r.font.size = Pt(f.get("size", size))
            r.font.bold = f.get("bold", bold)
            r.font.italic = f.get("italic", italic)
            r.font.color.rgb = f.get("color", color)
    return tb


def rect(sl, x, y, w, h, fill, edge=None, ew=0.75, shape=MSO_SHAPE.RECTANGLE):
    r = sl.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = fill
    if edge is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = edge
        r.line.width = Pt(ew)
    r.shadow.inherit = False
    return r


def cell_border(ce, side, color="FFFFFF", w=12700):
    tcPr = ce._tc.get_or_add_tcPr()
    for old in tcPr.findall(A + side):
        tcPr.remove(old)
    ln = etree.SubElement(tcPr, A + side, w=str(w))
    sf = etree.SubElement(ln, A + "solidFill")
    etree.SubElement(sf, A + "srgbClr", val=color)


def new_slide(prs, layout, title):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if "TITLE" in str(ph.placeholder_format.type):
            ph.text_frame.text = title
        else:
            ph._element.getparent().remove(ph._element)
    return s


def gain_label(sl):
    text(sl, 10, 52, 200, 14, "(+) gain  |  BRL ’000", size=9, italic=True)


def panel_blocks(sl, x, y, w, h, title, items, size=10):
    rect(sl, x, y, w, h, PANEL, LIGHT)
    rect(sl, x, y, 4, h, MIDBLUE)
    text(sl, x + 16, y + 10, w - 28, 20, title, size=13, bold=True, color=NAVY)
    top = y + 38
    bh = (y + h - 8 - top) / len(items)
    for i, (head, body) in enumerate(items):
        by = top + i * bh
        if i:
            line(sl, x + 16, by, x + w - 14, by, color=RGBColor(0xC9, 0xD3, 0xE6), w=0.5)
        text(sl, x + 16, by + 4, w - 30, bh - 8,
             [[(head, {"bold": True, "color": NAVY, "size": size + 0.5})],
              [(body, {"color": DGREY, "size": size})]],
             anchor=MSO_ANCHOR.MIDDLE, spacing=1.05, after=3)


def bridge(sl, x0, x1, y_top, y_bot, start, steps, start_label, end_label, bracket=None):
    """Bridge em formas nativas, base zero (aceita PBIT negativo). steps = [(rótulo, delta)]."""
    n = len(steps) + 2
    dx = (x1 - x0) / n
    bw = min(46.0, dx * 0.66)
    cx = [x0 + dx * (i + 0.5) for i in range(n)]
    lv = [start]
    for _, v in steps:
        lv.append(lv[-1] + v)
    end = lv[-1]
    vmax, vmin = max(max(lv), 0), min(min(lv), 0)
    cat_h = 30
    top_pad = 46 if bracket else 24
    bot_pad = 16 if vmin < 0 else 0
    plot_top, plot_bot = y_top + top_pad, y_bot - cat_h - bot_pad
    sc = (plot_bot - plot_top) / max(vmax - vmin, 1)
    yz = plot_top + vmax * sc
    Y = lambda v: yz - v * sc
    line(sl, x0, yz, x1, yz, color=BLACK, w=0.75)

    def total(i, v):
        top, bot = min(Y(v), yz), max(Y(v), yz)
        rect(sl, cx[i] - bw / 2, top, bw, max(bot - top, 1.2), NAVY)
        ly = top - 19 if v >= 0 else bot + 3
        text(sl, cx[i] - 32, ly, 64, 16, fmt(v), size=13, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.BOTTOM if v >= 0 else MSO_ANCHOR.TOP)

    total(0, start)
    for i, (lab, v) in enumerate(steps, start=1):
        a, b = Y(lv[i - 1]), Y(lv[i])
        top, bot = min(a, b), max(a, b)
        rect(sl, cx[i] - bw / 2, top, bw, max(bot - top, 1.2), LIGHT, EDGE, 0.75)
        line(sl, cx[i - 1] + bw / 2, a, cx[i] - bw / 2, a, color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
        text(sl, cx[i] - 32, top - 14, 64, 12, fmt(v, sign=True), size=9, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.BOTTOM)
    total(n - 1, end)
    line(sl, cx[-2] + bw / 2, Y(lv[-2]), cx[-1] - bw / 2, Y(lv[-2]), color=BLACK, w=0.25, dash=MSO_LINE.LONG_DASH)
    labels = [start_label] + [l for l, _ in steps] + [end_label]
    for c, lab in zip(cx, labels):
        text(sl, c - dx / 2, y_bot - cat_h + 6, dx, cat_h, lab, size=8.5,
             bold=lab in (start_label, end_label), align=PP_ALIGN.CENTER)
    if bracket:
        ytop = y_top + 6
        line(sl, cx[0], ytop, cx[0], min(Y(start), yz) - 22 if start >= 0 else yz - 6)
        line(sl, cx[0], ytop, cx[-1], ytop)
        line(sl, cx[-1], ytop, cx[-1], (min(Y(end), yz) - 22) if end >= 0 else yz - 6, arrow=True)
        ov = rect(sl, (cx[0] + cx[-1]) / 2 - 36, ytop - 10, 72, 20, WHITE, GREY, shape=MSO_SHAPE.OVAL)
        tf = ov.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = bracket
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = BLACK


# ------------------------------------------------------------------ tabela P&L (padrão Fitted/CP)
CW = [140, 86, 86, 86, 10, 80, 80]
HDR = ["", "MP’26", "Fcst R9", "MP’27", "", "vs. R9", "vs. MP’26"]
SCN = ("MP26", "R09", "MP27")


def pl_table(sl, e, x=14, y=40):
    """e = D[sheet]; colunas MP26 / R9 / MP27 e deltas (MP27 - ref)."""
    def trio(fn):
        return [fn(e[s]) for s in SCN]

    rows = [
        ("Net Sales", trio(lambda d: d["net"]), "main", "n"),
        ("Tons (’000)", trio(lambda d: d["tons"] / 1000), "sub", "t"),
        ("Variable cost", trio(lambda d: d["var"]), "main", "n"),
        ("CM", trio(lambda d: d["net"] + d["var"]), "cm", "n"),
        ("Fixed cost", trio(lambda d: d["fix"]), "main", "n"),
        ("PBIT", trio(pbit), "main", "n"),
        ("ROS", trio(lambda d: pbit(d) / d["net"] * 100 if d["net"] else None), "ros", "p"),
        ("o/w Labour", trio(lambda d: d["lab"]), "sub", "n"),
        ("o/w D&A", trio(lambda d: d["da"]), "sub", "n"),
        ("o/w Other costs", trio(lambda d: cost(d) - d["lab"] - d["da"]), "sub", "n"),
        ("HC (avg.)", trio(lambda d: d["hc"]), "sub", "h"),
    ]
    heights = [18, 15] + [17 if k in ("main", "cm") else (14 if k == "ros" else 12.5) for _, _, k, _ in rows]
    nrows = len(heights)
    gt = sl.shapes.add_table(nrows, len(CW), Pt(x), Pt(y), Pt(sum(CW)), Pt(sum(heights)))
    tbl = gt.table
    pr = gt._element.graphic.graphicData.tbl.tblPr
    pr.set("firstRow", "0")
    pr.set("bandRow", "0")
    sid = pr.find(A + "tableStyleId")
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(CW):
        tbl.columns[i].width = Pt(w)
    for i, h in enumerate(heights):
        tbl.rows[i].height = Pt(h)

    def put(r_, c_, t_, size=8, bold=False, color=BLACK, align=PP_ALIGN.RIGHT, fill=None, italic=False):
        ce = tbl.cell(r_, c_)
        ce.margin_top = ce.margin_bottom = Pt(0)
        ce.margin_left = Pt(3)
        ce.margin_right = Pt(6)
        ce.vertical_anchor = MSO_ANCHOR.MIDDLE
        ce.text = ""
        p = ce.text_frame.paragraphs[0]
        p.alignment = align
        rr = p.add_run()
        rr.text = t_
        rr.font.size = Pt(size)
        rr.font.bold = bold
        rr.font.italic = italic
        rr.font.color.rgb = color
        etree.SubElement(p._p, A + "endParaRPr").set("sz", str(int(size * 100)))
        if fill is None:
            ce.fill.background()
        else:
            ce.fill.solid()
            ce.fill.fore_color.rgb = fill

    for c_ in range(len(CW)):
        put(0, c_, "")
        put(1, c_, "")
    tbl.cell(0, 1).merge(tbl.cell(0, 3))
    put(0, 1, "Full Year", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, fill=NAVY)
    tbl.cell(0, 5).merge(tbl.cell(0, 6))
    put(0, 5, "Δ", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER, fill=SKY)
    for c_, h in enumerate(HDR):
        if c_ in (0, 4):
            continue
        put(1, c_, h, size=8.5, color=WHITE if c_ <= 3 else NAVY, align=PP_ALIGN.CENTER,
            fill=NAVY if c_ <= 3 else SKY, bold=(c_ == 3))
    for i, (lab, vals, kind, f) in enumerate(rows):
        r_ = 2 + i
        sz = 10 if kind == "main" else (9.5 if kind == "cm" else (8.5 if kind == "ros" else 7.5))
        put(r_, 0, lab, size=sz, bold=kind == "main", align=PP_ALIGN.RIGHT, italic=kind == "sub")
        put(r_, 4, "")
        for k in range(3):
            v = vals[k]
            if v is None or (f == "t" and not any(vals)):
                t_ = "–"
            elif f == "p":
                t_ = p1(v, False) + "%"
            elif f == "t":
                t_ = fmt(v, dec=1)
            else:
                t_ = fmt(v, paren=True)
            put(r_, 1 + k, t_, size=sz, bold=kind == "main" or (k == 2 and kind == "cm"), align=PP_ALIGN.CENTER,
                fill=MPFILL if k == 2 else None)
        for j, k in enumerate((1, 0)):           # vs R9, vs MP26
            if vals[2] is None or vals[k] is None or (f == "t" and not any(vals)):
                t_ = "–"
            elif f == "p":
                t_ = p1(vals[2] - vals[k]) + " p.p."
            elif f == "t":
                t_ = fmt(vals[2] - vals[k], dec=1)
            else:
                t_ = fmt(vals[2] - vals[k], paren=True)
            put(r_, 5 + j, t_, size=sz if f != "p" else 7.5, bold=kind == "main",
                color=NAVY if kind == "main" else GREY, align=PP_ALIGN.CENTER)
    for r_ in range(nrows):
        for c_ in range(len(CW)):
            for side in ("lnL", "lnR", "lnT", "lnB"):
                cell_border(tbl.cell(r_, c_), side, "FFFFFF", 0)
    sep = {2 + 3, 2 + 4, 2 + 5, 2 + 7}           # antes de CM, Fixed cost, PBIT, memo
    for r_ in sep:
        for c_ in (1, 2, 3, 5, 6):
            cell_border(tbl.cell(r_, c_), "lnT", "A6A6A6", 6350)
    return y + sum(heights)


def steps_for(e):
    """Bridge R9 -> MP27 do PBIT por driver (Δ positivo = ganho)."""
    a, b = e["R09"], e["MP27"]
    d_net = b["net"] - a["net"]
    d_lab = b["lab"] - a["lab"]
    d_da = b["da"] - a["da"]
    d_oth = (cost(b) - b["lab"] - b["da"]) - (cost(a) - a["lab"] - a["da"])
    st = [("Net Sales", d_net), ("Labour", d_lab), ("D&A", d_da), ("Other costs", d_oth)]
    assert abs(pbit(a) + sum(v for _, v in st) - pbit(b)) < 0.5
    return st


def footer(sl, extra=""):
    text(sl, 36, 516, 760, 14,
         f"Fonte: {FONTE}. PBIT antes de itens não recorrentes (lay-off). Fcst R9 e MP’26: arquivos de resultado de origem. {extra}",
         size=7, italic=True, color=GREY)


def entity_slide(prs, layout, title, key, items, note="", extra_notes=""):
    e = D[key]
    sl = new_slide(prs, layout, title)
    gain_label(sl)
    yb = pl_table(sl, e)
    st = steps_for(e)
    d = pbit(e["MP27"]) - pbit(e["R09"])
    bridge(sl, 40, 640, yb + 14, 512, pbit(e["R09"]), st, "R9", "MP’27", bracket=fmt(d, sign=True))
    panel_blocks(sl, 652, 44, 296, 468, "Mensagens-chave", items, size=9.5)
    footer(sl, note)
    rank = sorted(st, key=lambda x: -abs(x[1]))
    efeitos = "; ".join(f"{l} {fmt(v, sign=True)} ({'ganho' if v > 0 else 'perda'})" for l, v in rank if abs(v) >= 0.5)
    sl.notes_slide.notes_text_frame.text = (
        f"PRINCIPAIS EFEITOS MP'27 vs R9 (PBIT {fmt(pbit(e['R09']))} -> {fmt(pbit(e['MP27']))}, "
        f"{fmt(d, sign=True)}), do maior para o menor: {efeitos or 'sem variação'}.\n" + extra_notes +
        f"\nBridge PBIT (BRL mil): R9 {pbit(e['R09']):,.1f} | " + " | ".join(f"{l} {v:+,.1f}" for l, v in st) +
        f" | MP'27 {pbit(e['MP27']):,.1f}")
    return sl


# ------------------------------------------------------------------ textos por slide
def ros(e, sc):
    d = e[sc]
    return pbit(d) / d["net"] * 100 if d["net"] else 0.0


def tarifa(e, sc):
    d = e[sc]
    return d["net"] / d["tons"] * 1000 if d["tons"] else 0.0     # R$/t


def d_(e, k, a="MP27", b="R09"):
    return e[a][k] - e[b][k]


def items_gen(e, nome, extra):
    d_pbit = pbit(e["MP27"]) - pbit(e["R09"])
    st = dict(steps_for(e))
    net_pct = pct(e["MP27"]["net"], e["R09"]["net"]) if e["R09"]["net"] else 0
    cost_d = cost(e["MP27"]) - cost(e["R09"])
    big = max((k for k in st if k != "Net Sales"), key=lambda k: abs(st[k]))
    return [
        (f"PBIT {fmt(pbit(e['MP27']))}  |  {fmt(d_pbit, sign=True)} vs R9",
         f"ROS de {p1(ros(e, 'MP27'), False)}% (R9: {p1(ros(e, 'R09'), False)}%; MP’26 PBIT {fmt(pbit(e['MP26']))})."),
        (f"Net Sales {fmt(e['MP27']['net'])}  |  {fmt(d_(e, 'net'), sign=True)} vs R9",
         f"{p1(net_pct)}% vs R9. " + extra["net"]),
        (f"Custo total {fmt(-cost(e['MP27']))}  |  {fmt(-cost_d, sign=True) if cost_d else '0'} vs R9",
         f"Trabalho {fmt(st['Labour'], sign=True)}, D&A {fmt(st['D&A'], sign=True)}, outros {fmt(st['Other costs'], sign=True)} "
         f"(Δ PBIT; maior efeito: {big}). " + extra["cost"]),
        (extra["h4"], extra["b4"]),
    ]


# ================================================================== deck
prs = Presentation(TEMPLATE)
layout = prs.slides[1].slide_layout
cover = prs.slides[0]
for ph in cover.placeholders:
    if "CIRCUITO" in ph.text_frame.text.upper():
        ph.text_frame.paragraphs[0].runs[0].text = "TLM – ECONOMICS BUDGET’27"
        for r in ph.text_frame.paragraphs[0].runs[1:]:
            r.text = ""
ids = prs.slides._sldIdLst
for sid in list(ids)[1:]:
    prs.part.drop_rel(sid.rId)
    ids.remove(sid)

T = D["P&L - TLM"]
tm, tr = T["MP27"], T["R09"]
unid = [("BA", "BATotal"), ("CP", "CPLog"), ("FS", "FSLog"), ("CB", "CBTotal"),
        ("OA", "P&L Outras Atividades"), ("FAB", "P&L Fab")]
sbgv_r = pbit(D["SBLog"]["R09"]) + pbit(D["GVLog"]["R09"])
sbgv_m = pbit(D["SBLog"]["MP27"]) + pbit(D["GVLog"]["MP27"])
net_ba = d_(D["BATotal"], "net")
net_cp = d_(D["CPLog"], "net")
net_fs = d_(D["FSLog"], "net")
net_cb = d_(D["CBTotal"], "net")
net_oa = d_(D["P&L Outras Atividades"], "net")
net_fab = d_(D["P&L Fab"], "net")
da_d = (tm["da"] - tr["da"])

# ---------------- slide 2: total
items = [
    (f"PBIT {fmt(pbit(tm))}  |  {fmt(pbit(tm) - pbit(tr), sign=True)} vs R9",
     f"ROS de {p1(ros(T, 'MP27'), False)}% (R9 {p1(ros(T, 'R09'), False)}%; MP’26 {p1(ros(T, 'MP26'), False)}%). "
     f"Lay-off {fmt(LAYOFF['MP27'])} segue abaixo do PBIT (R9: {fmt(LAYOFF['R09'])})."),
    (f"Net Sales {fmt(tm['net'])}  |  {p1(pct(tm['net'], tr['net']))}% vs R9",
     f"Barueri {fmt(net_ba, sign=True)} (tarifa média por tonelada R$ {fmt(tarifa(D['BATotal'], 'MP27'))} vs "
     f"{fmt(tarifa(D['BATotal'], 'R09'))}, volume estável) e Cabreúva {fmt(net_cb, sign=True)} (sem operação no MP’27) "
     f"(desconto de tarifa em Barueri é decisão de eficiência para a Pirelli) superam Campinas {fmt(net_cp, sign=True)} e Feira {fmt(net_fs, sign=True)}. Outras Atividades {fmt(net_oa, sign=True)}."),
    (f"Custo total {fmt(-cost(tm))}  |  {p1(pct(cost(tm), cost(tr)))}% vs R9",
     f"Trabalho {fmt(tm['lab'] - tr['lab'], sign=True)} (HC {fmt(tm['hc'])} vs {fmt(tr['hc'])}). "
     f"D&A {fmt(da_d, sign=True)} vs R9: fim do contrato IFRS16 do armazém de Barueri (CC 7287) em jul/26, "
     f"sem aluguel/amortização na base MP’27."),
    (f"Volume {fmt(tm['tons'] / 1000, dec=1)} mil t  |  {p1(pct(tm['tons'], tr['tons']))}% vs R9",
     f"Campinas {fmt(D['CPLog']['MP27']['tons'] / 1000 - D['CPLog']['R09']['tons'] / 1000, sign=True, dec=1)} mil t; "
     f"Cabreúva {fmt(-D['CBTotal']['R09']['tons'] / 1000, sign=True, dec=1)} mil t (sem operação)."),
]
s2 = entity_slide(prs, layout, "TLM Economics MP’27 – Total", "P&L - TLM", items,
                  note="Ecolabel/Outras Atividades com premissas de simulação (sem volume oficial).",
                  extra_notes="Roteiro: (1) a queda do PBIT vem do desconto de tarifa em Barueri, decisão de eficiência para a Pirelli; "
                            "(2) Trabalho piora pelo HC e reajuste; (3) D&A melhora pelo fim do IFRS16 do armazém de Barueri, sem custo novo na base; "
                            "(4) a saída de Cabreúva melhora o PBIT, mas tira faturamento; (5) lay-off fica abaixo do PBIT.\n"
                            "Total = Logística + Outras Atividades + Fábricas. Net Sales R9 = Resumo FY do R9 (R$ 93.023,97 mil).")

# ---------------- slide 3: consolidado por unidade (bridge de PBIT por unidade + tabela com mensagem)
s3 = new_slide(prs, layout, "TLM | Bridge consolidado por unidade R9 vs MP’27")
gain_label(s3)
text(s3, 36, 62, 900, 24, [[(f"PBIT de BRL {mi(pbit(tm))} mi: {mi(pbit(tm) - pbit(tr), True)} mi vs R9",
                             {"bold": True, "color": NAVY, "size": 16})]])
steps_u = []
rows_u = []
spec = [("Barueri", "BATotal", "BA Log"), ("Campinas", "CPLog", "CP Log"), ("Feira de\nSantana", "FSLog", "FS Log"),
        ("Cabreúva", "CBTotal", "CB"), ("Outras Ativ.", "P&L Outras Atividades", "OA"), ("Fábricas", "P&L Fab", "Fab")]
unit_msg = {
    "BATotal": f"Faturamento {mi(net_ba, True)} mi com volume estável; custo {mi(-(cost(D['BATotal']['MP27']) - cost(D['BATotal']['R09'])), True)} mi",
    "CPLog": f"Volume {p1(pct(D['CPLog']['MP27']['tons'], D['CPLog']['R09']['tons']))}%; faturamento fecha com o custo (PBIT zero)",
    "FSLog": f"Faturamento {mi(net_fs, True)} mi cobre custo de {mi(-d_(D['FSLog'], 'var') - d_(D['FSLog'], 'fix'))} mi a mais",
    "CBTotal": "Sem operação no MP’27: sem receita nem custo",
    "P&L Outras Atividades": f"BAMP {mi(d_(D['BAMP'], 'net'), True)} mi de receita; Campneus {mi(d_(D['Campneus'], 'net'), True)} mi",
    "P&L Fab": f"Faturamento {mi(net_fab, True)} mi e custo {mi(-(cost(D['P&L Fab']['MP27']) - cost(D['P&L Fab']['R09'])), True)} mi",
}
for lab, k, _ in spec:
    steps_u.append((lab, pbit(D[k]["MP27"]) - pbit(D[k]["R09"])))
steps_u.insert(4, ("SB/GV Log", sbgv_m - sbgv_r))
assert abs(pbit(tr) + sum(v for _, v in steps_u) - pbit(tm)) < 0.5
bridge(s3, 30, 560, 90, 500, pbit(tr), steps_u, "R9 TLM", "MP’27 TLM", bracket=fmt(pbit(tm) - pbit(tr), sign=True))
# tabela por unidade
cw = [76, 44, 44, 44, 160]
hdr = ["Unidade", "R9", "MP’27", "Δ", "Mensagem"]
body = []
for lab, k, _ in spec[:4]:
    body.append((lab.replace("\n", " "), pbit(D[k]["R09"]), pbit(D[k]["MP27"]), unit_msg[k]))
body.append(("SB/GV Log", sbgv_r, sbgv_m, "Só custo (estrutura de suporte)"))
for lab, k, _ in spec[4:]:
    body.append((lab.replace("\n", " "), pbit(D[k]["R09"]), pbit(D[k]["MP27"]), unit_msg[k]))
rows = [(n, fmt(a_, paren=False), fmt(b_, paren=False), (fmt(b_ - a_, sign=True), b_ - a_), m) for n, a_, b_, m in body]
text(s3, 580, 98, 360, 14, f"Δ PBIT: {fmt(pbit(tm) - pbit(tr), sign=True)}  (PBIT por unidade, BRL ’000)", size=10, bold=True, color=NAVY)
gt = s3.shapes.add_table(len(rows) + 2, len(cw), Pt(580), Pt(116), Pt(sum(cw)), Pt(28 * len(rows) + 40))
tbl = gt.table
pr = gt._element.graphic.graphicData.tbl.tblPr
pr.set("firstRow", "0")
pr.set("bandRow", "0")
sid = pr.find(A + "tableStyleId")
if sid is not None:
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
for i, w in enumerate(cw):
    tbl.columns[i].width = Pt(w)


def tput(r_, c_, t_, size=8, bold=False, color=BLACK, align=PP_ALIGN.LEFT, fill=None):
    ce = tbl.cell(r_, c_)
    ce.margin_top = ce.margin_bottom = Pt(2)
    ce.margin_left = ce.margin_right = Pt(4)
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    ce.text = ""
    p = ce.text_frame.paragraphs[0]
    p.alignment = align
    rr = p.add_run()
    rr.text = t_
    rr.font.size = Pt(size)
    rr.font.bold = bold
    rr.font.color.rgb = color
    if fill is None:
        ce.fill.background()
    else:
        ce.fill.solid()
        ce.fill.fore_color.rgb = fill
    for side in ("lnL", "lnR", "lnT"):
        cell_border(ce, side, "FFFFFF", 0)
    cell_border(ce, "lnB", "D9D9D9", 6350)


tbl.rows[0].height = Pt(24)
for c_, h in enumerate(hdr):
    tput(0, c_, h, size=8.5, bold=True, color=WHITE, align=PP_ALIGN.LEFT if c_ in (0, 4) else PP_ALIGN.RIGHT, fill=NAVY)
for i, (n, a_, b_, dd, m) in enumerate(rows, start=1):
    tbl.rows[i].height = Pt(28)
    f = MPFILL if i % 2 == 0 else None
    tput(i, 0, n, size=9, bold=True, fill=f)
    tput(i, 1, a_, size=9, align=PP_ALIGN.RIGHT, fill=f)
    tput(i, 2, b_, size=9, bold=True, align=PP_ALIGN.RIGHT, fill=f)
    tput(i, 3, dd[0], size=9, bold=True, color=GREEN if dd[1] > 0.5 else (RED if dd[1] < -0.5 else GREY),
         align=PP_ALIGN.RIGHT, fill=f)
    tput(i, 4, m, size=7.5, color=DGREY, fill=f)
last = len(rows) + 1
tbl.rows[last].height = Pt(24)
tput(last, 0, "TLM", size=9, bold=True, fill=LIGHT)
tput(last, 1, fmt(pbit(tr)), size=9, bold=True, align=PP_ALIGN.RIGHT, fill=LIGHT)
tput(last, 2, fmt(pbit(tm)), size=9, bold=True, align=PP_ALIGN.RIGHT, fill=LIGHT)
tput(last, 3, fmt(pbit(tm) - pbit(tr), sign=True), size=9, bold=True, color=RED, align=PP_ALIGN.RIGHT, fill=LIGHT)
tput(last, 4, f"Barueri concentra a queda; saída de Cabreúva compensa em parte", size=7.5, bold=True, color=NAVY, fill=LIGHT)
footer(s3, "SB/GV Log = SBLog + GVLog (centros só com custo).")
s3.notes_slide.notes_text_frame.text = "Bridge PBIT por unidade (BRL mil): R9 %.1f | " % pbit(tr) + " | ".join(
    f"{l.replace(chr(10), ' ')} {v:+,.1f}" for l, v in steps_u) + " | MP'27 %.1f" % pbit(tm)

# ---------------- slides por unidade
D_ba, D_cp, D_fs, D_cb, D_oa, D_fab = (D[k] for k in ("BATotal", "CPLog", "FSLog", "CBTotal", "P&L Outras Atividades", "P&L Fab"))


def sub_line(keys, field="net"):
    return " · ".join(f"{n} {fmt(D[k]['MP27'][field])} ({fmt(D[k]['R09'][field])})" for n, k in keys)


def sub_pbit(keys):
    return " · ".join(f"{n} {fmt(pbit(D[k]['MP27']))} ({fmt(pbit(D[k]['R09']))})" for n, k in keys)


oa_keys = [("BAMP", "BAMP"), ("Campneus", "Campneus"), ("SB Transp", "SBTransp"), ("Hevea", "Hevea"),
           ("CP Eng", "CPEng"), ("CP Check", "CPCheck"), ("BAQA", "BAQA")]
fab_keys = [("CP Fab", "CPFab"), ("FS Fab", "FSFab")]

entity_slide(prs, layout, "TLM Economics MP’27 – Barueri (BA)", "BATotal", items_gen(D_ba, "Barueri", {
    "net": f"Tarifa média por tonelada R$ {fmt(tarifa(D_ba, 'MP27'))} vs R$ {fmt(tarifa(D_ba, 'R09'))} no R9 "
           f"({p1(pct(tarifa(D_ba, 'MP27'), tarifa(D_ba, 'R09')))}%) com volume praticamente estável "
           f"({p1(pct(D_ba['MP27']['tons'], D_ba['R09']['tons']))}%).",
    "cost": "D&A inclui todo o efeito: fim do contrato IFRS16 do armazém (CC 7287) em jul/26, sem aluguel/amortização na base MP’27.",
    "h4": f"Volume {fmt(D_ba['MP27']['tons'] / 1000, dec=1)} mil t  |  HC {fmt(D_ba['MP27']['hc'])}",
    "b4": f"HC {fmt(D_ba['MP27']['hc'])} vs {fmt(D_ba['R09']['hc'])} no R9. Desconto de tarifa deliberado: ganho de eficiência para a Pirelli, com o TLM buscando o menor resultado."}))
entity_slide(prs, layout, "TLM Economics MP’27 – Campinas (CP)", "CPLog", items_gen(D_cp, "Campinas", {
    "net": f"Volume {p1(pct(D_cp['MP27']['tons'], D_cp['R09']['tons']))}% ({fmt(D_cp['MP27']['tons'] / 1000, dec=1)} mil t). "
           f"Faturamento fecha com o custo da unidade.",
    "cost": "",
    "h4": f"Tarifa média R$ {fmt(tarifa(D_cp, 'MP27'))}/t  |  R9 R$ {fmt(tarifa(D_cp, 'R09'))}/t",
    "b4": f"HC {fmt(D_cp['MP27']['hc'])} vs {fmt(D_cp['R09']['hc'])} no R9. PBIT zero reflete a tarifa calculada a custo; "
          f"a margem real depende da tarifa praticada."}))
entity_slide(prs, layout, "TLM Economics MP’27 – Feira de Santana (FS)", "FSLog", items_gen(D_fs, "Feira", {
    "net": f"Tarifa média R$ {fmt(tarifa(D_fs, 'MP27'))}/t vs R$ {fmt(tarifa(D_fs, 'R09'))}/t no R9; volume "
           f"{p1(pct(D_fs['MP27']['tons'], D_fs['R09']['tons']))}%.",
    "cost": "",
    "h4": f"Volume {fmt(D_fs['MP27']['tons'] / 1000, dec=1)} mil t  |  HC {fmt(D_fs['MP27']['hc'])}",
    "b4": f"HC {fmt(D_fs['MP27']['hc'])} vs {fmt(D_fs['R09']['hc'])} no R9. Única unidade logística com PBIT positivo no MP’27."}))
entity_slide(prs, layout, "TLM Economics MP’27 – Cabreúva (CB)", "CBTotal", [
    (f"PBIT {fmt(pbit(D_cb['MP27']))}  |  {fmt(pbit(D_cb['MP27']) - pbit(D_cb['R09']), sign=True)} vs R9",
     f"Sem operação no MP’27: sem receita, custo ou volume. R9: PBIT {fmt(pbit(D_cb['R09']))} (ROS {p1(ros(D_cb, 'R09'), False)}%)."),
    (f"Net Sales {fmt(D_cb['MP27']['net'])}  |  {fmt(net_cb, sign=True)} vs R9",
     f"R9 previa {fmt(D_cb['R09']['net'])} e {fmt(D_cb['R09']['tons'] / 1000, dec=1)} mil t; MP’26 {fmt(D_cb['MP26']['net'])}."),
    (f"Custo total {fmt(-cost(D_cb['MP27']))}  |  {fmt(-(cost(D_cb['MP27']) - cost(D_cb['R09'])), sign=True)} vs R9",
     f"Saem {fmt(-D_cb['R09']['lab'])} de custo do trabalho e {fmt(-(cost(D_cb['R09']) - D_cb['R09']['lab']))} dos demais custos do R9."),
    ("Efeito no TLM",
     f"Sair de Cabreúva melhora o PBIT em {fmt(pbit(D_cb['MP27']) - pbit(D_cb['R09']), sign=True)} vs R9 (unidade deficitária), "
     f"mas retira {fmt(-net_cb)} de faturamento."),
])
entity_slide(prs, layout, "TLM Economics MP’27 – Outras Atividades", "P&L Outras Atividades", items_gen(D_oa, "OA", {
    "net": f"Net Sales por atividade MP’27 (R9): {sub_line(oa_keys)}.",
    "cost": "",
    "h4": "PBIT por atividade MP’27 (R9)",
    "b4": sub_pbit(oa_keys) + ". Ecolabel/Outras Atividades ainda com premissas de simulação."}),
    note="Ecolabel/Outras Atividades com premissas de simulação (sem volume oficial).")
entity_slide(prs, layout, "TLM Economics MP’27 – Fábricas", "P&L Fab", items_gen(D_fab, "Fab", {
    "net": f"Net Sales MP’27 (R9): {sub_line(fab_keys)}.",
    "cost": "",
    "h4": "PBIT por fábrica MP’27 (R9)",
    "b4": sub_pbit(fab_keys) + f". HC {fmt(D_fab['MP27']['hc'])} vs {fmt(D_fab['R09']['hc'])} no R9."}))

prs.save(OUT)
print("salvo", OUT, len(prs.slides), "slides")
