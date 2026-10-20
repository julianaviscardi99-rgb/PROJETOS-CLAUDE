"""MP27 TLM - evidencia da sobreposicao CC x classe entre a base das areas (Rafaela/Jefferson) e o RH fixo (Iana).

Uso: python mp27_evidencia_sobreposicao.py <pasta_MP27> <saida.xlsx>
  <pasta_MP27> = pasta '1a Versao MP27 vs R09' (tem TLM_Automation e Arquivos Recebidos)

Le a aba Sobreposicao da base unica mais recente, traz as linhas de cada fonte (com a linha do arquivo original da
Rafaela e os comentarios dela) e monta um xlsx para anexar ao e-mail de duvida. Nao altera nenhum arquivo de origem.
"""
import glob
import os
import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MESES = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']
MES_PT = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
AZUL = '1F3864'


def ler(pasta):
    base = sorted(glob.glob(os.path.join(pasta, 'TLM_Automation', 'Draft MP27 Despesas - * - base unica.xlsx')))[-1]
    sob = pd.read_excel(base, sheet_name='Sobreposicao', header=3, dtype=str)
    sob = sob.iloc[:, :2]
    sob.columns = ['CC', 'Classe']
    sob = sob.dropna()
    d = pd.read_excel(base, sheet_name='Draft_MP27', header=3,
                      dtype={'Centro de Custo': str, 'Conta Fiscal': str, 'Conta Gestorial': str})
    areas = pd.read_excel(base, sheet_name='Base_Areas', header=3,
                          dtype={'Centro de Custo': str, 'Conta Fiscal': str, 'Conta Gestorial': str})
    iana_arq = glob.glob(os.path.join(pasta, 'Arquivos Recebidos', 'MP27 RH TLM 0981*.xlsx'))[0]
    iana = pd.read_excel(iana_arq, sheet_name='Base_0981', header=3, dtype={'CC': str})
    iana['Linha na aba Base_0981'] = iana.index + 5
    return base, sob, d, areas, iana


def main(pasta, saida):
    base, sob, d, areas, iana = ler(pasta)
    linhas_a, linhas_i, resumo, mensal = [], [], [], []
    for cc, cl in sob.itertuples(index=False):
        a = areas[(areas['Centro de Custo'] == cc) & (areas['Conta Fiscal'] == cl)].copy()
        i = iana[(iana['CC'] == cc) & (iana['CLASSE DE CUSTO'] == cl)].copy()
        a['Total'] = a[MESES].sum(axis=1)
        mes_i = [c for c in iana.columns if str(c).upper()[:3] in ('JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC')][:12]
        i['Total'] = i[mes_i].sum(axis=1)
        for _, r in a.iterrows():
            row = {'CC': cc, 'Classe': cl, 'Linha no arquivo da Rafaela (aba Base)': r.iloc[0], 'Area': r.iloc[1],
                   'Comentario da area': r.iloc[2], 'WH': r.iloc[4], 'CC Desc': r.iloc[6], 'Conta gestorial': r.iloc[8],
                   'Classe Desc': r.iloc[11], 'VAR.': r.iloc[12], 'ORIG.': r.iloc[13]}
            row.update({MES_PT[k]: r[MESES[k]] for k in range(12)})
            row['Total MP27'] = r['Total']
            linhas_a.append(row)
        for _, r in i.iterrows():
            row = {'CC': cc, 'Classe': cl, 'Linha na aba Base_0981': r['Linha na aba Base_0981'],
                   'Descricao da classe': r.iloc[7], 'VAR.': r.iloc[8], 'ORIG.': r.iloc[9]}
            row.update({MES_PT[k]: r[mes_i[k]] for k in range(12)})
            row['Total MP27'] = r['Total']
            linhas_i.append(row)
        wh = a['WH'].iloc[0] if len(a) else (d[(d['Centro de Custo'] == cc)]['WH'].iloc[0] if len(d) else '')
        desc = a['Classe Desc'].iloc[0] if len(a) else ''
        resumo.append({'CC': cc, 'Classe': cl, 'WH': wh, 'Descricao da classe': desc,
                       'Linhas nas areas': len(a), 'Total areas (R$)': a['Total'].sum(),
                       'Linhas na Iana': len(i), 'Total Iana (R$)': i['Total'].sum()})
        m = {'CC': cc, 'Classe': cl}
        for k in range(12):
            m['Areas ' + MES_PT[k]] = a[MESES[k]].sum()
        for k in range(12):
            m['Iana ' + MES_PT[k]] = i[mes_i[k]].sum()
        mensal.append(m)

    res = pd.DataFrame(resumo)
    res['Soma hoje no MP27 (R$)'] = res['Total areas (R$)'] + res['Total Iana (R$)']
    res['Se for duplicidade e sair a Iana (R$)'] = res['Total areas (R$)']
    res['Se for duplicidade e sair as areas (R$)'] = res['Total Iana (R$)']
    res['Maximo a ajustar (R$)'] = res[['Total areas (R$)', 'Total Iana (R$)']].min(axis=1)

    leia = pd.DataFrame({'Item': [
        'Objetivo', 'Regra atual', 'O que cada fonte e', 'Base de calculo', 'O que NAO prova', 'Pergunta', 'Arquivos'],
        'Descricao': [
            'Confirmar se as 5 combinacoes CC x classe que aparecem em mais de uma base sao o mesmo custo informado duas vezes ou itens diferentes.',
            'Decisao do Caetano (15/10): as tres bases somam, mesmo na mesma CC x classe, ate as areas responderem. Hoje o MP27 soma as duas fontes abaixo.',
            'Areas = 20261009 - MP\'27 Custos TLM_envio - Sem RH.xlsx (Jefferson, 16/10). Iana = MP27 RH TLM 0981, versao de 15/10 18h27 (so empresa 981).',
            'Totais MP27 do ano em R$, ja com a inflacao do arquivo de cada fonte. Os meses estao na aba Mes a mes.',
            'Valores diferentes entre as fontes nao descartam duplicidade (uma fonte pode trazer so parte do custo); valores iguais, em geral, confirmariam.',
            'O valor das areas ja inclui o que o RH (Iana) informou nessas classes? Se sim, qual das duas fontes deve ficar?',
            'Origem: ' + os.path.basename(base)]})

    with pd.ExcelWriter(saida, engine='openpyxl') as xw:
        leia.to_excel(xw, sheet_name='Leia-me', index=False)
        res.to_excel(xw, sheet_name='Resumo', index=False)
        pd.DataFrame(linhas_a).to_excel(xw, sheet_name='Linhas areas', index=False)
        pd.DataFrame(linhas_i).to_excel(xw, sheet_name='Linhas Iana', index=False)
        pd.DataFrame(mensal).to_excel(xw, sheet_name='Mes a mes', index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor=AZUL)
                c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.freeze_panes = 'A2'
            for j, col in enumerate(ws.columns, 1):
                larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, larg + 2), 60 if ws.title != 'Leia-me' else 120)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = '#,##0.00'
            if ws.title == 'Leia-me':
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical='top')
    print('Gravado:', saida)
    print(res[['CC', 'Classe', 'Total areas (R$)', 'Total Iana (R$)', 'Soma hoje no MP27 (R$)']].to_string())
    print('Soma:', round(res['Soma hoje no MP27 (R$)'].sum(), 2), '| maximo a ajustar:', round(res['Maximo a ajustar (R$)'].sum(), 2))


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
