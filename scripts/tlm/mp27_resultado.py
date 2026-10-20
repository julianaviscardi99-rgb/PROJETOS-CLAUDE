# -*- coding: utf-8 -*-
"""mp27_resultado.py -- monta o `Resultado TLM - MP27 vs R09.xlsm` com as formulas lendo a base do MP27.

OBJETIVO: o P&L do MP27 no mesmo esqueleto do Resultado mensal, sem redigitar nada -- quem muda o
          draft de despesas ou preenche o faturamento ve o efeito no P&L.
OPERACAO:
  [1] `Base Total MP27.xlsx` -- tabela `Intermediaria` (mesmas colunas que os 1.994 SUMIFS de cada
      aba de WH consomem: WH, Variabilidade, Desc.Resultado, Cenario, Mes/Ano, January..December),
      alimentada pela aba Draft_MP27 do `Draft MP27 Despesas - <data>.xlsx`. Cenario = Mes/Ano = "MP27".
  [2] copia o `Resultado TLM - Flash Set '26.xlsm` (ja com o LF R09) para TLM_Automation e, via Excel:
      troca o link [1] Base Total_Act_Flash_Fct -> Base Total MP27; poe "MP27" nas linhas 5 e 6 do
      'P&L - TLM' (as abas de WH apontam para la); liga as linhas Tons / Tarifa / Fixo / Ecolabel e as
      Outras Atividades da aba Faturamento a planilha `Faturamento MP27 - base R09` da pasta Volume;
      zera Cabreuva e MP Barueri (nao existem em 2027); recalcula e grava.
ALINHAMENTO: e o que o MP26 fez a mao (`Resultado TLM - MP26 vs R09.xlsx` ligado ao `Base Total
          MP26.XLSX`), agora reproduzivel. Nada e escrito no SAP; o original do Flash nao e tocado.

Pede Excel instalado (win32com) -- a gravacao com openpyxl quebraria os 24 links externos.
"""
import glob
import os
import re
import shutil
import sys
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(BASE_DIR, "scripts", "antidumping"))
from caderno import LINHA_CABECALHO, escrever_caderno                  # noqa: E402

sys.path.insert(0, os.path.join(BASE_DIR, "scripts", "tlm"))
import mp27_tarifa                                                     # noqa: E402

REDE = r"\\SRV054-01BR.group.pirelli.com\DAC_TLM"
ANO_MP = datetime.now().year + 1
ANO_BASE = ANO_MP - 1
AA_MP = ANO_MP % 100
CENARIO = "MP%d" % AA_MP
PASTA_MP = os.path.join(REDE, "MP%d" % AA_MP, "1\u00aa Vers\u00e3o MP%d vs R09" % AA_MP)
PASTA_SAIDA = os.path.join(PASTA_MP, "TLM_Automation")
PASTA_VOLUME = os.path.join(PASTA_MP, "Volume")
FLASH_ORIGEM = os.path.join(REDE, str(ANO_BASE), "2. Flash", "09.Setembro", "Resultado TLM - Flash Set '%d.xlsm" % (ANO_BASE % 100))
BASE_ACT_FLASH = os.path.join(REDE, str(ANO_BASE), "Base Total_Act_Flash_Fct.xlsx")
BASE_MP = os.path.join(PASTA_SAIDA, "Base Total MP%d.xlsx" % AA_MP)
RESULTADO = os.path.join(PASTA_SAIDA, "Resultado TLM - MP%d vs R09.xlsm" % AA_MP)

MESES = ["January", "February", "March", "April", "May", "June",
         "July", "August", "September", "October", "November", "December"]
COLS_TABELA = ["Source.Name", "Conta Gestorial", "Descri\u00e7\u00e3o", "Conta Fiscal", "Descri\u00e7\u00e3o_1",
               "Centro de Custo", "Mini-F\u00e1brica", "Descri\u00e7\u00e3o_2", "Tp.Custo (Pr\u00f3prio /  Aprop.)"] + MESES + [
               "Total Ano", "Usu\u00e1rio", "Data", "Hora", "Variabilidade", "WH", "Gestorial 2", "Cenario",
               "Desc.Resultado", "M\u00eas Despesa", "M\u00eas/Ano", "Chave", "Check Mdo"]

# Aba Faturamento do Resultado <- planilha de preenchimento (Unidade/Cliente, BU/Atividade, Linha).
# fator 1000: a planilha esta em R$ mil / k ton; a aba Faturamento em R$ / ton.
LIGACOES_FATURAMENTO = [
    (4, ("Barueri", "Car", "Tons"), 1000), (5, ("Barueri", "Car", "Tarifa"), 1),
    (12, ("Barueri", "Car", "Valor Liq.Fix"), 1000), (22, ("Barueri", "Car", "Ecolabel"), 1000),
    (33, ("Barueri", "Moto", "Tons"), 1000), (34, ("Barueri", "Moto", "Tarifa"), 1),
    (41, ("Barueri", "Moto", "Valor Liq.Fix"), 1000),
    (56, ("Campinas", "Car", "Tons"), 1000), (57, ("Campinas", "Car", "Tarifa"), 1),
    (64, ("Campinas", "Car", "Valor Liq.Fix"), 1000),
    (67, ("Campinas", "Moto", "Tons"), 1000), (68, ("Campinas", "Moto", "Tarifa"), 1),
    (75, ("Campinas", "Moto", "Valor Liq.Fix"), 1000),
    (137, ("Feira de Santana", "Car", "Tons"), 1000), (138, ("Feira de Santana", "Car", "Tarifa"), 1),
    (145, ("Feira de Santana", "Car", "Valor Liq.Fix"), 1000), (150, ("Feira de Santana", "Car", "Ecolabel"), 1000),
    (162, ("Comercial", "SB.Transp", "Valor Liq."), 1000), (166, ("Comercial", "Campneus", "Valor Liq."), 1000),
    (170, ("Pneus", "SB.Transp", "Valor Liq."), 1000), (189, ("Pneus", "Campneus BA", "Valor Liq."), 1000),
    (195, ("Pneus", "HEVEA", "Valor Liq."), 1000), (199, ("Pneus", "Q.A", "Valor Liq."), 1000),
    (203, ("Pneus", "FS.Check", "Valor Liq."), 1000), (212, ("Pneus", "Fitted", "Valor Liq."), 1000),
]
AJUSTES_FLASH_HERDADOS = [("BAADM", "J105"), ("BALog", "D24")]

# Aprendizes do RH (Iana) vem em classes de fabrica (M) sem conta gestorial na TLM e ficavam fora do P&L.
# No Actual 2026 (Base_KSB1) a TLM lanca aprendiz em N14211S001 -> 4236020 "Prestados - Servicos", fixo
# (decisao da Juliana, 20/10/2026).
CLASSES_APRENDIZ_IANA = ["N14215M001", "N14219M000"]
APRENDIZ_COMO_ACTUAL = {"Conta Fiscal": "N14211S001", "Classe Desc": "COM DESP Mao de obra (aprendiz)",
                        "Conta Gestorial": "4236020", "Conta Gestorial Desc": "Prestados - Serviços", "VAR.": "F"}

# Cabreuva (82-127) e MP Barueri (180-184) nao existem em 2027: zera o que alimenta os blocos.
# (181/184 e nao 182: a Area zerada vira divisor das formulas de custo e faturamento do MP BA)
ZERAR_FATURAMENTO = [82, 87, 92, 97, 106, 111, 119, 124, 181, 184]

# HC do MP (linha 19 das abas de WH): o Flash le o HC R09 (e agosto do HC R07) por SUMIF em pasta externa,
# que da #VALUE! com a pasta fechada. O HC do MP vai para a aba interna "HC MP27"; a chave da aba (WH2:
# BAADM, BALog, CPFab...) sai do CC pelo HC R09, que e onde o Flash ja fazia esse de-para.
PADRAO_HC_MP = "* - Base Geral HC*MP%d*.xlsx" % AA_MP
VERSAO_HC_MP = "10. MP%d_v0" % AA_MP
HC_R09 = os.path.join(REDE, str(ANO_BASE), "1. FCST", "R09", "HC R09 TLM.xlsx")
ABA_HC = "HC MP%d" % AA_MP
MESES_HC = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

# Coluna O (MP x R09): o Flash deixou diferencas de UM mes (ex. =I19-W19); no MP vale total N x total AB.
ABAS_SEM_COL_O = ("Faturamento", "Resumo", "Resumo Resultado", "Delta Armazém", "Parametros")
O_DE_UM_MES = re.compile(r"^=(SUM\()?[B-M]\d+(:[B-M]\d+\))?-(SUM\()?([P-Z]|AA)\d+(:([P-Z]|AA)\d+\))?$")
FIRST_COL_MP = 19          # coluna S = MP27 Jan na planilha de preenchimento (5 chaves + 12 meses de 2026 + total)



def contar_erros(ws):
    """Celulas com erro (#VALUE!, #REF!, #DIV/0!...) lendo o intervalo usado de uma vez: celula a celula via COM
    levava horas no Resultado (ida e volta Python x Excel para cada uma das centenas de milhares de celulas)."""
    v = ws.UsedRange.Value
    if not isinstance(v, tuple):
        v = ((v,),)
    return sum(1 for linha in v for x in linha if isinstance(x, int) and x < -2146826000)

def _ultimo(pasta, padrao):
    arqs = sorted(glob.glob(os.path.join(pasta, padrao)))
    if not arqs:
        raise FileNotFoundError("nenhum %s em %s" % (padrao, pasta))
    return arqs[-1]


# ----------------------------------------------------------------------------- [1] Base Total MP27
def desc_resultado_por_gestorial():
    """Conta Gestorial -> Desc.Resultado mais usado na tabela Intermediaria da base oficial 2026."""
    import collections
    from openpyxl import load_workbook as _lw
    ws = _lw(BASE_ACT_FLASH, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    cab = next(r for r in it if r and "Conta Gestorial" in r and "Desc.Resultado" in r)
    ig, idr = cab.index("Conta Gestorial"), cab.index("Desc.Resultado")
    cont = collections.defaultdict(collections.Counter)
    for r in it:
        if r and r[ig] is not None and r[idr]:
            cont[str(r[ig])][r[idr]] += 1
    return {cg: c.most_common(1)[0][0] for cg, c in cont.items()}


def montar_base(draft):
    d = pd.read_excel(draft, sheet_name="Draft_%s" % CENARIO, header=LINHA_CABECALHO - 1,
                      dtype={"Centro de Custo": str, "Conta Gestorial": str, "Conta Fiscal": str})
    apr = d["Conta Fiscal"].isin(CLASSES_APRENDIZ_IANA)
    for col, val in APRENDIZ_COMO_ACTUAL.items():
        d.loc[apr, col] = val
    print("  aprendizes da Iana reclassificados para %s / %s: %d linhas, R$ %.0f"
          % (APRENDIZ_COMO_ACTUAL["Conta Fiscal"], APRENDIZ_COMO_ACTUAL["Conta Gestorial"], apr.sum(), d.loc[apr, MESES].astype(float).sum().sum()))
    mapa = pd.read_excel(BASE_ACT_FLASH, sheet_name="Parametros", usecols=[0, 1], dtype=str).dropna()
    mapa = dict(zip(mapa.iloc[:, 0].str.strip(), mapa.iloc[:, 1]))
    desc_res = d["Conta Fiscal"].map(mapa)
    sem_mapa = d.loc[desc_res.isna(), "Conta Fiscal"].unique()
    # conta fiscal fora do mapa (ex.: Pyplan 4201000P): usa o Desc.Resultado da base 2026 para a mesma conta GESTORIAL
    # (de-para validado em 16/10; sem isso R$ 44,85 mi de MO cai em nome cru e fica fora do Resultado)
    desc_res = desc_res.fillna(d["Conta Gestorial"].map(desc_resultado_por_gestorial()))
    desc_res = desc_res.fillna(d["Conta Gestorial Desc"])
    agora = datetime.now()
    t = pd.DataFrame({
        "Source.Name": (d["Enviado por"] + " | " + d["Fonte"]) if "Fonte" in d else os.path.basename(draft),
        "Conta Gestorial": d["Conta Gestorial"], "Descri\u00e7\u00e3o": d["Conta Gestorial Desc"],
        "Conta Fiscal": d["Conta Fiscal"], "Descri\u00e7\u00e3o_1": d["Classe Desc"],
        "Centro de Custo": d["Centro de Custo"], "Mini-F\u00e1brica": "",
        "Descri\u00e7\u00e3o_2": d["Mini-F\u00e1brica Desc"], "Tp.Custo (Pr\u00f3prio /  Aprop.)": d["ORIG."],
    })
    for m in MESES:
        t[m] = d[m].astype(float).round(2)
    t["Total Ano"] = t[MESES].sum(axis=1)
    t["Usu\u00e1rio"] = os.environ.get("USERNAME", "").upper()
    t["Data"] = agora.strftime("%d/%m/%Y")
    t["Hora"] = agora.strftime("%H:%M:%S")
    t["Variabilidade"] = d["VAR."].where(d["Conta Fiscal"] != "N14301S056", "V")   # mesma excecao da base oficial
    t["WH"] = d["WH"]
    t["Gestorial 2"] = ""
    t["Cenario"] = CENARIO
    t["Desc.Resultado"] = desc_res
    t["M\u00eas Despesa"] = t["Total Ano"]
    t["M\u00eas/Ano"] = CENARIO
    t["Chave"] = t["Conta Fiscal"] + t["Centro de Custo"]
    t["Check Mdo"] = t["Conta Fiscal"].str[:4]
    t = t[COLS_TABELA]

    mp27_tarifa.guardar_versao(BASE_MP)
    escrever_caderno(BASE_MP, [("Intermediaria", "Base Total %s" % CENARIO,
                                "tabela Intermediaria lida pelo Resultado TLM - %s vs R09 - fonte: %s" % (CENARIO, os.path.basename(draft)), t)])
    wb = load_workbook(BASE_MP)
    ws = wb["Intermediaria"]
    ws.auto_filter.ref = None                       # a tabela traz o proprio filtro; os dois juntos corrompem o arquivo
    ref = "A%d:%s%d" % (LINHA_CABECALHO, get_column_letter(len(COLS_TABELA)), LINHA_CABECALHO + len(t))
    tab = Table(displayName="Intermediaria", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)
    wb.save(BASE_MP)
    print("  Base Total: %d linhas, total %s R$ %.0f, conta fiscal sem Desc.Resultado no mapa: %d %s"
          % (len(t), CENARIO, t["Total Ano"].sum(), len(sem_mapa), list(sem_mapa)[:8]))
    return t


# ----------------------------------------------------------------------------- [2] Resultado
def linhas_da_planilha(fat):
    """(Unidade/Cliente, BU/Atividade, Linha) -> numero da linha na aba Faturamento da planilha de preenchimento."""
    f = pd.read_excel(fat, sheet_name="Faturamento", header=LINHA_CABECALHO - 1, usecols=[0, 1, 2, 3])
    return {(r["Unidade / Cliente"], r["BU / Atividade"], r["Linha"]): LINHA_CABECALHO + 1 + i for i, r in f.iterrows()}


def hc_mp():
    """Linhas do HC do MP (versao MP27_v0) com a chave de aba WH2 pelo CC (de-para do HC R09)."""
    arq = _ultimo(PASTA_SAIDA, PADRAO_HC_MP)
    r09 = pd.read_excel(HC_R09, sheet_name="Base dados", header=1).dropna(subset=["WH2"])
    mapa = r09.groupby("CC")["WH2"].agg(lambda s: sorted(set(s)))
    if (mapa.str.len() > 1).any():
        raise RuntimeError("CC com mais de um WH2 no HC R09: %s" % mapa[mapa.str.len() > 1].to_dict())
    hc = pd.read_excel(arq, sheet_name="Base")
    hc.columns = [str(c).strip() for c in hc.columns]
    hc = hc[hc["VERSÃO"] == VERSAO_HC_MP].copy()
    hc.insert(0, "WH2", hc["CC"].map(mapa.str[0]))
    if hc["WH2"].isna().any():
        raise RuntimeError("CC do HC %s sem aba no HC R09: %s" % (CENARIO, sorted(hc.loc[hc["WH2"].isna(), "CC"].unique())))
    hc[MESES_HC] = hc[MESES_HC].fillna(0)
    return arq, hc[["WH2", "WH", "CC", "Descrição CC", "F/V", "Serviço", "Categoria", "Cargo"] + MESES_HC]


def gravar_hc(wb, arq, hc):
    """Aba interna com o HC do MP e a linha 19 (HC) das abas de WH somando dela, no lugar do SUMIF externo."""
    ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
    ws.Name = ABA_HC
    dados = [list(hc.columns)] + [[None if pd.isna(x) else (x.item() if hasattr(x, "item") else x) for x in r]
                                  for r in hc.itertuples(index=False)]
    ws.Range(ws.Cells(1, 1), ws.Cells(len(dados), len(hc.columns))).Value = dados
    ws.Cells(1, len(hc.columns) + 2).Value = "Fonte: %s (versão %s); WH2 pelo CC do %s" % (
        os.path.basename(arq), VERSAO_HC_MP, os.path.basename(HC_R09))
    ult = len(dados)
    trocadas = 0
    for aba in wb.Worksheets:
        if aba.Range("A19").Value != "HC" or "SUMIF(" not in str(aba.Range("B19").Formula):
            continue
        for j in range(12):                         # B..M da aba <- I..T (Jan..Dez) da aba do HC
            col = get_column_letter(9 + j)
            aba.Cells(19, 2 + j).Formula = "=SUMIF('%s'!$A$2:$A$%d,$A$3,'%s'!%s$2:%s$%d)" % (ABA_HC, ult, ABA_HC, col, col, ult)
        trocadas += 1
    print("  HC %s: %d linhas de %s na aba '%s'; linha 19 religada em %d abas" % (
        CENARIO, len(hc), os.path.basename(arq), ABA_HC, trocadas))


def coluna_o_total(wb):
    """Coluna O = total MP (N) - total R09 (AB) onde o Flash deixou a diferenca de um mes so."""
    trocadas = 0
    for aba in wb.Worksheets:
        if aba.Name in ABAS_SEM_COL_O:
            continue
        ult = aba.UsedRange.Row + aba.UsedRange.Rows.Count - 1
        for i, (f,) in enumerate(aba.Range("O1:O%d" % ult).Formula, 1):
            if isinstance(f, str) and O_DE_UM_MES.match(f.replace("$", "")):
                aba.Cells(i, 15).Formula = "=N%d-AB%d" % (i, i)
                trocadas += 1
    print("  coluna O: %d formulas de um mes trocadas por total N - AB" % trocadas)


def montar_resultado(fat):
    import win32com.client
    mp27_tarifa.guardar_versao(RESULTADO)
    shutil.copy2(FLASH_ORIGEM, RESULTADO)
    linhas = linhas_da_planilha(fat)
    arq_hc, hc = hc_mp()
    ref_fat = "'%s\\[%s]Faturamento'" % (os.path.dirname(fat), os.path.basename(fat))

    xl = mp27_tarifa.abrir_excel()
    xl.EnableEvents = False
    try:
        base = xl.Workbooks.Open(BASE_MP, UpdateLinks=0, ReadOnly=True)     # aberta: referencia estruturada externa so resolve com a fonte aberta
        wb = xl.Workbooks.Open(RESULTADO, UpdateLinks=0)
        antigo = [s for s in (wb.LinkSources(1) or []) if os.path.basename(str(s)).lower() == "base total_act_flash_fct.xlsx"]
        if not antigo:
            raise RuntimeError("link da Base Total_Act_Flash_Fct nao encontrado no Resultado")
        wb.ChangeLink(antigo[0], BASE_MP, 1)
        print("  link trocado:", os.path.basename(antigo[0]), "->", os.path.basename(BASE_MP))

        # ajustes digitados no fim da formula (herdados do Flash Set): BAADM!J105 "-390", BALog!D24 "-0,98"
        for aba, cel in AJUSTES_FLASH_HERDADOS:
            rng = wb.Worksheets(aba).Range(cel)
            nova = re.sub(r"(/1000\))\s*[+-]\s*\d+(\.\d+)?\s*$", r"\1", rng.Formula)
            if nova != rng.Formula:
                rng.Formula = nova
                print("  ajuste do Flash removido: %s!%s" % (aba, cel))
        gravar_hc(wb, arq_hc, hc)
        coluna_o_total(wb)

        pl = wb.Worksheets("P&L - TLM")
        pl.Range("B5:M5").Value = CENARIO
        pl.Range("B6:M6").Value = CENARIO
        wb.Worksheets("Parametros").Range("J3").Value = CENARIO

        ws = wb.Worksheets("Faturamento")
        ligadas = 0
        for linha, chave, fator in LIGACOES_FATURAMENTO:
            if chave not in linhas:
                print("  AVISO: %s nao existe na planilha de preenchimento - linha %d mantida" % (chave, linha))
                continue
            for j in range(12):
                origem = "%s!%s%d" % (ref_fat, get_column_letter(FIRST_COL_MP + j), linhas[chave])
                ws.Cells(linha, 2 + j).Formula = "=%s*%d" % (origem, fator) if fator != 1 else "=%s" % origem
            ligadas += 1
        for linha in ZERAR_FATURAMENTO:
            ws.Range(ws.Cells(linha, 2), ws.Cells(linha, 13)).Value = 0
        print("  Faturamento: %d linhas ligadas a planilha de preenchimento, %d zeradas (Cabreuva / MP BA)" % (ligadas, len(ZERAR_FATURAMENTO)))

        # SUMIF/SUMIFS em pasta externa FECHADA devolve #VALUE! (HC [21], KSB1 [23], faturamentos mensais):
        # abre as fontes que existem antes de recalcular, como estavam quando o Flash foi salvo.
        abertas = []
        for fonte in (wb.LinkSources(1) or []):
            fonte = str(fonte)
            if os.path.normcase(fonte) in (os.path.normcase(BASE_MP), os.path.normcase(fat)) or not os.path.exists(fonte):
                continue
            # outro Resultado (MP26 do bloco de comparacao) fica FECHADO: so ha referencia direta a celula, que le o valor
            # gravado; aberto, ele se recalcula com as fontes DELE fechadas e devolve #VALUE! em todo o bloco AD:AP
            if os.path.basename(fonte).lower().startswith("resultado tlm"):
                continue
            try:
                abertas.append(xl.Workbooks.Open(fonte, UpdateLinks=0, ReadOnly=True))
            except Exception as e:                                       # noqa: BLE001 - fonte velha/corrompida nao trava o resto
                print("  AVISO: nao abriu %s (%s)" % (os.path.basename(fonte), str(e)[:60]))
        modelo = xl.Workbooks.Open(fat, UpdateLinks=0, ReadOnly=True)
        print("  fontes abertas para o recalculo: %d" % (len(abertas) + 2))

        xl.CalculateFull()
        checks = {
            "Faturamento liquido FY (N13)": pl.Range("N13").Value,
            "Custos variaveis FY (N21)": pl.Range("N21").Value,
            "Custos fixos FY (N50)": pl.Range("N50").Value,
            "TOTAL COST FY (N135)": pl.Range("N135").Value,
            "PBIT FY (N124)": pl.Range("N124").Value,
        }
        erros = {}
        for nome in ("P&L - TLM", "Faturamento", "BALog", "CPLog", "P&L Log"):
            erros[nome] = contar_erros(wb.Worksheets(nome))
        wb.Save()
        wb.Close(SaveChanges=False)
        base.Close(SaveChanges=False)
        modelo.Close(SaveChanges=False)
        for a in abertas:
            if a is not None:               # Workbooks.Open devolve None para fonte que ja estava aberta (xlsm com macro)
                a.Close(SaveChanges=False)
    finally:
        try:
            xl.Quit()
        except Exception as e_quit:             # nao mascarar a excecao original (Excel ocupado: 'Call was rejected by callee')
            print('  AVISO: Excel nao fechou (%s)' % str(e_quit)[:60])
    return checks, erros


def main():
    draft = _ultimo(PASTA_SAIDA, "Draft %s Despesas - *.xlsx" % CENARIO)
    fat = _ultimo(PASTA_VOLUME, "Faturamento %s - base R09 - *.xlsx" % CENARIO)
    print("[1] Base Total %s <- %s" % (CENARIO, os.path.basename(draft)))
    t = montar_base(draft)
    print("[2] Resultado <- %s" % os.path.basename(FLASH_ORIGEM))
    checks, erros = montar_resultado(fat)
    print("\nSaida:", RESULTADO)
    for k, v in checks.items():
        print("  %-32s %14s" % (k, ("%.0f" % v) if isinstance(v, (int, float)) else v))
    print("  celulas de erro por aba:", erros)
    print("  base %s (R$ mil, custo positivo): %.0f" % (CENARIO, t["Total Ano"].sum() / 1000))


if __name__ == "__main__":
    main()
