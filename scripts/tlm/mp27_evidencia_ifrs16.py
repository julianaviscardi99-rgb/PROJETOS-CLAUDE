"""MP27 TLM - evidencia das linhas identicas (IFRS16 e outras) na base das areas (Rafaela/Jefferson).

Uso: python mp27_evidencia_ifrs16.py <pasta_MP27> <saida.xlsx>
  <pasta_MP27> = pasta '1a Versao MP27 vs R09' (tem Arquivos Recebidos)

Para cada par de linhas identicas (mesmo WH, conta fiscal, CC e os 12 meses) na versao 'MP27 Inf' do arquivo
'... Custos TLM_envio - Sem RH.xlsx', traz TODAS as linhas do bloco WH x conta x CC, o liquido do bloco com e sem a
linha repetida, e o FY do R09 (Base FY (R09)) na mesma chave, para a area dizer se a repeticao e intencional
(lancamentos de reversao/reclass) ou erro. Nao altera nenhum arquivo de origem.
"""
import glob
import os
import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AZUL = '1F3864'
MES_PT = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']


def main(pasta, saida):
    arq = glob.glob(os.path.join(pasta, 'Arquivos Recebidos', "*Custos TLM_envio - Sem RH.xlsx"))[0]
    x = pd.read_excel(arq, sheet_name='Base', header=3, dtype=str)
    infl = float(pd.read_excel(arq, sheet_name='Base', header=None, nrows=2).iloc[1, 0])
    cols = list(x.columns)
    x['Linha'] = x.index + 5
    meses = cols[9:21]
    for m in meses:
        x[m] = pd.to_numeric(x[m], errors='coerce').fillna(0.0)
    x['Total'] = x[meses].sum(axis=1)
    cv = [c for c in cols if 'rs' in c.lower() and 'o' in c.lower()][-1]
    inf = x[x[cv] == 'MP27 Inf'].copy()

    chave = ['WH', 'Conta Fisc', 'CC']
    rep = inf[inf.duplicated(chave + meses, keep=False) & (inf['Total'] != 0)]
    blocos = rep[chave].drop_duplicates().itertuples(index=False)

    r9 = pd.read_excel(glob.glob(os.path.join(pasta, 'Arquivos Recebidos', '*Base FY (R09)*.xlsx'))[0], sheet_name='Base', header=None)
    r9 = r9.iloc[3:]
    r9['cc'] = r9[6].astype(str).str.replace('.0', '', regex=False)

    resumo, linhas = [], []
    for wh, conta, cc in blocos:
        b = inf[(inf['WH'] == wh) & (inf['Conta Fisc'] == conta) & (inf['CC'] == cc)]
        grupo_rep = rep[(rep['WH'] == wh) & (rep['Conta Fisc'] == conta) & (rep['CC'] == cc)]
        l_rep = grupo_rep['Linha'].tolist()
        val_rep = grupo_rep['Total'].iloc[0]
        r = r9[(r9[1] == wh) & (r9[4] == conta) & (r9['cc'] == str(cc))]
        fy9 = pd.to_numeric(r[20], errors='coerce').sum() if len(r) else None
        liq_com = b['Total'].sum()
        liq_sem = liq_com - val_rep
        resumo.append({
            'WH': wh, 'Conta fiscal': conta, 'Descricao': b['Descri��o2'].iloc[0] if 'Descri��o2' in b else '', 'CC': cc,
            'Linhas do bloco': len(b), 'Linhas identicas (arquivo)': ', '.join(map(str, l_rep)),
            'Valor da linha repetida (R$)': val_rep,
            'Liquido do bloco como esta (R$)': liq_com,
            'Liquido se sair UMA repetida (R$)': liq_sem,
            'R09 FY 2026 mesma chave (R$)': fy9,
            'R09 x inflacao (R$)': None if fy9 is None else fy9 * infl,
            'Distancia ao R09 - como esta': None if fy9 is None else abs(liq_com - fy9 * infl),
            'Distancia ao R09 - sem a repetida': None if fy9 is None else abs(liq_sem - fy9 * infl)})
        for _, l in b.iterrows():
            row = {'WH': wh, 'Conta fiscal': conta, 'CC': cc, 'Linha no arquivo da Rafaela (aba Base)': l['Linha'],
                   'Repetida?': 'SIM' if l['Linha'] in l_rep else '', 'Comentario': l.iloc[6], 'VAR.': l.iloc[8]}
            row.update({MES_PT[k]: l[meses[k]] for k in range(12)})
            row['Total Ano'] = l['Total']
            linhas.append(row)

    res = pd.DataFrame(resumo)
    leia = pd.DataFrame({'Item': ['Objetivo', 'O que foi encontrado', 'Como ler o Resumo', 'Limite da evidencia', 'Pergunta'],
                         'Descricao': [
        'Confirmar com a area se as linhas identicas da base de custos (principalmente IFRS16 de Barueri) sao lancamentos intencionais ou repeticao por engano.',
        'Pares de linhas com os 12 meses iguais na versao MP27 Inf (a que entra no Resultado). Em varios casos o par faz parte de um bloco de 4 lancamentos (+X, -X, -X, +Y), tipico de reversao/reclassificacao.',
        'Compare a distancia ao R09: se tirar a linha repetida afasta o valor do R09 (que tem uma linha so), a repeticao provavelmente e intencional; se aproxima, e candidata a duplicidade.',
        'O R09 e FY 2026 e usa outra estrutura; serve como referencia de ordem de grandeza, nao como prova. Quem confirma e a area que montou a base.',
        'Estas linhas repetidas sao intencionais? Se nao, qual delas deve sair?']})

    with pd.ExcelWriter(saida, engine='openpyxl') as xw:
        leia.to_excel(xw, sheet_name='Leia-me', index=False)
        res.to_excel(xw, sheet_name='Resumo', index=False)
        pd.DataFrame(linhas).to_excel(xw, sheet_name='Linhas do bloco', index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor=AZUL)
                c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.freeze_panes = 'A2'
            for j, col in enumerate(ws.columns, 1):
                larg = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(j)].width = min(max(11, larg + 2), 120 if ws.title == 'Leia-me' else 45)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = '#,##0.00'
                    if ws.title == 'Leia-me':
                        c.alignment = Alignment(wrap_text=True, vertical='top')
            if ws.title == 'Linhas do bloco':
                amarelo = PatternFill('solid', fgColor='FFF2CC')
                for row in ws.iter_rows(min_row=2):
                    if row[4].value == 'SIM':
                        for c in row:
                            c.fill = amarelo
    print('Gravado:', saida, '| inflacao', infl)
    print(res.drop(columns=['Descricao']).round(0).to_string())


if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
