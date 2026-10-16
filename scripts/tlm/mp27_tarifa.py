# -*- coding: utf-8 -*-
"""mp27_tarifa.py -- raciocinio da tarifa fixa e variavel do MP27 dentro da planilha de faturamento.

OBJETIVO: quem abre o `Faturamento MP27 - base R09` ve como a tarifa nasce do custo (igual ao MP26),
          de onde sai cada despesa, o % de BU e o comparativo contra o R09 -- e, quando a tecnica
          entregar as toneladas, a tarifa e o faturamento saem sozinhos.
OPERACAO: acrescenta 5 abas a planilha ja gravada pelo mp27_faturamento_base.py:
            Custos      -- linhas 21 / 50 / 83 das abas de WH do `Resultado TLM - MP27 vs R09.xlsm`
                           (MP27 = colunas B:M, R09 = colunas P:AA), em valor, com a celula de origem
            Area_m2     -- m2 utilizados por unidade x BU: R09 (aba Faturamento do Resultado R09) e MP27
                           (amarelo, comeca igual a dezembro do R09)
            BU          -- % de BU: coeficiente de empenho de M.O do MP26 x % pela tonelada carregada
            Tarifa      -- a conta, unidade por unidade, por formula
            Comparativo -- R09 x MP27 (toneladas, area, custo, tarifa, faturamento)
          e troca as celulas MP27 de Tarifa e Valor Liq.Fix da aba Faturamento por formula ligada a aba Tarifa.
ALINHAMENTO: e a mesma regra das abas `Tarifa` / `Tarifa por B.U` do `Resultado TLM - MP26 vs R09.xlsx`
          (out/25), so que explicada e rastreavel. Nada e escrito no SAP; os Resultados so sao lidos.

Regra do MP26 (lida nas formulas do arquivo de out/25):
  Tarifa variavel (R$/t)   = Custo variavel da WH (linha 21) x % BU / toneladas da BU x (1 + markup)
                             markup: Barueri 2 %, Campinas 0 %, Feira 0 %. % BU = coeficiente de empenho de
                             M.O (Barueri 76/24, Campinas 50/50). Questoes trabalhistas ficam fora.
  Tarifa fixa (R$/m2/mes)  = (Custo fixo da WH (linha 50) - Questoes trabalhistas (linha 83) - receita de
                             Ecolabel) / m2 utilizados do ano. Uma tarifa por unidade, igual para Car e Moto.
  Faturamento bruto        = liquido / (1 - impostos): 11,25 % Barueri, 14,25 % Campinas e Feira.
"""
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import sys
BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(BASE_DIR, "scripts", "copc"))
from copc_estilo import AZUL, MOEDA, THIN                              # noqa: E402

MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
AMARELO = PatternFill("solid", fgColor="FFF2CC")
VERDE = PatternFill("solid", fgColor="E2EFDA")         # calculado pela aba Tarifa (pode sobrescrever)
CINZA = PatternFill("solid", fgColor="EDEDED")
SECAO = PatternFill("solid", fgColor="D9E1F2")
CINZA_TEXTO = "595959"

# Unidade -> WH do P&L que pagam a tarifa, BU, coeficiente MP26, markup, imposto, linhas da aba Faturamento
# do Resultado R09 com o m2 de cada BU.
UNIDADES = [
    {"nome": "Barueri", "wh": ["BALog", "BAADM"], "bus": ["Car", "Moto"], "coef": {"Car": 0.76, "Moto": 0.24},
     "markup": 0.02, "imposto": 0.1125, "m2_r09": {"Car": 9, "Moto": 38}, "ecolabel": True},
    {"nome": "Campinas", "wh": ["CPLog"], "bus": ["Car", "Moto"], "coef": {"Car": 0.5, "Moto": 0.5},
     "markup": 0.0, "imposto": 0.1425, "m2_r09": {"Car": 61, "Moto": 72}, "ecolabel": False},
    {"nome": "Feira de Santana", "wh": ["FSLog"], "bus": ["Car"], "coef": {"Car": 1.0},
     "markup": 0.0, "imposto": 0.1425, "m2_r09": {"Car": 142}, "ecolabel": False},
]
LINHAS_PL = [(21, "Custos variaveis"), (50, "Custos fixos"), (83, "Questoes trabalhistas")]
COLS_CENARIO = {"MP27": 2, "R09": 16}          # B:M = cenario do arquivo (MP27), P:AA = LF (R09)


# ----------------------------------------------------------------------------- leitura

def contar_erros(ws):
    """Celulas com erro (#VALUE!, #REF!, #DIV/0!...) lendo o intervalo usado de uma vez: celula a celula via COM
    levava horas no Resultado (ida e volta Python x Excel para cada uma das centenas de milhares de celulas)."""
    v = ws.UsedRange.Value
    if not isinstance(v, tuple):
        v = ((v,),)
    return sum(1 for linha in v for x in linha if isinstance(x, int) and x < -2146826000)

def ler_custos(resultado_mp):
    """{(wh, linha, cenario): [12 meses]} em R$ mil, sinal do P&L (custo negativo)."""
    wb = load_workbook(resultado_mp, data_only=True, read_only=True)
    out = {}
    for u in UNIDADES:
        for wh in u["wh"]:
            ws = wb[wh]
            linhas = {i: r for i, r in enumerate(ws.iter_rows(min_row=1, max_row=90, max_col=27, values_only=True), 1)}
            for ln, _ in LINHAS_PL:
                for cen, c0 in COLS_CENARIO.items():
                    out[(wh, ln, cen)] = [float(linhas[ln][c0 - 1 + j] or 0) for j in range(12)]
    wb.close()
    return out


def ler_area_r09(resultado_r09):
    """{(unidade, bu): [12 meses de m2]} da aba Faturamento do Resultado R09."""
    wb = load_workbook(resultado_r09, data_only=True, read_only=True)
    ws = wb["Faturamento"]
    linhas = {i: r for i, r in enumerate(ws.iter_rows(min_row=1, max_row=150, max_col=13, values_only=True), 1)}
    out = {}
    for u in UNIDADES:
        for bu, ln in u["m2_r09"].items():
            out[(u["nome"], bu)] = [float(linhas[ln][j + 1] or 0) for j in range(12)]
    wb.close()
    return out


# ----------------------------------------------------------------------------- estilo
def _titulo(ws, titulo, subtitulo):
    ws.sheet_view.showGridLines = False
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, size=10, color=CINZA_TEXTO)


def _cabecalho(ws, linha, nomes):
    for j, n in enumerate(nomes, 1):
        c = ws.cell(row=linha, column=j, value=n)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN


def _cel(ws, linha, col, valor, fmt=MOEDA, fill=None, bold=False):
    c = ws.cell(row=linha, column=col, value=valor)
    c.border = THIN
    if fmt and not isinstance(valor, str) or (isinstance(valor, str) and valor.startswith("=")):
        c.number_format = fmt or MOEDA
    if fill:
        c.fill = fill
    if bold:
        c.font = Font(bold=True)
    return c


def _larguras(ws, larguras):
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


# ----------------------------------------------------------------------------- abas
def aba_custos(wb, custos, nome_resultado):
    ws = wb.create_sheet("Custos")
    _titulo(ws, "Custos que formam a tarifa - MP27 e R09 (R$ mil, sinal do P&L: custo negativo)",
            "Fonte: %s - abas de WH, linhas 21 (custos variaveis), 50 (custos fixos) e 83 (questoes trabalhistas); "
            "MP27 = colunas B:M, R09 = colunas P:AA (LF). Valores copiados: rodar o script de novo se o draft mudar." % nome_resultado)
    cols = ["Unidade", "WH", "Linha P&L", "Cenario"] + MESES + ["Total", "Origem"]
    _cabecalho(ws, 4, cols)
    ref = {}                                     # (wh, linha, cenario) -> celula do total
    r = 5
    for u in UNIDADES:
        for wh in u["wh"]:
            for ln, desc in LINHAS_PL:
                for cen, c0 in COLS_CENARIO.items():
                    vals = custos[(wh, ln, cen)]
                    for j, v in enumerate([u["nome"], wh, "%d - %s" % (ln, desc), cen]):
                        _cel(ws, r, j + 1, v, fmt=None)
                    for m in range(12):
                        _cel(ws, r, 5 + m, round(vals[m], 6))
                    _cel(ws, r, 17, "=SUM(E%d:P%d)" % (r, r), bold=True)
                    _cel(ws, r, 18, "%s!%s%d:%s%d" % (wh, get_column_letter(c0), ln, get_column_letter(c0 + 11), ln), fmt=None)
                    ref[(wh, ln, cen)] = "Custos!$Q$%d" % r
                    r += 1
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:R%d" % (r - 1)
    _larguras(ws, [18, 9, 26, 9] + [11] * 12 + [13, 22])
    return ref


def aba_area(wb, area_r09):
    ws = wb.create_sheet("Area_m2")
    _titulo(ws, "Area utilizada (m2) - R09 x MP27",
            "R09 = aba Faturamento do Resultado TLM - R09 '26 (linhas 'm2 utilizada'). MP27 em AMARELO = digitar; "
            "comeca igual a dezembro do R09. A soma do ano (m2-mes) e o divisor da tarifa fixa.")
    cols = ["Unidade", "BU"] + ["R09 %s" % m for m in MESES] + ["R09 Total"] + ["MP27 %s" % m for m in MESES] + ["MP27 Total"]
    _cabecalho(ws, 4, cols)
    ref = {}
    r = 5
    for u in UNIDADES:
        for bu in u["bus"]:
            vals = area_r09[(u["nome"], bu)]
            _cel(ws, r, 1, u["nome"], fmt=None)
            _cel(ws, r, 2, bu, fmt=None)
            for m in range(12):
                _cel(ws, r, 3 + m, vals[m], fmt="#,##0")
            _cel(ws, r, 15, "=SUM(C%d:N%d)" % (r, r), fmt="#,##0", bold=True)
            for m in range(12):
                _cel(ws, r, 16 + m, vals[11], fmt="#,##0", fill=AMARELO)
            _cel(ws, r, 28, "=SUM(P%d:AA%d)" % (r, r), fmt="#,##0", bold=True)
            ref[(u["nome"], bu)] = {"r09": "Area_m2!$O$%d" % r, "mp": "Area_m2!$AB$%d" % r, "linha": r}
            r += 1
    ws.freeze_panes = "C5"
    _larguras(ws, [18, 8] + [9] * 13 + [9] * 13)
    return ref


def aba_bu(wb, fat_ref):
    ws = wb.create_sheet("BU")
    _titulo(ws, "% de BU - coeficiente de empenho de M.O (MP26) x tonelada carregada",
            "O custo variavel da unidade e dividido entre Car e Moto pelo % de BU. Criterio em C4: 'Empenho M.O' = regra do MP26 "
            "(padrao); 'Toneladas' = % pela tonelada MP27 da aba Faturamento (vazia = vale o coeficiente). Colunas J e K: a tarifa nos dois.")
    ws["A4"] = "Criterio do % BU"
    ws["A4"].font = Font(bold=True)
    c = ws["C4"]
    c.value = "Empenho M.O"
    c.fill = AMARELO
    c.border = THIN
    c.font = Font(bold=True)
    dv = DataValidation(type="list", formula1='"Toneladas,Empenho M.O"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C4")
    ws["D4"] = "Toneladas | Empenho M.O"
    ws["D4"].font = Font(italic=True, color=CINZA_TEXTO)

    cols = ["Unidade", "BU", "Coef. empenho M.O (MP26)", "Tons R09 (k t)", "% tons R09",
            "Tons MP27 (k t)", "% tons MP27", "% BU usado no MP27", "Leitura",
            "Tarifa var. pelo coeficiente (R$/t)", "Tarifa var. pela tonelada (R$/t)"]
    _cabecalho(ws, 6, cols)
    ref = {}
    r = 7
    for u in UNIDADES:
        r0 = r
        rows = {}
        for bu in u["bus"]:
            rows[bu] = r
            r += 1
        r1 = r - 1
        for bu, rr in rows.items():
            f = fat_ref[(u["nome"], bu, "Tons")]
            _cel(ws, rr, 1, u["nome"], fmt=None)
            _cel(ws, rr, 2, bu, fmt=None)
            _cel(ws, rr, 3, u["coef"][bu], fmt="0.0%", fill=AMARELO)
            _cel(ws, rr, 4, "=Faturamento!$R$%d" % f)
            _cel(ws, rr, 5, "=IFERROR(D%d/SUM(D$%d:D$%d),0)" % (rr, r0, r1), fmt="0.0%")
            _cel(ws, rr, 6, "=Faturamento!$AE$%d" % f)
            _cel(ws, rr, 7, "=IFERROR(F%d/SUM(F$%d:F$%d),0)" % (rr, r0, r1), fmt="0.0%")
            _cel(ws, rr, 8, '=IF(AND($C$4="Toneladas",SUM(F$%d:F$%d)>0),G%d,C%d)' % (r0, r1, rr, rr), fmt="0.0%", bold=True)
            ref[(u["nome"], bu)] = "BU!$H$%d" % rr
            ref[("_linha", u["nome"], bu)] = rr
        leitura = ("pelas toneladas a tarifa variavel por tonelada fica IGUAL em Car e Moto; o coeficiente "
                   "carrega mais custo na BU que empenha mais M.O por tonelada (Moto)") if len(u["bus"]) > 1 else "uma BU so: 100 %"
        _cel(ws, r0, 9, leitura, fmt=None)
        _cel(ws, r, 1, "Total %s" % u["nome"], fmt=None, bold=True)
        for col in (3, 4, 5, 6, 7, 8):
            L = get_column_letter(col)
            _cel(ws, r, col, "=SUM(%s%d:%s%d)" % (L, r0, L, r1), fmt="0.0%" if col in (3, 5, 7, 8) else MOEDA, bold=True)
        r += 2
    _larguras(ws, [20, 8, 16, 13, 11, 14, 11, 15, 80, 16, 16])
    return ref


def bu_tarifas(wb, bu_ref, tar_ref):
    """Na aba BU, a tarifa variavel MP27 pelos dois criterios - para ver o efeito antes de trocar o C4."""
    ws = wb["BU"]
    for u in UNIDADES:
        T = tar_ref[u["nome"]]
        cv, mk, tons = "Tarifa!$D$%d" % T["cv"], "Tarifa!$D$%d" % T["mk"], "Tarifa!$D$%d" % T["tons"]
        for bu in u["bus"]:
            rr = bu_ref[("_linha", u["nome"], bu)]
            _cel(ws, rr, 10, '=IFERROR(%s*C%d/F%d*(1+%s),"")' % (cv, rr, rr, mk))
            _cel(ws, rr, 11, '=IFERROR(%s/%s*(1+%s),"")' % (cv, tons, mk))


def aba_tarifa(wb, fat_ref, custos_ref, area_ref, bu_ref):
    """Bloco por unidade. Colunas: A rotulo | B Car | C Moto | D Unidade | E unid. | F de onde vem."""
    ws = wb.create_sheet("Tarifa")
    _titulo(ws, "Como a tarifa MP27 e montada - regra do MP26 (R$ mil, k t, R$/t, R$/m2)",
            "Tarifa variavel = custo variavel x %% BU / toneladas x (1 + markup).  Tarifa fixa = (custo fixo - questoes "
            "trabalhistas - Ecolabel) / m2 do ano.  Amarelo = premissa editavel. A aba Faturamento le daqui (celulas verdes).")
    ref = {}
    r = 4
    for u in UNIDADES:
        nome = u["nome"]
        for col in range(1, 7):
            ws.cell(row=r, column=col).fill = SECAO
        ws.cell(row=r, column=1, value=nome.upper()).font = Font(bold=True, size=12)
        r += 1
        _cabecalho(ws, r, ["Passo", "Car", "Moto" if "Moto" in u["bus"] else "-", "Unidade", "Unid.", "De onde vem / raciocinio"])
        r += 1
        L = {}                                  # nome do passo -> linha

        def linha(chave, rotulo, unid, origem, d=None, por_bu=None, fmt=MOEDA, fill=None, bold=False):
            nonlocal r
            L[chave] = r
            _cel(ws, r, 1, rotulo, fmt=None, bold=bold)
            for j, bu in enumerate(["Car", "Moto"]):
                if por_bu and bu in u["bus"]:
                    _cel(ws, r, 2 + j, por_bu(bu), fmt=fmt, fill=fill, bold=bold)
                else:
                    _cel(ws, r, 2 + j, None, fmt=None)
            _cel(ws, r, 4, d, fmt=fmt, fill=fill if d is not None and por_bu is None else None, bold=bold) if d is not None else _cel(ws, r, 4, None, fmt=None)
            _cel(ws, r, 5, unid, fmt=None)
            _cel(ws, r, 6, origem, fmt=None)
            r += 1

        whs = u["wh"]
        soma = lambda ln, cen: "-(" + "+".join(custos_ref[(w, ln, cen)] for w in whs) + ")"   # noqa: E731
        col_bu = {"Car": "B", "Moto": "C"}

        # --- variavel
        linha("cv", "1. Custo variavel MP27", "R$ mil",
              "Resultado MP27 vs R09, %s linha 21 CUSTOS VARIAVEIS (aba Custos)" % " + ".join(whs), d="=" + soma(21, "MP27"))
        linha("pbu", "2. % BU", "%", "aba BU (criterio em BU!C4)", por_bu=lambda bu: "=" + bu_ref[(nome, bu)], fmt="0.0%")
        linha("cvbu", "3. Custo variavel da BU = 1 x 2", "R$ mil", "",
              por_bu=lambda bu: "=$D$%d*%s%d" % (L["cv"], col_bu[bu], L["pbu"]), d="=SUM(B%d:C%d)" % (r, r))
        linha("tons", "4. Toneladas MP27", "k t", "aba Faturamento, linha Tons (Total MP27) - vem do Volume Movimentado MP27 da tecnica",
              por_bu=lambda bu: "=Faturamento!$AE$%d" % fat_ref[(nome, bu, "Tons")], d="=SUM(B%d:C%d)" % (r, r))
        linha("cvt", "5. Custo variavel por tonelada = 3 / 4", "R$/t", "",
              por_bu=lambda bu: '=IFERROR(%s%d/%s%d,"")' % (col_bu[bu], L["cvbu"], col_bu[bu], L["tons"]),
              d='=IFERROR(D%d/D%d,"")' % (L["cv"], L["tons"]))
        linha("mk", "6. Markup sobre o custo variavel", "%",
              "MP26: Barueri +2 %, Campinas e Feira 0 % (Tarifa por B.U, coluna AF)", d=u["markup"], fmt="0.0%", fill=AMARELO)
        linha("tv", "7. TARIFA VARIAVEL MP27 = 5 x (1 + 6)", "R$/t", "vai para a linha Tarifa da aba Faturamento (verde)",
              por_bu=lambda bu: '=IFERROR(%s%d*(1+$D$%d),"")' % (col_bu[bu], L["cvt"], L["mk"]),
              d='=IFERROR(D%d/D%d,"")' % (r + 7, L["tons"]), bold=True)

        # --- fixo
        linha("cf", "8. Custo fixo MP27", "R$ mil",
              "Resultado MP27 vs R09, %s linha 50 CUSTOS FIXOS (aba Custos)" % " + ".join(whs), d="=" + soma(50, "MP27"))
        linha("qt", "9. (-) Questoes trabalhistas", "R$ mil",
              "linha 83 - fica FORA da tarifa (como no MP26); entra no resultado como custo da TLM", d="=" + soma(83, "MP27"))
        if u["ecolabel"]:
            linha("eco", "10. (-) Receita Ecolabel MP27", "R$ mil",
                  "aba Faturamento, Ecolabel Barueri Car (Total MP27): a receita de etiquetas abate o fixo a recuperar (MP26)",
                  d="=Faturamento!$AE$%d" % fat_ref[(nome, "Car", "Ecolabel")])
        else:
            linha("eco", "10. (-) Receita Ecolabel MP27", "R$ mil", "nao se aplica nesta unidade no MP26", d=0)
        linha("cfr", "11. Custo fixo a recuperar = 8 - 9 - 10", "R$ mil", "", d="=D%d-D%d-D%d" % (L["cf"], L["qt"], L["eco"]), bold=True)
        linha("m2", "12. Area utilizada MP27 (soma do ano)", "m2-mes", "aba Area_m2 (MP27 em amarelo)",
              por_bu=lambda bu: "=" + area_ref[(nome, bu)]["mp"], d="=SUM(B%d:C%d)" % (r, r), fmt="#,##0")
        linha("tf", "13. TARIFA FIXA MP27 = 11 x 1000 / 12", "R$/m2/mes",
              "uma tarifa por unidade, igual para Car e Moto; vai para Valor Liq.Fix = tarifa x m2 do mes",
              d='=IFERROR(D%d*1000/D%d,"")' % (L["cfr"], L["m2"]), bold=True)

        # --- faturamento resultante
        linha("fv", "14. Faturamento liquido variavel = 7 x 4", "R$ mil", "",
              por_bu=lambda bu: '=IFERROR(%s%d*%s%d,0)' % (col_bu[bu], L["tv"], col_bu[bu], L["tons"]),
              d="=SUM(B%d:C%d)" % (r, r))
        # a linha 7 (unidade) aponta para a 14: tarifa media ponderada
        ws.cell(row=L["tv"], column=4).value = '=IFERROR(D%d/D%d,"")' % (L["fv"], L["tons"])
        linha("ff", "15. Faturamento liquido fixo = 13 x 12 / 1000", "R$ mil", "",
              por_bu=lambda bu: '=IFERROR($D$%d*%s%d/1000,0)' % (L["tf"], col_bu[bu], L["m2"]),
              d="=SUM(B%d:C%d)" % (r, r))
        linha("fl", "16. Faturamento liquido total", "R$ mil", "",
              por_bu=lambda bu: "=%s%d+%s%d" % (col_bu[bu], L["fv"], col_bu[bu], L["ff"]),
              d="=SUM(B%d:C%d)" % (r, r), bold=True)
        linha("imp", "17. Impostos sobre o faturamento", "%",
              "gross-up do MP26: Barueri 11,25 %, Campinas e Feira 14,25 %", d=u["imposto"], fmt="0.00%", fill=AMARELO)
        linha("fb", "18. Faturamento bruto = 16 / (1 - 17)", "R$ mil", "",
              por_bu=lambda bu: "=%s%d/(1-$D$%d)" % (col_bu[bu], L["fl"], L["imp"]), d="=SUM(B%d:C%d)" % (r, r))
        linha("mg", "19. Margem = 16 - 1 - 11", "R$ mil",
              "o que sobra e so o markup do variavel (+ Ecolabel em Barueri); zero = tarifa a custo",
              d="=D%d-D%d-D%d" % (L["fl"], L["cv"], L["cfr"]))
        ref[nome] = dict(L)
        r += 1

    ws.freeze_panes = "B4"
    _larguras(ws, [44, 13, 13, 14, 10, 95])
    return ref


def aba_comparativo(wb, fat_ref, custos_ref, area_ref, tar_ref):
    ws = wb.create_sheet("Comparativo_R09")
    _titulo(ws, "MP27 x R09 - toneladas, area, custo, tarifa e faturamento",
            "R09 'praticado' = o que esta no R09 '26 (Act jan-ago + Fcst set-dez). R09 'pelo metodo' = a mesma conta da "
            "aba Tarifa aplicada ao custo e ao volume do R09 - mostra o quanto o R09 esta acima/abaixo do custo.")
    cols = ["Unidade", "BU", "Indicador", "Unid.", "R09 praticado", "R09 pelo metodo MP26", "MP27",
            "Delta MP27 - R09", "Delta %", "Leitura"]
    _cabecalho(ws, 4, cols)
    r = 5
    for u in UNIDADES:
        nome, T = u["nome"], tar_ref[u["nome"]]
        whs = u["wh"]
        soma = lambda ln, cen: "-(" + "+".join(custos_ref[(w, ln, cen)] for w in whs) + ")"   # noqa: E731
        col_bu = {"Car": "B", "Moto": "C"}
        cf_r09 = "(%s-%s-%s)" % (soma(50, "R09"), soma(83, "R09"),
                                 ("Faturamento!$R$%d" % fat_ref[(nome, "Car", "Ecolabel")]) if u["ecolabel"] else "0")
        m2_r09_un = "(" + "+".join(area_ref[(nome, b)]["r09"] for b in u["bus"]) + ")"
        tf_r09_metodo = "%s*1000/%s" % (cf_r09, m2_r09_un)

        def lin(bu, ind, unid, r09, r09m, mp, leitura="", fmt=MOEDA, bold=False):
            nonlocal r
            for j, v in enumerate([nome, bu, ind, unid]):
                _cel(ws, r, j + 1, v, fmt=None, bold=bold)
            _cel(ws, r, 5, r09, fmt=fmt, bold=bold)
            _cel(ws, r, 6, r09m, fmt=fmt, bold=bold) if r09m else _cel(ws, r, 6, None, fmt=None)
            _cel(ws, r, 7, mp, fmt=fmt, bold=bold)
            _cel(ws, r, 8, '=IFERROR(G%d-E%d,"")' % (r, r), fmt=fmt, bold=bold)
            _cel(ws, r, 9, '=IFERROR(G%d/E%d-1,"")' % (r, r), fmt="0.0%", bold=bold)
            _cel(ws, r, 10, leitura, fmt=None)
            r += 1

        for bu in u["bus"]:
            ft = fat_ref[(nome, bu, "Tons")]
            fv = fat_ref[(nome, bu, "Valor Liq.Var")]
            ff = fat_ref[(nome, bu, "Valor Liq.Fix")]
            c = col_bu[bu]
            coef = u["coef"][bu]
            lin(bu, "Toneladas", "k t", "=Faturamento!$R$%d" % ft, None, "=Faturamento!$AE$%d" % ft,
                "MP27 vazio ate a tecnica entregar o Volume Movimentado MP27")
            lin(bu, "Area utilizada (soma do ano)", "m2-mes", "=" + area_ref[(nome, bu)]["r09"], None,
                "=" + area_ref[(nome, bu)]["mp"], fmt="#,##0")
            lin(bu, "Tarifa variavel", "R$/t", '=IFERROR(Faturamento!$R$%d/Faturamento!$R$%d,"")' % (fv, ft),
                '=IFERROR(%s*%s/Faturamento!$R$%d*(1+%s),"")' % (soma(21, "R09"), coef, ft, "Tarifa!$D$%d" % T["mk"]),
                "=Tarifa!%s%d" % (c, T["tv"]), "R09 pelo metodo usa o coeficiente de empenho do MP26", bold=True)
            lin(bu, "Tarifa fixa", "R$/m2/mes", '=IFERROR(Faturamento!$R$%d*1000/%s,"")' % (ff, area_ref[(nome, bu)]["r09"]),
                '=IFERROR(%s,"")' % tf_r09_metodo, "=Tarifa!$D$%d" % T["tf"], bold=True)
            lin(bu, "Faturamento liquido variavel", "R$ mil", "=Faturamento!$R$%d" % fv, None, "=Tarifa!%s%d" % (c, T["fv"]))
            lin(bu, "Faturamento liquido fixo", "R$ mil", "=Faturamento!$R$%d" % ff, None, "=Tarifa!%s%d" % (c, T["ff"]))
        # unidade
        lin("Total", "Custo variavel", "R$ mil", "=" + soma(21, "R09"), None, "=Tarifa!$D$%d" % T["cv"],
            "linha 21 das WH (aba Custos)")
        lin("Total", "Custo fixo a recuperar", "R$ mil", "=" + cf_r09, None, "=Tarifa!$D$%d" % T["cfr"],
            "fixo - questoes trabalhistas - Ecolabel")
        lin("Total", "Questoes trabalhistas (fora da tarifa)", "R$ mil", "=" + soma(83, "R09"), None, "=Tarifa!$D$%d" % T["qt"])
        fl_r09 = "+".join("Faturamento!$R$%d+Faturamento!$R$%d" % (fat_ref[(nome, b, "Valor Liq.Var")], fat_ref[(nome, b, "Valor Liq.Fix")])
                          for b in u["bus"])
        lin("Total", "Faturamento liquido (var + fixo)", "R$ mil", "=" + fl_r09, None, "=Tarifa!$D$%d" % T["fl"], bold=True)
        lin("Total", "Margem = faturamento - custo var - fixo a recuperar", "R$ mil",
            "=E%d-E%d-E%d" % (r - 1, r - 4, r - 3), None, "=Tarifa!$D$%d" % T["mg"],
            "R09 praticado positivo = tarifa do R09 acima do custo")
        r += 1
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:J%d" % (r - 1)
    _larguras(ws, [18, 7, 46, 10, 14, 14, 14, 14, 9, 60])


def ligar_faturamento(wb, fat_ref, area_ref, tar_ref, col_mp0):
    """Tarifa e Valor Liq.Fix MP27 da aba Faturamento passam a ler a aba Tarifa (verde = calculado, pode sobrescrever)."""
    ws = wb["Faturamento"]
    col_bu = {"Car": "B", "Moto": "C"}
    dica_col = None
    for c in ws[4]:
        if c.value == "Como preencher":
            dica_col = c.column
    for u in UNIDADES:
        T = tar_ref[u["nome"]]
        for bu in u["bus"]:
            rt = fat_ref[(u["nome"], bu, "Tarifa")]
            rf = fat_ref[(u["nome"], bu, "Valor Liq.Fix")]
            la = area_ref[(u["nome"], bu)]["linha"]
            for m in range(12):
                c = ws.cell(row=rt, column=col_mp0 + m, value='=IFERROR(Tarifa!$%s$%d*1,0)' % (col_bu[bu], T["tv"]))
                c.fill = VERDE
                c = ws.cell(row=rf, column=col_mp0 + m,
                            value='=IFERROR(Tarifa!$D$%d*Area_m2!%s%d/1000,0)' % (T["tf"], get_column_letter(16 + m), la))
                c.fill = VERDE
            if dica_col:
                ws.cell(row=rt, column=dica_col, value="R$/t - CALCULADO na aba Tarifa (verde); sobrescrever so se houver contrato")
                ws.cell(row=rf, column=dica_col, value="R$ mil - CALCULADO: tarifa fixa (aba Tarifa) x m2 do mes (aba Area_m2)")
    ws["A2"] = ("R$ mil e k ton - amarelo = digitar - verde = calculado na aba Tarifa (regra MP26) - "
                "Valor Liq.Var = Tarifa x Tons")


def abrir_excel(tentativas=6, espera=45):
    """DispatchEx com nova tentativa. O Excel as vezes sobe ocupado ('application is busy') enquanto carrega
    suplementos / recuperacao de documento; esperando, a tentativa seguinte abre em ~2 s (medido 08/10/2026).
    NAO encerrar a instancia que ficou presa: cada taskkill vira 'fechamento inesperado' e a proxima abertura
    sobe lenta de novo -- foi esse ciclo que fez 26 tentativas seguidas falharem."""
    import time
    import pythoncom
    import win32com.client

    for n in range(tentativas):
        try:
            xl = win32com.client.DispatchEx("Excel.Application")
            xl.Visible = False
            xl.DisplayAlerts = False
            xl.AskToUpdateLinks = False
            return xl
        except pythoncom.com_error as e:
            print("  Excel ocupado na abertura (%s) - tentativa %d de %d, esperando %ds" % (str(e)[:40], n + 1, tentativas, espera))
            time.sleep(espera)
    raise RuntimeError("Excel nao abriu depois de %d tentativas" % tentativas)


def recalcular(destino):
    """Abre no Excel, recalcula e grava: o Resultado MP27 le os valores ja calculados desta planilha."""
    xl = abrir_excel()
    try:
        wb = xl.Workbooks.Open(destino, UpdateLinks=0)
        xl.CalculateFull()
        erros = {}
        for ws in wb.Worksheets:
            erros[ws.Name] = contar_erros(ws)
        fat = wb.Worksheets("Faturamento")
        total = next((fat.Cells(r, 31).Value for r in range(5, 200) if fat.Cells(r, 4).Value == "Total TLM"), None)
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
    return erros, total


def acrescentar(destino, linha_de, resultado_mp, resultado_r09, col_mp0):
    """linha_de: (unidade, bu, linha) -> numero da linha na aba Faturamento."""
    custos = ler_custos(resultado_mp)
    area_r09 = ler_area_r09(resultado_r09)
    wb = load_workbook(destino)
    custos_ref = aba_custos(wb, custos, os.path.basename(resultado_mp))
    area_ref = aba_area(wb, area_r09)
    bu_ref = aba_bu(wb, linha_de)
    tar_ref = aba_tarifa(wb, linha_de, custos_ref, area_ref, bu_ref)
    bu_tarifas(wb, bu_ref, tar_ref)
    aba_comparativo(wb, linha_de, custos_ref, area_ref, tar_ref)
    ligar_faturamento(wb, linha_de, area_ref, tar_ref, col_mp0)
    wb.save(destino)
    tot = {k: sum(v) for k, v in custos.items()}
    return tot, area_r09
