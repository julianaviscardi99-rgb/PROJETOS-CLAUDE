# -*- coding: utf-8 -*-
"""caderno.py -- grava DataFrame no padrao visual azul do Cockpit.

REGRA (Caetano, 24/08/2026): **nada e gravado sem formatacao**. Vale para caderno de leitura e
tambem para base de conferencia -- nao existe "espelho cru".

O `copc_estilo.py` e a fonte da verdade das cores e da borda; aqui so se acrescenta o que ele nao
faz: escrever um DataFrame inteiro. O `tabela()` de la desenha celula a celula, o que e certo para
caderno de check (dezenas de linhas) e inviavel para base analitica (dezenas de milhares).

Padrao (identico ao do MP_Preparacao):
  · sem linhas de grade
  · A1 titulo em negrito 15 · A2 subtitulo italico 10 cinza
  · tabela a partir da linha 4, cabecalho azul 1F4E79 com texto branco e borda fina BFBFBF
  · painel congelado abaixo do cabecalho, autofiltro, largura de coluna pelo conteudo

Desempenho: os valores vao pelo caminho rapido do pandas e a formatacao e aplicada por cima, por
COLUNA. Escrever celula a celula custava ~20 s por arquivo de 11 mil linhas -- inviavel para os
120 arquivos da base 602.
"""
import os
import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(BASE_DIR, "scripts", "copc"))

from copc_estilo import AZUL, MOEDA, THIN                              # noqa: E402

LINHA_CABECALHO = 4
LIMITE_BORDA = 5000          # acima disso a borda por celula incha o arquivo sem ganho
LARGURA_MIN, LARGURA_MAX = 10, 52
CINZA_TEXTO = "595959"


def _formato_da_coluna(nome, serie):
    nome = str(nome)
    if nome.startswith("%") or (len(nome) > 1 and nome[0] == "P" and nome[1:].isdigit()):
        return "0.0000"                         # percentual de rateio, como a area le
    if pd.api.types.is_float_dtype(serie):
        return MOEDA
    if pd.api.types.is_integer_dtype(serie):
        return "#,##0"
    return None


def _largura(df, coluna):
    amostra = df[coluna].head(200).astype(str)
    maior = max([len(str(coluna))] + [len(v) for v in amostra]) + 2
    return max(LARGURA_MIN, min(LARGURA_MAX, maior))


def formatar(ws, df, titulo, subtitulo):
    """Aplica o padrao numa aba que ja recebeu os dados a partir de LINHA_CABECALHO."""
    ws.sheet_view.showGridLines = False
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=15)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True, size=10, color=CINZA_TEXTO)

    for j in range(1, len(df.columns) + 1):
        c = ws.cell(row=LINHA_CABECALHO, column=j)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN

    primeira, ultima = LINHA_CABECALHO + 1, LINHA_CABECALHO + len(df)
    com_borda = len(df) <= LIMITE_BORDA
    for j, coluna in enumerate(df.columns, 1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = _largura(df, coluna)
        fmt = _formato_da_coluna(coluna, df[coluna])
        if not fmt and not com_borda:
            continue
        for linha in range(primeira, ultima + 1):
            c = ws.cell(row=linha, column=j)
            if fmt:
                c.number_format = fmt
            if com_borda:
                c.border = THIN

    ws.freeze_panes = ws.cell(row=primeira, column=1)
    if len(df):
        ws.auto_filter.ref = "A%d:%s%d" % (LINHA_CABECALHO,
                                           get_column_letter(len(df.columns)), ultima)
    return ws


def escrever_caderno(destino, abas):
    """abas = [(nome_da_aba, titulo, subtitulo, DataFrame), ...]"""
    pasta = os.path.dirname(destino)
    if pasta:
        os.makedirs(pasta, exist_ok=True)
    with pd.ExcelWriter(destino, engine="openpyxl") as w:
        for nome, titulo, subtitulo, df in abas:
            aba = nome[:31]
            df.to_excel(w, sheet_name=aba, index=False, startrow=LINHA_CABECALHO - 1)
            formatar(w.sheets[aba], df, titulo, subtitulo)
            print("     aba %-28s %7d linhas" % (nome, len(df)))
    return destino
