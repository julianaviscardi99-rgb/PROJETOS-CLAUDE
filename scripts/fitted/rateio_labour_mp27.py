"""Monta o Labour Cost FITTED (MP) no padrao da v2: LOCALIDADE + pivots + rateio da Gerencia.

Uso: python rateio_labour_mp27.py <modelo_v2.xlsx> <labour_nova.xlsx> <saida.xlsx>
- modelo_v2: ultima versao ja montada (traz Sheet1 com pivots, % do rateio e checks)
- labour_nova: versao crua recebida (sem LOCALIDADE, Cost Center em texto)
Usa Excel (COM) para manter as pivots reais. Nunca grava sobre os originais.
"""
import os
import sys
import openpyxl
import win32com.client as win32

modelo, nova, saida = (os.path.abspath(p) for p in sys.argv[1:4])
if saida in (modelo, nova):
    sys.exit("saida nao pode sobrescrever um original")

# mapa Cost Center -> LOCALIDADE, tirado do modelo (v2)
ws_m = openpyxl.load_workbook(modelo, read_only=True)["Labour Cost"]
mapa = {}
for r in ws_m.iter_rows(min_row=2, values_only=True):
    if r[2] is not None:
        mapa[str(r[2])] = r[3]

linhas = []
for r in openpyxl.load_workbook(nova, read_only=True)["Labour Cost"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    cc = str(r[2])
    if cc not in mapa:
        sys.exit(f"Cost Center {cc} sem LOCALIDADE no modelo - definir antes de seguir")
    # nova: A..C, D=conta... -> modelo: insere LOCALIDADE depois do Cost Center
    linhas.append([r[0], r[1], int(cc), mapa[cc], *r[3:]])
n = len(linhas)

xl = win32.DispatchEx("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    wb = xl.Workbooks.Open(modelo)
    ws = wb.Worksheets("Labour Cost")
    ws.Range(f"A2:V{max(ws.UsedRange.Rows.Count, n + 1)}").ClearContents()
    ws.Range(f"A2:V{n + 1}").Value = [list(l) for l in linhas]
    for pc in wb.PivotCaches():
        if not pc.SourceData.endswith(f"R{n + 1}C22"):  # intervalo muda so se o no de linhas mudou
            pc.SourceData = f"'Labour Cost'!R1C1:R{n + 1}C22"
        pc.Refresh()
    wb.SaveAs(saida, FileFormat=51)
    s1 = wb.Worksheets("Sheet1")
    print("linhas:", n, "| total MO:", round(s1.Range("N9").Value, 2))
    print("checks rateio (devem ser 0):",
          [round(s1.Range(f"{c}{r}").Value or 0, 6) for r in (40, 50, 51) for c in ("B", "M")])
    wb.Close(False)
finally:
    xl.Quit()
