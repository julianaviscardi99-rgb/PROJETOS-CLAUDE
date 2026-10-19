"""Variacao da MO FITTED: MP27 v2 x R9 FY26 x run-rate R9 (set-dez x3), por unidade e grupo de custo.

Uso: python variacao_mo_r9_runrate_mp27.py <labour_v2.xlsx> <labour_r9.xlsx> <saida.xlsx>
Somente leitura nas fontes. Unidade do R9 vem do centro de custo no MP27 (CC so no R9 = ENC, fora da comparacao).
Grupos (apresentacao): Salarios e PPR (inclui HE e provisoes de 13o/ferias) | Encargos | Assistencia medica | Outros beneficios.
"""
import sys
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

f2, f9, out = sys.argv[1:4]
MS = [f"M{i:02d}" for i in range(1, 13)]
RR = ["M09", "M10", "M11", "M12"]
UN = ["SJP", "IBI", "GOI", "RES", "GER"]

GRUPO = {  # conta -> (grupo apresentacao, grupo detalhe)
    **{c: ("Salarios e PPR", "Salarios + horas extras") for c in ("4200000P", "4200100P", "4201000P", "4201100P")},
    **{c: ("Salarios e PPR", "Provisoes 13o e ferias") for c in ("4200110P", "4200120P", "4201110P", "4201120P")},
    **{c: ("Salarios e PPR", "PPR / PLR") for c in ("4200130P", "4201130P")},
    **{c: ("Encargos", "Encargos") for c in ("4200150P", "4200160P", "4200170P", "4200180P", "4201150P", "4201160P", "4201170P", "4201180P")},
    **{c: ("Assistencia medica", "Assistencia medica") for c in ("4245000P", "4245100P")},
    **{c: ("Outros beneficios", "Outros beneficios") for c in ("4247100P", "4254100P", "4254110P")},
}
G1 = ["Salarios e PPR", "Encargos", "Assistencia medica", "Outros beneficios"]
G2 = ["Salarios + horas extras", "Provisoes 13o e ferias", "PPR / PLR", "Encargos", "Assistencia medica", "Outros beneficios"]


def carregar(path):
    df = pd.read_excel(path, sheet_name="Labour Cost", dtype={"SAP Code": str, "Cost Center": str, "SAP Account": str})
    df = df[df["SAP Code"].notna()].copy()
    cols = [c for c in df.columns if str(c)[:2] == "20"]
    df = df.rename(columns={c: f"M{str(c)[-2:]}" for c in cols})
    df["SAP Account"] = df["SAP Account"].str.strip()
    return df


v2, r9 = carregar(f2), carregar(f9)
cc2u = v2.drop_duplicates("Cost Center").set_index("Cost Center")["LOCALIDADE"].to_dict()
r9["LOCALIDADE"] = r9["Cost Center"].map(cc2u).fillna("ENC")
for d in (v2, r9):
    d["G1"] = d["SAP Account"].map(lambda a: GRUPO[a][0])
    d["G2"] = d["SAP Account"].map(lambda a: GRUPO[a][1])
    assert d["SAP Account"].isin(GRUPO).all(), d.loc[~d["SAP Account"].isin(GRUPO), "SAP Account"].unique()


def valor(d, un, gcol, g, meses, mult=1):
    s = d[(d.LOCALIDADE == un) & (d[gcol] == g)] if un else d[(d.LOCALIDADE.isin(UN)) & (d[gcol] == g)]
    return float(s[meses].sum().sum()) * mult / 1e3


HDR = PatternFill("solid", fgColor="1F3864"); SUB = PatternFill("solid", fgColor="D9E1F2")
NUM = '#,##0;[Red]-#,##0;"-"'; NUM1 = '#,##0.0;[Red]-#,##0.0;"-"'; PCT = '0.0%;[Red]-0.0%;"-"'
BLUE = Font(color="0000FF")


def head(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row, i, t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def aba(ws, titulo, nota, gcol, grupos):
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = nota; ws["A2"].font = Font(italic=True, color="666666")
    head(ws, 4, ["Unidade", "Grupo", "R9 FY26", "R9 run-rate", "MP27 v2", "Δ vs R9 FY", "Δ % vs R9 FY",
                 "Δ vs run-rate", "Δ % vs run-rate", "Mix MP27", "Médio/mês R9 run-rate", "Médio/mês MP27"])
    r = 5; tot_rows = []
    for u in UN + ["TOTAL FITTED"]:
        ini = r
        if u == "TOTAL FITTED":
            for g in grupos:
                refs = [(i, ws.cell(i, 2).value) for i in range(5, r) if ws.cell(i, 2).value == g and ws.cell(i, 1).value in UN]
                ws.cell(r, 1, u); ws.cell(r, 2, g)
                for col in "CDE":
                    ws[f"{col}{r}"] = "=" + "+".join(f"{col}{i}" for i, _ in refs)
                r += 1
        else:
            for g in grupos:
                ws.cell(r, 1, u); ws.cell(r, 2, g)
                for j, (d, mm, mult) in enumerate([(r9, MS, 1), (r9, RR, 3), (v2, MS, 1)]):
                    c = ws.cell(r, 3 + j, round(valor(d, u, gcol, g, mm, mult), 1)); c.font = BLUE
                r += 1
        ws.cell(r, 1, u); ws.cell(r, 2, "Total MO").font = Font(bold=True)
        for col in "CDE":
            ws[f"{col}{r}"] = f"=SUM({col}{ini}:{col}{r-1})"
        for i in range(ini, r + 1):
            ws[f"F{i}"] = f"=E{i}-C{i}"; ws[f"G{i}"] = f'=IF(C{i}=0,"n/a",F{i}/C{i})'
            ws[f"H{i}"] = f"=E{i}-D{i}"; ws[f"I{i}"] = f'=IF(D{i}=0,"n/a",H{i}/D{i})'
            ws[f"J{i}"] = f'=IF(E${r}=0,"n/a",E{i}/E${r})'
            ws[f"K{i}"] = f"=D{i}/12"; ws[f"L{i}"] = f"=E{i}/12"
            for col in "CDEFHKL": ws[f"{col}{i}"].number_format = NUM1 if col in "KL" else NUM
            for col in "GIJ": ws[f"{col}{i}"].number_format = PCT; ws[f"{col}{i}"].alignment = Alignment(horizontal="right")
        for j in range(1, 13):
            ws.cell(r, j).font = Font(bold=True); ws.cell(r, j).fill = SUB
        tot_rows.append(r)
        r += 2
    for i, w in enumerate([14, 24, 11, 12, 11, 12, 13, 13, 14, 10, 17, 15], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "C5"
    return tot_rows


wb = Workbook()
ws = wb.active; ws.title = "Resumo (4 grupos)"
aba(ws, "Variação da MO FITTED - MP27 v2 x R9 FY26 x run-rate R9 (R$ mil)",
    "Run-rate = média set-dez/26 x 12. Custo direto por unidade, sem rateio da Gerência (GER à parte); exclui ENC (PLR jan/26 sem unidade). Azul = valor das fontes; preto = fórmula.",
    "G1", G1)
ws = wb.create_sheet("Detalhe (6 grupos)")
aba(ws, "Variação da MO FITTED - abertura: salário / provisões / PPR separados (R$ mil)",
    "Mesma base da aba anterior; permite escolher a apresentação (ex.: salário+PPR juntos ou separados).", "G2", G2)

# aba de apresentacao: uma tabela por unidade no mesmo formato da tabela-resumo
ws = wb.create_sheet("Tabelas por unidade", 0)
ws["A1"] = "MO FITTED - MP27 v2 x R9 ano x run-rate R9, por unidade (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "Run-rate = média set-dez/26 x 12. Custo direto, sem rateio da Gerência; GER à parte. Δ % vs R9 ano de RES reflete a rampa de 2026 (unidade iniciada no meio do ano)."
ws["A2"].font = Font(italic=True, color="666666")
r = 4
for u in ["TOTAL FITTED"] + UN:
    ws.cell(r, 1, u).font = Font(bold=True, size=12); r += 1
    head(ws, r, ["Grupo", "R9 ano", "Run-rate R9", "MP27 v2", "Δ vs R9 ano", "Δ % vs R9 ano", "Δ vs run-rate", "Δ % vs run-rate"]); r += 1
    ini = r
    for g in G1:
        ws.cell(r, 1, g)
        for j, (d, mm, mult) in enumerate([(r9, MS, 1), (r9, RR, 3), (v2, MS, 1)]):
            x = sum(valor(d, uu, "G1", g, mm, mult) for uu in UN) if u == "TOTAL FITTED" else valor(d, u, "G1", g, mm, mult)
            c = ws.cell(r, 2 + j, round(x, 1)); c.font = BLUE
        r += 1
    ws.cell(r, 1, "Total MO")
    for col in "BCD": ws[f"{col}{r}"] = f"=SUM({col}{ini}:{col}{r-1})"
    for i in range(ini, r + 1):
        ws[f"E{i}"] = f"=D{i}-B{i}"; ws[f"F{i}"] = f'=IF(B{i}=0,"n/a",E{i}/B{i})'
        ws[f"G{i}"] = f"=D{i}-C{i}"; ws[f"H{i}"] = f'=IF(C{i}=0,"n/a",G{i}/C{i})'
        for col in "BCDEG": ws[f"{col}{i}"].number_format = NUM
        for col in "FH": ws[f"{col}{i}"].number_format = PCT; ws[f"{col}{i}"].alignment = Alignment(horizontal="right")
    for j in range(1, 9): ws.cell(r, j).font = Font(bold=True); ws.cell(r, j).fill = SUB
    r += 3
for i, w in enumerate([24, 11, 13, 11, 13, 14, 14, 15], 1): ws.column_dimensions[L(i)].width = w

# aba de contas, para auditoria da classificacao
ws = wb.create_sheet("Contas x grupo")
head(ws, 1, ["Conta SAP", "Descrição", "Grupo (resumo)", "Grupo (detalhe)", "R9 FY26", "R9 run-rate", "MP27 v2"])
desc = pd.concat([v2, r9]).drop_duplicates("SAP Account").set_index("SAP Account")["SAP Account Description"].fillna("")
r = 2
for conta in sorted(GRUPO):
    sel = lambda d: d[(d["SAP Account"] == conta) & d.LOCALIDADE.isin(UN)]
    vals = [sel(r9)[MS].sum().sum() / 1e3, sel(r9)[RR].sum().sum() * 3 / 1e3, sel(v2)[MS].sum().sum() / 1e3]
    for j, x in enumerate([conta, desc.get(conta, ""), GRUPO[conta][0], GRUPO[conta][1]] + [round(float(v), 1) for v in vals], 1):
        c = ws.cell(r, j, x)
        if j >= 5: c.number_format = NUM; c.font = BLUE
    r += 1
ws.cell(r, 2, "TOTAL").font = Font(bold=True)
for col in "EFG": ws[f"{col}{r}"] = f"=SUM({col}2:{col}{r-1})"; ws[f"{col}{r}"].number_format = NUM
for i, w in enumerate([12, 34, 20, 26, 11, 12, 11], 1): ws.column_dimensions[L(i)].width = w
wb.save(out)
print("ok", out)
