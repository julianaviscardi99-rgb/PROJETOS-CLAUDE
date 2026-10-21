"""PPT MP27 Fitted v7 - ajuste de 21/10 (pedido da Juliana), em cima da v7 que ela salvou (10:13):

- slide 2: bridge ganha a coluna Sorocaba (-0,2 = venda de materiais de consumo; R09 Brasil = 13,0), tabela ganha a linha
  Sorocaba, textos de 4 unidades -> Brasil (EBIT -1,5 | -11,6%);
- slide 3 (leitura economica) removido;
- slide 6 (resumo executivo): linha Sorocaba na tabela;
- slide 12 (mao de obra) segue oculto (ja estava).

Uso: python ajustar_ppt_mp27_v7_sorocaba.py <v7_ela.pptx> <basebook.xlsx> <saida.pptx>
"""
import copy
import sys

import openpyxl
from pptx import Presentation
from pptx.util import Emu

DECK, BB, OUT = sys.argv[1:4]
MINUS = '−'

bb = openpyxl.load_workbook(BB, data_only=True)
ws = bb['Month_FY_LC - MP27']
BR9, BR27 = float(ws.cell(55, 42).value), float(ws.cell(55, 43).value)
U = {}
for u, sh in (('IBI', 'IBIRITE'), ('SJP', 'SJP'), ('GOI', 'GOIANA'), ('RES', 'RESENDE')):
    w = bb[sh]
    U[u] = (float(w.cell(55, 42).value), float(w.cell(55, 43).value))
SOR9 = BR9 - sum(v[0] for v in U.values())
SOR27 = BR27 - sum(v[1] for v in U.values())
assert abs(SOR27) < 0.002 and abs(SOR9 - 0.188) < 0.01, (SOR9, SOR27)


def num(x, d=1):
    return f"{abs(x):,.{d}f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def sg(x, d=1):
    return num(0, d) if round(abs(x), d) == 0 else (MINUS if x < 0 else '+') + num(x, d)


def par(x):
    return f"({num(x)})" if round(abs(x), 1) != 0 and x < 0 else num(x)


prs = Presentation(DECK)
s2 = prs.slides[1]
assert 'Bridge consolidado' in s2.shapes[0].text_frame.text
by = {x.name: x for x in s2.shapes}
tree = s2.shapes._spTree
next_id = [max(x.shape_id for x in s2.shapes) + 1]


def clone(sh):
    el = copy.deepcopy(sh._element)
    el.xpath('.//p:cNvPr')[0].set('id', str(next_id[0]))
    next_id[0] += 1
    tree.append(el)
    return s2.shapes[-1]


def settext(sh, new):
    p = sh.text_frame.paragraphs[0]
    runs = p.runs
    runs[0].text = new
    for r in runs[1:]:
        r.text = ''


def sub(sh, old, new):
    for p in sh.text_frame.paragraphs:
        full = ''.join(r.text for r in p.runs)
        if old in full:
            rs = p.runs
            rs[0].text = full.replace(old, new)
            for r in rs[1:]:
                r.text = ''
            return
    raise KeyError(old)


# ------------------------------------------------------------ bridge: 5 deltas
deltas = [('IBI', U['IBI'][1] - U['IBI'][0]), ('SJP', U['SJP'][1] - U['SJP'][0]), ('GOI', U['GOI'][1] - U['GOI'][0]),
          ('RES', U['RES'][1] - U['RES'][0]), ('SOR', SOR27 - SOR9)]
SC, YB = 195000, 5093208
y = lambda v: int(YB - v * SC)
sor_bar, sor_lab, sor_cat, sor_con = clone(by['Shape 27']), clone(by['Text 29']), clone(by['Text 43']), clone(by['Shape 34'])
bars = {'IBI': by['Shape 24'], 'SJP': by['Shape 27'], 'GOI': by['Shape 30'], 'RES': by['Shape 33'], 'SOR': sor_bar}
labs = {'IBI': by['Text 26'], 'SJP': by['Text 29'], 'GOI': by['Text 32'], 'RES': by['Text 35'], 'SOR': sor_lab}
cats = {'IBI': by['Text 40'], 'SJP': by['Text 41'], 'GOI': by['Text 42'], 'RES': by['Text 43'], 'SOR': sor_cat}
cons = [by['Shape 25'], by['Shape 28'], by['Shape 31'], by['Shape 34'], sor_con, by['Shape 36']]   # R09>IBI ... SOR>final
CX0, PITCH, BW = 1260000, 930000, 600000
cx = lambda i: CX0 + i * PITCH
W0, WF = 700000, 800000          # largura das barras de total

r9b = by['Shape 22']
r9b.left, r9b.width = Emu(cx(0) - W0 // 2), Emu(W0)
r9b.top, r9b.height = Emu(y(BR9)), Emu(YB - y(BR9))
lab0 = by['Text 23']
lab0.left, lab0.top = Emu(cx(0) - lab0.width // 2), Emu(y(BR9) - 256032)
cat0 = by['Text 39']
cat0.left = Emu(cx(0) - cat0.width // 2)
settext(cat0, 'R09')
cat0.text_frame.paragraphs[1].runs[0].text = 'Brasil' if len(cat0.text_frame.paragraphs) > 1 else ''
lvl = BR9
levels = [BR9]
for i, (u, dv) in enumerate(deltas, start=1):
    new = lvl + dv
    lo, hi = min(lvl, new), max(lvl, new)
    b = bars[u]
    b.left, b.width = Emu(cx(i) - BW // 2), Emu(BW)
    b.top, b.height = Emu(y(hi)), Emu(max(y(lo) - y(hi), 12700))
    L = labs[u]
    L.left, L.top = Emu(cx(i) - L.width // 2), Emu(y(hi) - 256032)
    if u == 'SOR':
        settext(L, sg(dv))
    c = cats[u]
    c.left = Emu(cx(i) - c.width // 2)
    if u == 'SOR':
        c.top = by['Text 43'].top
        settext(c, 'Sorocaba')
    levels.append(new)
    lvl = new
final = by['Shape 37']
final.left, final.width = Emu(cx(6) - WF // 2), Emu(WF)
final.top, final.height = Emu(y(lvl)), Emu(YB - y(lvl))
lf = by['Text 38']
lf.left, lf.top = Emu(cx(6) - lf.width // 2), Emu(y(lvl) - 256032)
cf = by['Text 44']
cf.left = Emu(cx(6) - cf.width // 2)
widths = [W0] + [BW] * 5 + [WF]
for k, cn in enumerate(cons):
    x1 = cx(k) + widths[k] // 2
    x2 = cx(k + 1) - widths[k + 1] // 2
    cn.left, cn.width, cn.top, cn.height = Emu(x1), Emu(x2 - x1), Emu(y(levels[k])), Emu(1)
assert abs(lvl - BR27) < 0.01

# ------------------------------------------------------------ textos do slide 2
sub(by['Text 1'], '4 unidades: Resende, Ibirité, Goiana e São José dos Pinhais', 'Brasil: Ibirité, SJP, Goiana, Resende e Sorocaba (encerrada)')
sub(by['Text 3'], 'EBIT cai R$ 1,3 mi no perímetro das 4 unidades', f"EBIT cai R$ {num(BR9 - BR27)} mi vs R09")
sub(by['Text 13'], 'EBIT | 4 unidades', 'EBIT | Brasil')
sub(by['Text 15'], by['Text 15'].text_frame.text, f"−R$ {num(BR9 - BR27)} mi | −{num((BR9 - BR27) / BR9 * 100)}%")
for n in ('Text 46', 'Text 47'):
    sub(by[n], '−R$ 1,3 mi', f"−R$ {num(BR9 - BR27)} mi")
sub(by['Text 83'], by['Text 83'].text_frame.text,
    f"Reconciliação: R09 Brasil R$ {num(BR9)} mi = {num(sum(v[0] for v in U.values()))} das 4 unidades + {num(SOR9)} de Sorocaba "
    "(venda de materiais de consumo, sem EBIT no MP’27).")

# ------------------------------------------------------------ tabela do slide 2: 5 linhas
P, TOP = 420000, 2990088
rows = [('Shape 54', 'Text 55', 'Text 56', 'Text 57', 'Text 58', 'Text 59'),
        ('Shape 60', 'Text 61', 'Text 62', 'Text 63', 'Text 64', 'Text 65'),
        ('Shape 66', 'Text 67', 'Text 68', 'Text 69', 'Text 70', 'Text 71'),
        ('Shape 72', 'Text 73', 'Text 74', 'Text 75', 'Text 76', 'Text 77')]
sor_row = [clone(by['Shape 66'])] + [clone(by[n]) for n in rows[1][1:]]     # fundo alternado + textos da linha SJP (delta vermelho)
rows.append(tuple(x.name for x in sor_row))
objs = [[by[n] for n in r] for r in rows[:4]] + [sor_row]
for k, r in enumerate(objs):
    top = TOP + k * P
    r[0].top, r[0].height = Emu(top), Emu(P)
    for t in r[1:5]:
        t.top = Emu(top + 123000)
    r[5].top, r[5].height = Emu(top + 20000), Emu(P - 40000)
for t, v in zip(sor_row[1:], ['Sorocaba', num(SOR9), num(SOR27), sg(SOR27 - SOR9), 'Venda de materiais de consumo']):
    settext(t, v)
box = by['Shape 78']
box.top, box.height = Emu(TOP + 5 * P + 40000), Emu(380000)
t79 = by['Text 79']
t79.top = Emu(box.top + (380000 - t79.height) // 2)
sub(t79, '84%', f"{round((U['IBI'][1] + U['SJP'][1]) / BR27 * 100)}%")

# ------------------------------------------------------------ slide 6: linha Sorocaba
s6 = prs.slides[5]
assert 'Resumo executivo' in s6.shapes.title.text_frame.text
gf = [x for x in s6.shapes if x.name == 'Table 74'][0]
tbl = gf.table._tbl
trs = tbl.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tr')
assert len(trs) == 6
new_tr = copy.deepcopy(trs[4])          # copia da linha RES
trs[5].addprevious(new_tr)
trs = tbl.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tr')
for r_ in trs:
    r_.set('h', '290000')
trs[0].set('h', '330000')
gf.height = Emu(330000 + 6 * 290000)
t6 = gf.table


def cset(r, c, new):
    p = t6.cell(r, c).text_frame.paragraphs[0]
    rs = p.runs
    if c == 0:
        rs[-1].text = new
        return
    rs[0].text = new
    for x in rs[1:]:
        x.text = ''


cset(5, 0, 'Sorocaba')
for c, v in enumerate(['', '–', '–', num(SOR9), num(SOR27), par(SOR27 - SOR9), '–', '–', 'Venda de materiais de consumo']):
    if c:
        cset(5, c, v)
cset(6, 8, 'R09 = 4 unidades 12,8 + Sorocaba 0,2')
for sh in s6.shapes:
    if sh.name == 'Nota de rodape':
        sh.top = Emu(4745000)

# ------------------------------------------------------------ slide 3 (leitura economica) removido
lst = prs.slides._sldIdLst
sid = list(lst)[2]
assert any('Leitura econ' in x.text_frame.text for x in prs.slides[2].shapes if x.has_text_frame)
prs.part.drop_rel(sid.rId)
lst.remove(sid)

prs.save(OUT)
print('salvo', OUT, len(prs.slides), 'slides; ocultos:', [i + 1 for i, s in enumerate(prs.slides) if s._element.get('show') == '0'])
