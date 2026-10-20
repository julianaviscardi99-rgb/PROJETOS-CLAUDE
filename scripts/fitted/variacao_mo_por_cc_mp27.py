"""Variacao da MO FITTED por centro de custo (CC) e categoria, MP27 v2 x R9 FY26 x run-rate R9, com foco em SJP.

Uso: python variacao_mo_por_cc_mp27.py <labour_v2.xlsx> <labour_r9.xlsx> <saida.xlsx>
Somente leitura nas fontes. R$ mil. Mesmos grupos de variacao_mo_r9_runrate_mp27.py, com PPR/PLR separado dos salarios.
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

f2, f9, out = sys.argv[1:4]
MS = [f"M{i:02d}" for i in range(1, 13)]
RR = ["M09", "M10", "M11", "M12"]
UN = ["SJP", "IBI", "GOI", "RES", "GER"]
GRUPO = {}
for c in ("4200000P", "4200100P", "4201000P", "4201100P", "4200110P", "4200120P", "4201110P", "4201120P"):
    GRUPO[c] = ("Salarios e PPR", "Salarios, HE e provisoes 13o/ferias")
for c in ("4200130P", "4201130P"):
    GRUPO[c] = ("Salarios e PPR", "PPR / PLR")
for c in ("4200150P", "4200160P", "4200170P", "4200180P", "4201150P", "4201160P", "4201170P", "4201180P"):
    GRUPO[c] = ("Encargos", "Encargos")
for c in ("4245000P", "4245100P"):
    GRUPO[c] = ("Assistencia medica", "Assistencia medica")
for c in ("4247100P", "4254100P", "4254110P"):
    GRUPO[c] = ("Outros beneficios", "Outros beneficios")
G1 = ["Salarios e PPR", "Encargos", "Assistencia medica", "Outros beneficios"]
G2 = ["Salarios, HE e provisoes 13o/ferias", "PPR / PLR", "Encargos", "Assistencia medica", "Outros beneficios"]


def carregar(path):
    df = pd.read_excel(path, sheet_name="Labour Cost", dtype={"SAP Code": str, "Cost Center": str, "SAP Account": str})
    df = df[df["SAP Code"].notna()].copy()
    cols = [c for c in df.columns if str(c)[:2] == "20"]
    df = df.rename(columns={c: f"M{str(c)[-2:]}" for c in cols})
    df["SAP Account"] = df["SAP Account"].str.strip()
    assert df["SAP Account"].isin(GRUPO).all(), df.loc[~df["SAP Account"].isin(GRUPO), "SAP Account"].unique()
    df["G1"] = df["SAP Account"].map(lambda a: GRUPO[a][0])
    df["G2"] = df["SAP Account"].map(lambda a: GRUPO[a][1])
    return df


v2, r9 = carregar(f2), carregar(f9)
cc2u = v2.drop_duplicates("Cost Center").set_index("Cost Center")["LOCALIDADE"].to_dict()
r9["LOCALIDADE"] = r9["Cost Center"].map(cc2u).fillna("ENC")
FONTES = [(r9, MS, 1), (r9, RR, 3), (v2, MS, 1)]

HDR = PatternFill("solid", fgColor="1F3864"); SUB = PatternFill("solid", fgColor="D9E1F2")
NUM = '#,##0;[Red]-#,##0;"-"'; PCT = '0.0%;[Red]-0.0%;"-"'; BLUE = Font(color="0000FF")
COLS = ["R9 ano", "Run-rate R9", "MP27 v2", "Δ vs R9 ano", "Δ % vs R9 ano", "Δ vs run-rate", "Δ % vs run-rate",
        "Efeito base R9 (run-rate - R9 ano)"]


def vals(mask_fn, gcol=None, g=None):
    r = []
    for d, mm, mult in FONTES:
        s = d[mask_fn(d)]
        if gcol:
            s = s[s[gcol] == g]
        r.append(round(float(s[mm].sum().sum()) * mult / 1e3, 1))
    return r


def head(ws, row, first):
    for i, t in enumerate(first + COLS, 1):
        c = ws.cell(row, i, t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def formulas(ws, i, c0):
    a, b, c = (f"{L(c0 + k)}{i}" for k in range(3))
    ws[f"{L(c0 + 3)}{i}"] = f"={c}-{a}"
    ws[f"{L(c0 + 4)}{i}"] = f'=IF({a}=0,"n/a",{L(c0 + 3)}{i}/{a})'
    ws[f"{L(c0 + 5)}{i}"] = f"={c}-{b}"
    ws[f"{L(c0 + 6)}{i}"] = f'=IF({b}=0,"n/a",{L(c0 + 5)}{i}/{b})'
    ws[f"{L(c0 + 7)}{i}"] = f"={b}-{a}"
    for k in (0, 1, 2, 3, 5, 7):
        ws[f"{L(c0 + k)}{i}"].number_format = NUM
    for k in (4, 6):
        ws[f"{L(c0 + k)}{i}"].number_format = PCT
        ws[f"{L(c0 + k)}{i}"].alignment = Alignment(horizontal="right")


def bloco(ws, r, titulo, linhas, grupos, gcol):
    """linhas: [(rotulo, mask_fn)]; para cada linha, uma sub-tabela por grupo + total."""
    ws.cell(r, 1, titulo).font = Font(bold=True, size=12); r += 1
    head(ws, r, ["Linha", "Grupo"]); r += 1
    for rot, fn in linhas:
        ini = r
        for g in grupos:
            ws.cell(r, 1, rot); ws.cell(r, 2, g)
            for j, x in enumerate(vals(fn, gcol, g)):
                c = ws.cell(r, 3 + j, x); c.font = BLUE
            formulas(ws, r, 3); r += 1
        ws.cell(r, 1, rot); ws.cell(r, 2, "Total MO")
        for k in range(3):
            ws[f"{L(3 + k)}{r}"] = f"=SUM({L(3 + k)}{ini}:{L(3 + k)}{r - 1})"
        formulas(ws, r, 3)
        for j in range(1, 11):
            ws.cell(r, j).font = Font(bold=True); ws.cell(r, j).fill = SUB
        r += 2
    return r


def larg(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[L(i)].width = x


def cabecalho(ws, titulo):
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = NOTA; ws["A2"].font = Font(italic=True, color="666666")


NOTA = ("Run-rate = média set-dez/26 x 12. 'Efeito base R9' = quanto o R9 do ano está abaixo do seu próprio run-rate (jan-fev/26 com poucos horistas). "
        "CC = centro de custo; Categoria = Horista/Mensalista. Custo direto, sem rateio GER. Azul = fonte; preto = fórmula. R$ mil.")
sjp = lambda d: d.LOCALIDADE == "SJP"
CCS = sorted(v2[v2.LOCALIDADE == "SJP"]["Cost Center"].unique())

wb = Workbook()
ws = wb.active; ws.title = "SJP por CC"
cabecalho(ws, "SJP - MO por centro de custo e grupo: MP27 v2 x R9 ano x run-rate (R$ mil)")
linhas = [("SJP total", sjp)] + [(f"CC {cc}", (lambda cc: lambda d: sjp(d) & (d["Cost Center"] == cc))(cc)) for cc in CCS]
bloco(ws, 4, "SJP por CC (salários / PPR / encargos / assistência médica)", linhas, G2, "G2")
larg(ws, [16, 34, 10, 12, 10, 12, 12, 12, 13, 18]); ws.freeze_panes = "C4"

ws = wb.create_sheet("SJP por categoria")
cabecalho(ws, "SJP - MO por categoria (Horista x Mensalista) e grupo (R$ mil)")
cats = [(c.capitalize(), (lambda c: lambda d: sjp(d) & (d["Category"] == c))(c)) for c in ("HORISTA", "MENSALISTA")]
bloco(ws, 4, "SJP por categoria", cats, G2, "G2")
larg(ws, [16, 34, 10, 12, 10, 12, 12, 12, 13, 18])

ws = wb.create_sheet("Todas as unidades por CC")
cabecalho(ws, "MO FITTED por unidade e centro de custo (R$ mil)")
lin = []
for u in UN:
    for cc in sorted(v2[v2.LOCALIDADE == u]["Cost Center"].unique()):
        cat = "/".join(sorted(v2[(v2.LOCALIDADE == u) & (v2["Cost Center"] == cc)]["Category"].str.capitalize().unique()))
        lin.append((f"{u} - CC {cc} ({cat})", (lambda u, cc: lambda d: (d.LOCALIDADE == u) & (d["Cost Center"] == cc))(u, cc)))
bloco(ws, 4, "Por CC", lin, G1, "G1")
larg(ws, [34, 22, 10, 12, 10, 12, 12, 12, 13, 18]); ws.freeze_panes = "C4"

ws = wb.create_sheet("SJP por conta")
cabecalho(ws, "SJP - MO por CC, categoria e conta SAP (R$ mil)")
head(ws, 4, ["CC / Categoria", "Conta - descrição"])
base = pd.concat([v2[sjp(v2)], r9[sjp(r9)]]).drop_duplicates(["Cost Center", "Category", "SAP Account"])
base = base.sort_values(["Cost Center", "Category", "SAP Account"])
r = 5
for _, b in base.iterrows():
    fn = lambda d, b=b: sjp(d) & (d["Cost Center"] == b["Cost Center"]) & (d["Category"] == b["Category"]) & (d["SAP Account"] == b["SAP Account"])
    ws.cell(r, 1, f'{b["Cost Center"]} {b["Category"].capitalize()}')
    ws.cell(r, 2, f'{b["SAP Account"]} {str(b["SAP Account Description"]).strip()}')
    for j, x in enumerate(vals(fn)):
        c = ws.cell(r, 3 + j, x); c.font = BLUE
    formulas(ws, r, 3); r += 1
ws.cell(r, 2, "TOTAL SJP")
for k in range(3):
    ws[f"{L(3 + k)}{r}"] = f"=SUM({L(3 + k)}5:{L(3 + k)}{r - 1})"
formulas(ws, r, 3)
for j in range(1, 11):
    ws.cell(r, j).font = Font(bold=True); ws.cell(r, j).fill = SUB
larg(ws, [20, 42, 10, 12, 10, 12, 12, 12, 13, 18]); ws.freeze_panes = "C5"

ws = wb.create_sheet("SJP mensal")
ws["A1"] = "SJP - MO por mês e centro de custo, R9 x MP27 v2 (R$ mil)"; ws["A1"].font = Font(bold=True, size=13)
for i, t in enumerate(["CC", "Versão"] + [f"{m:02d}" for m in range(1, 13)] + ["Total"], 1):
    c = ws.cell(3, i, t); c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR; c.alignment = Alignment(horizontal="center")
r = 4
for cc in ["Total SJP"] + CCS:
    for nm, d in (("R9", r9), ("MP27 v2", v2)):
        s = d[sjp(d)] if cc == "Total SJP" else d[sjp(d) & (d["Cost Center"] == cc)]
        ws.cell(r, 1, cc); ws.cell(r, 2, nm)
        for j, m in enumerate(MS):
            c = ws.cell(r, 3 + j, round(float(s[m].sum()) / 1e3, 1)); c.font = BLUE; c.number_format = NUM
        ws.cell(r, 15, f"=SUM(C{r}:N{r})").number_format = NUM
        r += 1
    ws.cell(r, 1, cc); ws.cell(r, 2, "Δ MP27 - R9").font = Font(bold=True)
    for j in range(13):
        col = L(3 + j)
        ws[f"{col}{r}"] = f"={col}{r - 1}-{col}{r - 2}"; ws[f"{col}{r}"].number_format = NUM; ws[f"{col}{r}"].font = Font(bold=True)
    r += 2
larg(ws, [12, 14] + [8] * 12 + [9])
wb.save(out)
print("ok", out)
