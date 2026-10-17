"""Delta Labour Cost FITTED MP27: v1 vs v2, por conta contabil e centro de montagem (LOCALIDADE).

Uso: python delta_labour_v1_v2.py <v1.xlsx> <v2.xlsx> <saida.xlsx>
"""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

f1, f2, out = sys.argv[1:4]
KEY = ["SAP Code", "Category", "Cost Center", "LOCALIDADE", "SAP Account", "SAP Account Description", "Cost Center Variability"]


def load(p):
    df = pd.read_excel(p, sheet_name="Labour Cost", dtype={"SAP Code": str})
    return df.set_index(KEY)


a, b = load(f1), load(f2)
M = [c for c in a.columns if str(c).startswith("2027")]
a, b = a[M], b[M]
idx = a.index.union(b.index)
a, b = a.reindex(idx, fill_value=0), b.reindex(idx, fill_value=0)

HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E1F2")
FMT = '#,##0;[Red]-#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'


def head(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row, i, t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def grid(ws, start, df1, df2, keys):
    """Escreve blocos Total: V1 | V2 | Delta | Delta % (delta/% como formula)."""
    n = len(keys)
    head(ws, start, keys + ["Total V1", "Total V2", "Delta (V2-V1)", "Delta %"])
    r = start + 1
    for k in df1.index:
        vals = list(k) if isinstance(k, tuple) else [k]
        for i, v in enumerate(vals, 1):
            ws.cell(r, i, v)
        ws.cell(r, n + 1, float(df1.loc[k].sum())).number_format = FMT
        ws.cell(r, n + 2, float(df2.loc[k].sum())).number_format = FMT
        ws.cell(r, n + 3, f"={L(n+2)}{r}-{L(n+1)}{r}").number_format = FMT
        ws.cell(r, n + 4, f'=IF({L(n+1)}{r}=0,"",{L(n+3)}{r}/{L(n+1)}{r})').number_format = PCT
        r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    for c in range(1, n + 5):
        ws.cell(r, c).fill = SUB
    for c in (n + 1, n + 2, n + 3):
        x = ws.cell(r, c, f"=SUM({L(c)}{start+1}:{L(c)}{r-1})"); x.number_format = FMT; x.font = Font(bold=True)
    x = ws.cell(r, n + 4, f'=IF({L(n+1)}{r}=0,"",{L(n+3)}{r}/{L(n+1)}{r})'); x.number_format = PCT; x.font = Font(bold=True)
    for i in range(1, n + 5):
        ws.column_dimensions[L(i)].width = 14 if i > n else (34 if i == 2 and n > 2 else 16)
    ws.freeze_panes = ws.cell(start + 1, n + 1)
    return r


def agg(df, by):
    return df.groupby(by).sum()


wb = Workbook()

# 1. Resumo por centro de montagem
ws = wb.active; ws.title = "Resumo por Centro"
ws["A1"] = "Labour Cost FITTED MP27 - Delta V2 vs V1 (BRL)"; ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "V1 = Labour Cost FITTED - MP27 (15/10/2026) | V2 = Labour Cost FITTED - MP27_v2 (17/10/2026)"
grid(ws, 4, agg(a, ["LOCALIDADE"]), agg(b, ["LOCALIDADE"]), ["Centro de Montagem"])
r0 = 12
ws.cell(r0 - 1, 1, "Por centro e variabilidade").font = Font(bold=True)
grid(ws, r0, agg(a, ["LOCALIDADE", "Cost Center Variability"]), agg(b, ["LOCALIDADE", "Cost Center Variability"]), ["Centro de Montagem", "Variabilidade"])
ws.freeze_panes = None

# 2. Por conta contabil
d1 = agg(a, ["SAP Account", "SAP Account Description"]); d2 = agg(b, ["SAP Account", "SAP Account Description"])
ws = wb.create_sheet("Delta por Conta")
grid(ws, 1, d1, d2, ["SAP Account", "Descricao da Conta"])

# 3. Conta x Centro (principal)
c1 = agg(a, ["LOCALIDADE", "SAP Account", "SAP Account Description"]); c2 = agg(b, ["LOCALIDADE", "SAP Account", "SAP Account Description"])
ws = wb.create_sheet("Delta Conta x Centro")
grid(ws, 1, c1, c2, ["Centro de Montagem", "SAP Account", "Descricao da Conta"])

# 4. Mensal conta x centro (delta)
dm = (agg(b, ["LOCALIDADE", "SAP Account", "SAP Account Description"]) - agg(a, ["LOCALIDADE", "SAP Account", "SAP Account Description"]))
ws = wb.create_sheet("Delta Mensal Conta x Centro")
head(ws, 1, ["Centro de Montagem", "SAP Account", "Descricao da Conta"] + [str(m) for m in M] + ["Total"])
r = 2
for k, row in dm.iterrows():
    for i, v in enumerate(k, 1):
        ws.cell(r, i, v)
    for j, m in enumerate(M):
        ws.cell(r, 4 + j, float(row[m])).number_format = FMT
    ws.cell(r, 4 + len(M), f"=SUM(D{r}:{L(3+len(M))}{r})").number_format = FMT
    r += 1
ws.cell(r, 1, "TOTAL").font = Font(bold=True)
for j in range(len(M) + 1):
    c = 4 + j
    x = ws.cell(r, c, f"=SUM({L(c)}2:{L(c)}{r-1})"); x.number_format = FMT; x.font = Font(bold=True)
for i in range(1, 4 + len(M) + 1):
    ws.column_dimensions[L(i)].width = 14 if i > 3 else (34 if i == 3 else 16)
ws.freeze_panes = "D2"

# 5. Detalhe linha a linha (centro de custo)
ws = wb.create_sheet("Detalhe Linha a Linha")
tot1, tot2 = a.sum(axis=1), b.sum(axis=1)
dt = pd.DataFrame({"v1": tot1, "v2": tot2}); dt["d"] = dt.v2 - dt.v1
dt = dt[dt.d.abs() > 0.005].sort_values("d", key=abs, ascending=False)
head(ws, 1, ["SAP Code", "Category", "Cost Center", "Centro de Montagem", "SAP Account", "Descricao da Conta", "Variabilidade", "Total V1", "Total V2", "Delta", "Delta %"])
for r, (k, row) in enumerate(dt.iterrows(), 2):
    for i, v in enumerate(k, 1):
        ws.cell(r, i, v)
    ws.cell(r, 8, float(row.v1)).number_format = FMT
    ws.cell(r, 9, float(row.v2)).number_format = FMT
    ws.cell(r, 10, f"=I{r}-H{r}").number_format = FMT
    ws.cell(r, 11, f'=IF(H{r}=0,"",J{r}/H{r})').number_format = PCT
for i in range(1, 12):
    ws.column_dimensions[L(i)].width = 34 if i == 6 else 14
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:K{len(dt)+1}"

wb.save(out)
print("linhas com delta:", len(dt), "de", len(tot1))
print("Total V1", round(tot1.sum(), 2), "V2", round(tot2.sum(), 2), "Delta", round(tot2.sum() - tot1.sum(), 2))
print(agg(b, ["LOCALIDADE"]).sum(axis=1).sub(agg(a, ["LOCALIDADE"]).sum(axis=1)).round(0).to_string())
