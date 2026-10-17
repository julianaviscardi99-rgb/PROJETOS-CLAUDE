"""Rentabilidade por cliente - Ibirite (Fiat / Iveco / CNH) a partir da MENS FITTED MP27.

Faturamento e material de consumo = diretos por cliente (linhas ja abertas na aba IBI).
Demais custos = rateados por driver (Volume, Faturamento ou Fiat 100%), escolhido na aba Premissas.

Uso: python rentabilidade_cliente_ibi.py <MENS FITTED.xlsx> <saida.xlsx>
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

src, out = sys.argv[1:3]
ws = openpyxl.load_workbook(src, data_only=True)["IBI"]
MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
CLI = ["Fiat", "Iveco", "CNH"]


def linha(r, k=1.0):
    return [(ws.cell(r, 5 + i).value or 0) * k for i in range(12)]


# linhas da aba IBI (custos em R$ mil; volume em pecas; faturamento em R$)
VOL = {"Fiat": 53, "Iveco": 54, "CNH": 55}
FAT = {"Fiat": 65, "Iveco": 66, "CNH": 67}
MAT = {"Fiat": 75, "Iveco": 76, "CNH": 77}
CUSTOS = [  # (grupo, nome, linha, driver padrao)
    ("Variable Cost", "Labour", 19, "Volume"),
    ("Variable Cost", "Handling", 20, "Volume"),
    ("Variable Cost", "Transportation", 22, "Fiat 100%"),
    ("Variable Cost", "Other Variable", 23, "Volume"),
    ("Fixed Cost", "Labour", 32, "Volume"),
    ("Fixed Cost", "Depreciation", 33, "Volume"),
    ("Fixed Cost", "IFRS16", 34, "Volume"),
    ("Fixed Cost", "Rents", 35, "Volume"),
    ("Fixed Cost", "Condominio", 36, "Volume"),
    ("Fixed Cost", "Other Fixed", 37, "Volume"),
]
VARS = [c for c in CUSTOS if c[0] == "Variable Cost"]
FIXS = [c for c in CUSTOS if c[0] == "Fixed Cost"]

HDR = PatternFill("solid", fgColor="1F3864")
INP = PatternFill("solid", fgColor="FFF2CC")
N0 = '#,##0.0;[Red]-#,##0.0;"-"'
PC = '0.0%;[Red]-0.0%;"-"'
B = Font(bold=True)


def head(s, r, vals, c0=1):
    for i, t in enumerate(vals):
        c = s.cell(r, c0 + i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)


wb = openpyxl.Workbook()

# ---- Premissas
pr = wb.active
pr.title = "Premissas"
pr["A1"] = "Rentabilidade por cliente - IBIRITE - MP27 (R$ mil)"
pr["A1"].font = Font(bold=True, size=13)
pr["A2"] = ("Fonte: MENS FITTED MP27_v3_teste daniel.xlsx, aba IBI. Faturamento e Material de consumo = diretos por cliente. "
            "Demais custos = rateio pelo driver abaixo (celulas amarelas editaveis: Volume, Faturamento ou Fiat 100%). Frete (Transportation) = 100% Fiat, confirmado pela Juliana.")
head(pr, 4, ["Grupo", "Linha de custo", "Driver de rateio"])
dv = DataValidation(type="list", formula1='"Volume,Faturamento,Fiat 100%"', allow_blank=False)
pr.add_data_validation(dv)
PR_ROW = {}
for i, (g, n, r, d) in enumerate(CUSTOS):
    rr = 5 + i
    pr.cell(rr, 1, g)
    pr.cell(rr, 2, n)
    c = pr.cell(rr, 3, d)
    c.fill = INP
    dv.add(c)
    PR_ROW[r] = rr
pr.cell(5 + len(CUSTOS), 1, "Variable Cost")
pr.cell(5 + len(CUSTOS), 2, "Direct Materials")
pr.cell(5 + len(CUSTOS), 3, "Direto (material de consumo de cada cliente)")
for col, w in zip("ABC", (16, 22, 44)):
    pr.column_dimensions[col].width = w

# ---- Dados IBI (valores do arquivo)
dd = wb.create_sheet("Dados IBI")
head(dd, 1, ["Item", "Cliente"] + MESES)
R = {}
state = {"r": 2}


def put(key, item, cli, vals):
    r = state["r"]
    dd.cell(r, 1, item)
    dd.cell(r, 2, cli)
    for i, v in enumerate(vals):
        dd.cell(r, 3 + i, v).number_format = N0
    R[key] = r
    state["r"] += 1


def soma3(key):
    return [f"=SUM({L(3+i)}{R[(key,'Fiat')]}:{L(3+i)}{R[(key,'CNH')]})" for i in range(12)]


for c in CLI:
    put(("vol", c), "Volume (pecas)", c, linha(VOL[c]))
put(("vol", "T"), "Volume (pecas)", "Total", soma3("vol"))
for c in CLI:
    put(("fat", c), "Faturamento (R$ mil)", c, linha(FAT[c], 1 / 1000))
put(("fat", "T"), "Faturamento (R$ mil)", "Total", soma3("fat"))
for c in CLI:
    put(("mat", c), "Material de consumo (R$ mil)", c, linha(MAT[c], 1 / 1000))
for g, n, row, d in CUSTOS:
    put(("c", row), f"{g} - {n} (R$ mil)", "Total IBI", linha(row))
put("ebit", "EBIT total IBI no arquivo (R$ mil)", "Total IBI", linha(43))
dd.column_dimensions["A"].width = 42
dd.column_dimensions["B"].width = 12

# ---- Calculo mensal por cliente (R$ mil)
cm = wb.create_sheet("Calculo Mensal")
cm["A1"] = "Calculo mensal por cliente (R$ mil)"
cm["A1"].font = Font(bold=True, size=13)
BLK = {}


def fill_row(s, r, label, fn, bold=False, total=True):
    s.cell(r, 1, label).font = Font(bold=bold)
    for i in range(12):
        c = s.cell(r, 3 + i, fn(L(3 + i)))
        c.number_format = N0
        c.font = Font(bold=bold)
    if total:
        c = s.cell(r, 15, f"=SUM(C{r}:N{r})")
        c.number_format = N0
        c.font = Font(bold=bold)


row = 3
for c in CLI:
    head(cm, row, [c, ""] + MESES + ["Total"])
    row += 1
    fill_row(cm, row, "Volume (mil pecas)", lambda col: f"='Dados IBI'!{col}{R[('vol',c)]}/1000")
    BLK[(c, "vol")] = row
    row += 1
    fill_row(cm, row, "NET SALES", lambda col: f"='Dados IBI'!{col}{R[('fat',c)]}", True)
    BLK[(c, "ns")] = row
    row += 1
    first = row
    fill_row(cm, row, "Variable Cost - Direct Materials", lambda col: f"='Dados IBI'!{col}{R[('mat',c)]}")
    BLK[(c, 21)] = row
    row += 1
    for g, n, ln, d in CUSTOS:
        pcell = f"Premissas!$C${PR_ROW[ln]}"

        def fn(col, ln=ln, pcell=pcell):
            tot = f"'Dados IBI'!{col}{R[('c', ln)]}"
            vol = f"IFERROR('Dados IBI'!{col}{R[('vol', c)]}/'Dados IBI'!{col}{R[('vol','T')]},0)"
            fat = f"IFERROR('Dados IBI'!{col}{R[('fat', c)]}/'Dados IBI'!{col}{R[('fat','T')]},0)"
            fiat = "1" if c == "Fiat" else "0"
            return f'={tot}*IF({pcell}="Volume",{vol},IF({pcell}="Faturamento",{fat},{fiat}))'

        fill_row(cm, row, f"{g} - {n}", fn)
        BLK[(c, ln)] = row
        row += 1
    last = row - 1
    fill_row(cm, row, "TOTAL COST", lambda col: f"=SUM({col}{first}:{col}{last})", True)
    BLK[(c, "tc")] = row
    row += 1
    fill_row(cm, row, "EBIT", lambda col: f"={col}{BLK[(c,'ns')]}+{col}{BLK[(c,'tc')]}", True)
    BLK[(c, "ebit")] = row
    row += 2
cm.column_dimensions["A"].width = 38
cm.freeze_panes = "C3"

# ---- Rentabilidade por cliente (anual)
rc = wb.create_sheet("Rentabilidade por Cliente", 0)
rc["A1"] = "Rentabilidade por cliente - IBIRITE - MP27 (R$ mil)"
rc["A1"].font = Font(bold=True, size=13)
rc["A2"] = "Faturamento e material de consumo diretos; demais custos rateados (ver Premissas)."
head(rc, 4, ["Linha"] + CLI + ["Total IBI", "Arquivo (aba IBI)", "Check"])
TOT = "O"


def ref(c, key):
    return f"='Calculo Mensal'!{TOT}{BLK[(c,key)]}"


def line_clients(r, label, key, fmt=N0, bold=False, total=True):
    rc.cell(r, 1, label).font = Font(bold=bold)
    for j, c in enumerate(CLI):
        x = rc.cell(r, 2 + j, ref(c, key))
        x.number_format = fmt
        x.font = Font(bold=bold)
    if total:
        x = rc.cell(r, 5, f"=SUM(B{r}:D{r})")
        x.number_format = fmt
        x.font = Font(bold=bold)


def line_calc(r, label, fn, fmt, bold=False):
    rc.cell(r, 1, label).font = Font(bold=bold)
    for j in range(4):
        x = rc.cell(r, 2 + j, fn(L(2 + j)))
        x.number_format = fmt
        x.font = Font(bold=bold)


r = 5
rows = {}
rows["vol"] = r
line_clients(r, "Pieces (mil)", "vol")
r += 1
rows["ns"] = r
line_clients(r, "NET SALES", "ns", bold=True)
r += 1
line_calc(r, "Net Sales / Pc (R$)", lambda c: f"=IFERROR({c}{rows['ns']}/{c}{rows['vol']},0)", '0.00')
r += 1
vrow = r
r += 1
for g, n, ln, d in [("Variable Cost", "Direct Materials", 21, None)] + VARS:
    line_clients(r, "   " + n, ln)
    r += 1
line_calc(vrow, "Variable Cost", lambda c: f"=SUM({c}{vrow+1}:{c}{r-1})", N0, True)
rows["cm"] = r
line_calc(r, "CONTRIBUTION MARGIN", lambda c: f"={c}{rows['ns']}+{c}{vrow}", N0, True)
r += 1
line_calc(r, "% on sales", lambda c: f"=IFERROR({c}{rows['cm']}/{c}{rows['ns']},0)", PC)
r += 1
frow = r
r += 1
for g, n, ln, d in FIXS:
    line_clients(r, "   " + n, ln)
    r += 1
line_calc(frow, "Fixed Cost", lambda c: f"=SUM({c}{frow+1}:{c}{r-1})", N0, True)
rows["ebit"] = r
line_calc(r, "EBIT", lambda c: f"={c}{rows['cm']}+{c}{frow}", N0, True)
r += 1
line_calc(r, "ROS %", lambda c: f"=IFERROR({c}{rows['ebit']}/{c}{rows['ns']},0)", PC)
r += 1
line_calc(r, "EBIT / Pc (R$)", lambda c: f"=IFERROR({c}{rows['ebit']}/{c}{rows['vol']},0)", '0.00')
r += 1
rc.cell(r, 1, "% do EBIT total")
for j in range(3):
    x = rc.cell(r, 2 + j, f"=IFERROR({L(2+j)}{rows['ebit']}/$E{rows['ebit']},0)")
    x.number_format = PC

# conferencia contra os totais da aba IBI
fr, er = R[("fat", "T")], R["ebit"]
rc.cell(rows["ns"], 6, f"=SUM('Dados IBI'!C{fr}:N{fr})").number_format = N0
rc.cell(rows["ebit"], 6, f"=SUM('Dados IBI'!C{er}:N{er})").number_format = N0
partes = [f"SUM('Dados IBI'!C{R[('c',ln)]}:N{R[('c',ln)]})" for g, n, ln, d in VARS]
partes.append(f"SUM('Dados IBI'!C{R[('mat','Fiat')]}:N{R[('mat','CNH')]})")
rc.cell(vrow, 6, "=" + "+".join(partes)).number_format = N0
for k in (rows["ns"], rows["ebit"], vrow):
    rc.cell(k, 7, f'=IF(ABS(E{k}-F{k})<0.01,"OK","DIVERGE")')
rc.column_dimensions["A"].width = 30
for col in "BCDEFG":
    rc.column_dimensions[col].width = 16
rc.freeze_panes = "B5"

# ---- EBIT mensal por cliente
em = wb.create_sheet("EBIT Mensal", 1)
head(em, 1, ["EBIT (R$ mil)"] + MESES + ["Total"])
for k, c in enumerate(CLI):
    em.cell(2 + k, 1, c)
    for i in range(13):
        em.cell(2 + k, 2 + i, f"='Calculo Mensal'!{L(3+i)}{BLK[(c,'ebit')]}").number_format = N0
em.cell(5, 1, "Total IBI").font = B
for i in range(13):
    x = em.cell(5, 2 + i, f"=SUM({L(2+i)}2:{L(2+i)}4)")
    x.number_format = N0
    x.font = B
em.column_dimensions["A"].width = 16
wb.save(out)
print("ok", out)
