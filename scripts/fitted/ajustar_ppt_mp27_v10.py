"""PPT MP27 Fitted v10 = v9 (dela) com Cost Var. + ANTT juntos numa so barra nos bridges dos slides 3 (Brasil) e 7 (IBI).
Valores do MENS FITTED v4 (mesmas formulas da 'Tabela Variantes'); nada mais muda.
Uso: python ajustar_ppt_mp27_v10.py <v9.pptx> <mens_v4.xlsx> <saida.pptx>
"""
import sys

import openpyxl
from pptx import Presentation
from pptx.util import Emu

DECK, MENS, OUT = sys.argv[1:4]
MINUS = '−'
JAGUAR, ANTT = -0.55, -1.497
wb = openpyxl.load_workbook(MENS, data_only=True)


def bridge(sheet):
    ws = wb[sheet]
    T = lambda r: float(ws.cell(r, 20).value or 0)
    start, end = float(ws.cell(43, 19).value) / 1000, float(ws.cell(43, 17).value) / 1000
    d = [('Volume', T(8) / 1000), ('Price/ Mix', (T(13) + T(16)) / 1000 - JAGUAR), ('Encerr. Jaguar', JAGUAR),
         ('Labour', (T(19) + T(32)) / 1000),
         ('Cost Var.', sum(T(r) for r in range(20, 24)) / 1000),      # ja inclui o ANTT (−1,497)
         ('Cost Fix.', sum(T(r) for r in range(33, 42)) / 1000)]
    assert abs(start + sum(v for _, v in d) - end) < 0.003
    return {'start': start, 'deltas': d, 'end': end}


def f1(x):
    return (MINUS if round(x, 1) < 0 else '') + f"{abs(x):.1f}".replace('.', ',')


def sg(x):
    return '0,0' if round(abs(x), 1) == 0 else (MINUS if x < 0 else '+') + f"{abs(x):.1f}".replace('.', ',')


prs = Presentation(DECK)
YT, YB, Y_BR, LBL_D, LBL_T = 4950000, 6420000, 4480000, 150000, 210000
BW, L0, L_LAST = 561444, 472055, 7716500


def run(n, spec):
    s = prs.slides[n - 1]
    grp = [x for x in s.shapes if x.shape_type == 6 and x.name.startswith('Bridge MP27')][0]
    g = {x.name: x for x in grp.shapes}
    con = [x for x in grp.shapes if x.name == 'Conector']
    for nm in ('Barra ANTT', 'Valor ANTT', 'Categoria ANTT'):
        grp._element.remove(g[nm]._element)
    grp._element.remove(con[5]._element)            # conector Cost Var->ANTT... (6o); o ANTT->Cost Fix vira Cost Var->Cost Fix
    con = [x for x in grp.shapes if x.name == 'Conector']
    cols = [('Fcst R9', 'tot', spec['start'])] + [(k, 'd', v) for k, v in spec['deltas']] + [('MP’27', 'tot', spec['end'])]
    N = len(cols)
    pitch = (L_LAST - L0) / (N - 1)
    left = lambda i: int(L0 + i * pitch)
    bars, cum = [], 0.0
    for _, kind, v in cols:
        if kind == 'tot':
            lo, hi, cum = min(0, v), max(0, v), v
        else:
            lo, hi = min(cum, cum + v), max(cum, cum + v)
            cum += v
        bars.append((lo, hi, cum))
    vmax = max(max(b[1] for b in bars), 0)
    vmin = min(min(b[0] for b in bars), 0)
    sc = (YB - (200000 if vmin < 0 else 0) - YT) / (vmax - vmin)
    y = lambda v: int(YT + (vmax - v) * sc)
    for i, (lab, kind, v) in enumerate(cols):
        lo, hi, lvl = bars[i]
        b = g[f'Barra {lab}']
        b.left, b.top, b.height = Emu(left(i)), Emu(y(hi)), Emu(max(y(lo) - y(hi), 12700))
        cx = left(i) + BW // 2
        t = g[f'Valor {lab}']
        t.left = Emu(cx - t.width // 2)
        c = g[f'Categoria {lab}']
        c.left = Emu(cx - c.width // 2)
        if kind == 'tot':
            t.top = Emu((y(0) + 25000) if v < 0 else (y(hi) - 25000 - LBL_T))
            txt = f1(v)
        else:
            t.top = Emu(y(hi) - 25000 - LBL_D)
            txt = sg(v)
        p = t.text_frame.paragraphs[0]
        p.runs[0].text = txt
        for r in p.runs[1:]:
            r.text = ''
        if i < N - 1:
            k = con[i]
            k.left, k.width, k.top = Emu(left(i) + BW), Emu(left(i + 1) - left(i) - BW), Emu(y(lvl))
    lt = lambda i: (y(0) - 20000) if cols[i][2] < 0 else (y(bars[i][1]) - 25000 - LBL_T - 30000)
    xe, xd = left(0) + BW // 2, left(N - 1) + BW // 2
    g['Colchete esq'].left, g['Colchete esq'].height = Emu(xe), Emu(lt(0) - Y_BR)
    g['Colchete dir'].left, g['Colchete dir'].height = Emu(xd), Emu(lt(N - 1) - Y_BR)
    g['Colchete topo'].left, g['Colchete topo'].width = Emu(xe), Emu(xd - xe)
    v = g['Variacao']
    v.left = Emu((xe + xd) // 2 - v.width // 2)
    if s.has_notes_slide:
        notes = s.notes_slide.notes_text_frame
        vals = {lab.replace('’', "'"): val for lab, _, val in cols}
        out = []
        for ln in notes.text.split('\n'):
            k = ln.split('\t')[0]
            if '\t' in ln and k == 'ANTT':
                continue
            out.append(f"{k}\t{vals[k]:.3f}" if '\t' in ln and k in vals else ln)
        notes.text = '\n'.join(out)
    print('slide', n, [(k, round(v, 3)) for k, v in spec['deltas']])


run(3, bridge('TOTAL'))
run(7, bridge('IBI'))
prs.save(OUT)
print('salvo', OUT, len(prs.slides))
