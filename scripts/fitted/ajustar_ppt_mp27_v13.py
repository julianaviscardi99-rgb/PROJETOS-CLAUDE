"""PPT MP27 Fitted v13 = v12 (dela) com SO o slide 8 (IBI | Rentabilidade por cliente) atualizado: tabela + comentarios.
Fonte: Rentabilidade por Cliente IBI v3 (19:16) e MENS v4 (aba 'Impacto JLR', volume R9 do Iveco). Nenhum outro slide e tocado.
Uso: python ajustar_ppt_mp27_v13.py <v12.pptx> <rentab_ibi_v3.xlsx> <mens_v4.xlsx> <saida.pptx>
"""
import re
import sys

import openpyxl
from pptx import Presentation

DECK, RENT, MENS, OUT = sys.argv[1:5]
MINUS = '−'
R = openpyxl.load_workbook(RENT, data_only=True)['Rentabilidade por Cliente']
row = {R.cell(r, 1).value: r for r in range(1, 40)}
CL = ['Fiat', 'Iveco', 'CNH', 'Total']
v = lambda lab, c: float(R.cell(row[lab], 2 + CL.index(c)).value)
dec = lambda x, n=1: f"{x:.{n}f}".replace('.', ',')
mil = lambda x: f"{x:,.0f}".replace(',', '.')
par = lambda x: f"({dec(abs(x))})" if x < 0 else dec(x)

ns, eb, cm = {c: v('NET SALES', c) for c in CL}, {c: v('EBIT', c) for c in CL}, {c: v('CONTRIBUTION MARGIN', c) for c in CL}
pcs, fx = {c: v('Pieces (mil)', c) for c in CL}, {c: v('Fixed Cost', c) for c in CL}
ros = {c: eb[c] / ns[c] for c in CL}

ws = openpyxl.load_workbook(MENS, data_only=True)['Impacto JLR']
r_iv = [r for r in range(1, 40) if ws.cell(r, 3).value == 'Iveco'][0]       # 1o bloco: vol R9, vol MP27
iv = [ws.cell(r_iv, 4).value, ws.cell(r_iv, 5).value]
iveco_vol = iv[1] / iv[0] - 1

prs = Presentation(DECK)
s8 = prs.slides[7]
assert 'Rentabilidade por cliente' in s8.shapes.title.text_frame.text
sh = {x.name: x for x in s8.shapes}


def setcell(cell, txt):
    p = cell.text_frame.paragraphs[0]
    p.runs[0].text = txt
    for r in p.runs[1:]:
        r.text = ''


# ------------------------------------------------------------------ tabela (so os valores, formatacao intacta)
T = sh['Table 95'].table
body = {
    'Peças (mil)': [mil(pcs[c]) for c in CL],
    'Net Sales': [dec(ns[c] / 1000) for c in CL],
    'NS/pç (R$)': [dec(ns[c] / pcs[c], 2) for c in CL],
    'CM': [dec(cm[c] / 1000) for c in CL],
    'CM % da receita': [dec(cm[c] / ns[c] * 100) + '%' for c in CL],
    'Fixed Cost': [par(fx[c] / 1000) for c in CL],
    'EBIT': [dec(eb[c] / 1000) for c in CL],
    'EBIT/pç (R$)': [dec(eb[c] / pcs[c], 2) for c in CL],
    'ROS %': [dec(ros[c] * 100) + '%' for c in CL],
}
for r in T.rows:
    k = r.cells[0].text
    if k in body:
        for i, txt in enumerate(body[k]):
            setcell(r.cells[1 + i], txt)
        print(k, body[k])

# ------------------------------------------------------------------ comentarios (troca dentro dos runs, formatacao intacta)
com = sh['Comentarios'].text_frame.paragraphs


def sub(pi, rx, new):
    for r in com[pi].runs:
        if re.search(rx, r.text):
            r.text = re.sub(rx, new, r.text, count=1)
            return
    raise AssertionError((pi, rx, com[pi].text))


rec = lambda c: ns[c] / ns['Total'] * 100
ebs = lambda c: eb[c] / eb['Total'] * 100
sub(1, r'^\d+(?=% da )', f"{rec('Fiat'):.0f}")
sub(1, r'(?<= e )\d+(?=% do EBIT)', f"{ebs('Fiat'):.0f}")
sub(1, r'R\$ [\d,]+ mi', f"R$ {dec(eb['Fiat'] / 1000)} mi")
sub(2, r'[\d,]+(?=%)', dec(ros['Fiat'] * 100))
sub(4, r'ROS [\d,]+%', f"ROS {dec(ros['Iveco'] * 100)}%")
sub(4, r'R\$ [\d,]+ mi', f"R$ {dec(eb['Iveco'] / 1000)} mi")
sub(5, r'[−+][\d,]+%', f"{MINUS if iveco_vol < 0 else '+'}{dec(abs(iveco_vol) * 100)}%")
sub(7, r'ROS [\d,]+%', f"ROS {dec(ros['CNH'] * 100)}%")
sub(7, r'R\$ [\d,]+ mi', f"R$ {dec(eb['CNH'] / 1000)} mi")
ic_rec = (ns['Iveco'] + ns['CNH']) / ns['Total'] * 100
ic_eb = (eb['Iveco'] + eb['CNH']) / eb['Total'] * 100
lo, hi = sorted((cm['Iveco'] / ns['Iveco'], cm['CNH'] / ns['CNH']))
sub(11, r'^Iveco e CNH: \d+%', f"Iveco e CNH: {ic_rec:.0f}%")
sub(11, r'(?<= e )\d+(?=% do EBIT)', f"{ic_eb:.0f}")
sub(11, r'de \d+% a \d+%', f"de {lo * 100:.0f}% a {hi * 100:.0f}%")
sub(11, r'contra \d+%', f"contra {cm['Fiat'] / ns['Fiat'] * 100:.0f}%")
sub(13, r'R\$ [\d,]+ mi', f"R$ {dec(ns['Fiat'] / 100 / 1000)} mi")        # 1 p.p. de ROS da Fiat
sub(15, r'R\$ [\d,]+ mi', f"R$ {dec(eb['Total'] / 1000)} mi")
sub(15, r'ROS de [\d,]+%', f"ROS de {dec(ros['Total'] * 100)}%")
print('Iveco vol', round(iveco_vol * 100, 1), 'Iveco+CNH', round(ic_rec, 1), round(ic_eb, 1), round(lo * 100, 1), round(hi * 100, 1))
prs.save(OUT)
print('salvo', OUT, len(prs.slides))
