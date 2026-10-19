"""Rentabilidade por cliente - Ibirite (Fiat / Iveco / CNH) a partir da MENS FITTED MP27 - v3.

Diferencas da v2:
- aba "Dados IBI" LIGADA (link externo) a aba IBI da MENS FITTED (faturamento, volume, material e custos);
- custos especificos de CNH e Iveco (Labour e Handling) informados na aba Premissas e REALOCADOS do total
  de IBI (o EBIT total continua fechando com a MENS); o restante e rateado pelo driver escolhido;
- driver padrao = Faturamento (melhora a rentabilidade da Fiat vs. Volume);
- meta de ROS minimo da Fiat (padrao 10,5%): se ficar abaixo, o que falta e realocado para Iveco/CNH.

Uso: python rentabilidade_cliente_ibi_v3.py <MENS FITTED.xlsx> <saida.xlsx>
"""
import sys
import os
import shutil
import tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

src, out = sys.argv[1:3]
src = os.path.abspath(src)
out = os.path.abspath(out)
ws = openpyxl.load_workbook(src, data_only=True)["IBI"]
MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
CLI = ["Fiat", "Iveco", "CNH"]


def linha(r, k=1.0):
    return [(ws.cell(r, 5 + i).value or 0) * k for i in range(12)]


# linhas da aba IBI (custos em R$ mil; volume em pecas; faturamento em R$)
VOL = {"Fiat": 53, "Iveco": 54, "CNH": 55}
FAT = {"Fiat": 65, "Iveco": 66, "CNH": 67}
MAT = {"Fiat": 75, "Iveco": 76, "CNH": 77}
CUSTOS = [  # (grupo, nome, linha, driver padrao)
    ("Variable Cost", "Labour", 19, "Faturamento"),
    ("Variable Cost", "Handling", 20, "Faturamento"),
    ("Variable Cost", "Transportation", 22, "Fiat 100%"),
    ("Variable Cost", "Other Variable", 23, "Faturamento"),
    ("Fixed Cost", "Labour", 32, "Faturamento"),
    ("Fixed Cost", "Depreciation", 33, "Faturamento"),
    ("Fixed Cost", "IFRS16", 34, "Faturamento"),
    ("Fixed Cost", "Rents", 35, "Faturamento"),
    ("Fixed Cost", "Condominio", 36, "Faturamento"),
    ("Fixed Cost", "Other Fixed", 37, "Faturamento"),
]
VARS = [c for c in CUSTOS if c[0] == "Variable Cost"]
FIXS = [c for c in CUSTOS if c[0] == "Fixed Cost"]

HDR = PatternFill("solid", fgColor="1F3864")
INP = PatternFill("solid", fgColor="FFF2CC")
N0 = '#,##0.0;[Red]-#,##0.0;"-"'
PC = '0.0%;[Red]-0.0%;"-"'
B = Font(bold=True)


def head(s, r, vals, c0=1):
    for i, t in enumerate(vals):
        c = s.cell(r, c0 + i, t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)


wb = openpyxl.Workbook()

# ---- Premissas
pr = wb.active
pr.title = "Premissas"
pr["A1"] = "Rentabilidade por cliente - IBIRITE - MP27 (R$ mil)"
pr["A1"].font = Font(bold=True, size=13)
pr["A2"] = ("Fonte: MENS FITTED MP27_v3_teste daniel.xlsx, aba IBI. Faturamento e Material de consumo = diretos por cliente. "
            "Demais custos = rateio pelo driver abaixo (celulas amarelas editaveis: Volume, Faturamento ou Fiat 100%). Frete (Transportation) = 100% Fiat, confirmado pela Juliana.")
head(pr, 4, ["Grupo", "Linha de custo", "Driver de rateio"])
dv = DataValidation(type="list", formula1='"Volume,Faturamento,Fiat 100%"', allow_blank=False)
pr.add_data_validation(dv)
PR_ROW = {}
for i, (g, n, r, d) in enumerate(CUSTOS):
    rr = 5 + i
    pr.cell(rr, 1, g)
    pr.cell(rr, 2, n)
    c = pr.cell(rr, 3, d)
    c.fill = INP
    dv.add(c)
    PR_ROW[r] = rr
pr.cell(5 + len(CUSTOS), 1, "Variable Cost")
pr.cell(5 + len(CUSTOS), 2, "Direct Materials")
pr.cell(5 + len(CUSTOS), 3, "Direto (material de consumo de cada cliente)")
for col, w in zip("ABCD", (16, 22, 44, 14)):
    pr.column_dimensions[col].width = w

# custos especificos por cliente (R$ mil/ano) - realocados do total IBI
SPEC_ROW0 = 5 + len(CUSTOS) + 3
pr.cell(SPEC_ROW0 - 1, 1, "Custos especificos por cliente (R$ mil/ano) - realocados do total IBI (nao criam custo novo); o restante vai pelo driver").font = B
head(pr, SPEC_ROW0, ["Linha", "Fiat", "Iveco", "CNH"])
SPECS = [  # (linha aba IBI, nome, {cliente: R$ mil/ano})
    (19, "Variable - Labour", {"Fiat": 0, "Iveco": 110, "CNH": 240}),
    (20, "Variable - Handling", {"Fiat": 0, "Iveco": 63, "CNH": 400}),
]
SPEC_CELL = {}
for k, (ln, nome, d) in enumerate(SPECS):
    rr = SPEC_ROW0 + 1 + k
    pr.cell(rr, 1, nome)
    for j, c in enumerate(CLI):
        x = pr.cell(rr, 2 + j, d[c])
        x.fill = INP
        x.number_format = N0
        SPEC_CELL[(ln, c)] = f"Premissas!${L(2 + j)}${rr}"
pr.cell(SPEC_ROW0 + 3, 1, "CNH: Handling +400 (area utilizavel 2.470 m2 + empilhadeira); Labour +240 (2 pessoas dedicadas). "
        "Iveco: Labour +110; Handling +63 (area utilizavel 450 m2). Distribuidos igualmente nos 12 meses.")
ROS_ROW = SPEC_ROW0 + 5
pr.cell(ROS_ROW, 1, "ROS minimo Fiat").font = B
ROS_CELL = f"Premissas!$B${ROS_ROW}"
x = pr.cell(ROS_ROW, 2, 0.105)
x.fill = INP
x.number_format = "0.0%"
pr.cell(ROS_ROW + 1, 1, "Se o ROS da Fiat ficar abaixo da meta, a diferenca e realocada da Fiat para Iveco/CNH na proporcao abaixo.")
pr.cell(ROS_ROW + 2, 1, "% da realocacao -> Iveco").font = B
pr.cell(ROS_ROW + 3, 1, "% da realocacao -> CNH").font = B
SH_IVE = f"Premissas!$B${ROS_ROW + 2}"
SH_CNH = f"Premissas!$B${ROS_ROW + 3}"

# ---- Dados IBI (valores do arquivo)
dd = wb.create_sheet("Dados IBI")
head(dd, 1, ["Item", "Cliente"] + MESES)
R = {}
state = {"r": 2}
LINKS = []  # (linha em Dados IBI, linha na aba IBI da MENS, divisor)


def put(key, item, cli, vals, ibi=None, div=1):
    r = state["r"]
    dd.cell(r, 1, item)
    dd.cell(r, 2, cli)
    for i, v in enumerate(vals):
        dd.cell(r, 3 + i, v).number_format = N0
    if ibi:
        LINKS.append((r, ibi, div))
    R[key] = r
    state["r"] += 1


def soma3(key):
    return [f"=SUM({L(3+i)}{R[(key,'Fiat')]}:{L(3+i)}{R[(key,'CNH')]})" for i in range(12)]


for c in CLI:
    put(("vol", c), "Volume (pecas)", c, linha(VOL[c]), VOL[c])
put(("vol", "T"), "Volume (pecas)", "Total", soma3("vol"))
for c in CLI:
    put(("fat", c), "Faturamento (R$ mil)", c, linha(FAT[c], 1 / 1000), FAT[c], 1000)
put(("fat", "T"), "Faturamento (R$ mil)", "Total", soma3("fat"))
for c in CLI:
    put(("mat", c), "Material de consumo (R$ mil)", c, linha(MAT[c], 1 / 1000), MAT[c], 1000)
for g, n, row, d in CUSTOS:
    put(("c", row), f"{g} - {n} (R$ mil)", "Total IBI", linha(row), row)
put("ebit", "EBIT total IBI no arquivo (R$ mil)", "Total IBI", linha(43), 43)
dd.column_dimensions["A"].width = 42
dd.column_dimensions["B"].width = 12

# repartição da realocação adicional (meta ROS Fiat): proporcional ao faturamento de Iveco e CNH
fi, fc = R[("fat", "Iveco")], R[("fat", "CNH")]
pr[SH_IVE.split("!")[1].replace("$", "")] = f"=SUM('Dados IBI'!C{fi}:N{fi})/(SUM('Dados IBI'!C{fi}:N{fi})+SUM('Dados IBI'!C{fc}:N{fc}))"
pr[SH_CNH.split("!")[1].replace("$", "")] = f"=1-{SH_IVE.split('!')[1]}"
for a in (SH_IVE, SH_CNH):
    pr[a.split("!")[1].replace("$", "")].number_format = "0.0%"
    pr[a.split("!")[1].replace("$", "")].fill = INP

# ---- Calculo mensal por cliente (R$ mil)
cm = wb.create_sheet("Calculo Mensal")
cm["A1"] = "Calculo mensal por cliente (R$ mil)"
cm["A1"].font = Font(bold=True, size=13)
BLK = {}


def fill_row(s, r, label, fn, bold=False, total=True):
    s.cell(r, 1, label).font = Font(bold=bold)
    for i in range(12):
        c = s.cell(r, 3 + i, fn(L(3 + i)))
        c.number_format = N0
        c.font = Font(bold=bold)
    if total:
        c = s.cell(r, 15, f"=SUM(C{r}:N{r})")
        c.number_format = N0
        c.font = Font(bold=bold)


row = 3
for c in CLI:
    head(cm, row, [c, ""] + MESES + ["Total"])
    row += 1
    fill_row(cm, row, "Volume (mil pecas)", lambda col: f"='Dados IBI'!{col}{R[('vol',c)]}/1000")
    BLK[(c, "vol")] = row
    row += 1
    fill_row(cm, row, "NET SALES", lambda col: f"='Dados IBI'!{col}{R[('fat',c)]}", True)
    BLK[(c, "ns")] = row
    row += 1
    first = row
    fill_row(cm, row, "Variable Cost - Direct Materials", lambda col: f"='Dados IBI'!{col}{R[('mat',c)]}")
    BLK[(c, 21)] = row
    row += 1
    for g, n, ln, d in CUSTOS:
        pcell = f"Premissas!$C${PR_ROW[ln]}"
        espec = [(l2, c2) for (l2, c2) in SPEC_CELL if l2 == ln]

        def fn(col, ln=ln, pcell=pcell, espec=espec):
            tot = f"'Dados IBI'!{col}{R[('c', ln)]}"
            vol = f"IFERROR('Dados IBI'!{col}{R[('vol', c)]}/'Dados IBI'!{col}{R[('vol','T')]},0)"
            fat = f"IFERROR('Dados IBI'!{col}{R[('fat', c)]}/'Dados IBI'!{col}{R[('fat','T')]},0)"
            fiat = "1" if c == "Fiat" else "0"
            if espec:  # custos especificos saem do total (realocados) e o restante vai pelo driver
                soma = "+".join(SPEC_CELL[e] for e in espec)
                pool = f"({tot}+({soma})/12)"
                prop = f"-{SPEC_CELL[(ln, c)]}/12+"
            else:
                pool, prop = tot, ""
            return f'={prop}{pool}*IF({pcell}="Volume",{vol},IF({pcell}="Faturamento",{fat},{fiat}))'

        fill_row(cm, row, f"{g} - {n}", fn)
        BLK[(c, ln)] = row
        row += 1
    last = row - 1
    # realocacao adicional p/ meta de ROS minimo da Fiat (Fiat recebe +, Iveco/CNH pagam -)
    if c == "Fiat":
        ns_r = BLK[(c, "ns")]
        need = f"MAX(0,{ROS_CELL}*$O${ns_r}-($O${ns_r}+SUM($O${first}:$O${last})))"
        fill_row(cm, row, "Realocacao p/ meta ROS Fiat", lambda col: f"={need}*IFERROR({col}{ns_r}/$O${ns_r},0)")
    else:
        sh = SH_IVE if c == "Iveco" else SH_CNH
        fiat_x = BLK[("Fiat", "x")]
        fill_row(cm, row, "Realocacao p/ meta ROS Fiat", lambda col, sh=sh, fiat_x=fiat_x: f"=-{col}{fiat_x}*{sh}")
    BLK[(c, "x")] = row
    row += 1
    fill_row(cm, row, "TOTAL COST", lambda col: f"=SUM({col}{first}:{col}{row-1})", True)
    BLK[(c, "tc")] = row
    row += 1
    fill_row(cm, row, "EBIT", lambda col: f"={col}{BLK[(c,'ns')]}+{col}{BLK[(c,'tc')]}", True)
    BLK[(c, "ebit")] = row
    row += 2
cm.column_dimensions["A"].width = 38
cm.freeze_panes = "C3"

# ---- Rentabilidade por cliente (anual)
rc = wb.create_sheet("Rentabilidade por Cliente", 0)
rc["A1"] = "Rentabilidade por cliente - IBIRITE - MP27 (R$ mil)"
rc["A1"].font = Font(bold=True, size=13)
rc["A2"] = "Faturamento e material de consumo diretos; demais custos rateados (ver Premissas)."
head(rc, 4, ["Linha"] + CLI + ["Total IBI", "Arquivo (aba IBI)", "Check"])
TOT = "O"


def ref(c, key):
    return f"='Calculo Mensal'!{TOT}{BLK[(c,key)]}"


def line_clients(r, label, key, fmt=N0, bold=False, total=True):
    rc.cell(r, 1, label).font = Font(bold=bold)
    for j, c in enumerate(CLI):
        x = rc.cell(r, 2 + j, ref(c, key))
        x.number_format = fmt
        x.font = Font(bold=bold)
    if total:
        x = rc.cell(r, 5, f"=SUM(B{r}:D{r})")
        x.number_format = fmt
        x.font = Font(bold=bold)


def line_calc(r, label, fn, fmt, bold=False):
    rc.cell(r, 1, label).font = Font(bold=bold)
    for j in range(4):
        x = rc.cell(r, 2 + j, fn(L(2 + j)))
        x.number_format = fmt
        x.font = Font(bold=bold)


r = 5
rows = {}
rows["vol"] = r
line_clients(r, "Pieces (mil)", "vol")
r += 1
rows["ns"] = r
line_clients(r, "NET SALES", "ns", bold=True)
r += 1
line_calc(r, "Net Sales / Pc (R$)", lambda c: f"=IFERROR({c}{rows['ns']}/{c}{rows['vol']},0)", '0.00')
r += 1
vrow = r
r += 1
for g, n, ln, d in [("Variable Cost", "Direct Materials", 21, None)] + VARS:
    line_clients(r, "   " + n, ln)
    r += 1
line_calc(vrow, "Variable Cost", lambda c: f"=SUM({c}{vrow+1}:{c}{r-1})", N0, True)
rows["cm"] = r
line_calc(r, "CONTRIBUTION MARGIN", lambda c: f"={c}{rows['ns']}+{c}{vrow}", N0, True)
r += 1
line_calc(r, "% on sales", lambda c: f"=IFERROR({c}{rows['cm']}/{c}{rows['ns']},0)", PC)
r += 1
frow = r
r += 1
for g, n, ln, d in FIXS:
    line_clients(r, "   " + n, ln)
    r += 1
line_calc(frow, "Fixed Cost", lambda c: f"=SUM({c}{frow+1}:{c}{r-1})", N0, True)
rows["x"] = r
line_clients(r, "Realocacao p/ meta ROS Fiat", "x")
r += 1
rows["ebit"] = r
line_calc(r, "EBIT", lambda c: f"={c}{rows['cm']}+{c}{frow}+{c}{rows['x']}", N0, True)
r += 1
rows["ros"] = r
line_calc(r, "ROS %", lambda c: f"=IFERROR({c}{rows['ebit']}/{c}{rows['ns']},0)", PC)
r += 1
line_calc(r, "EBIT / Pc (R$)", lambda c: f"=IFERROR({c}{rows['ebit']}/{c}{rows['vol']},0)", '0.00')
r += 1
rc.cell(r, 1, "% do EBIT total")
for j in range(3):
    x = rc.cell(r, 2 + j, f"=IFERROR({L(2+j)}{rows['ebit']}/$E{rows['ebit']},0)")
    x.number_format = PC

# conferencia contra os totais da aba IBI
fr, er = R[("fat", "T")], R["ebit"]
rc.cell(rows["ns"], 6, f"=SUM('Dados IBI'!C{fr}:N{fr})").number_format = N0
rc.cell(rows["ebit"], 6, f"=SUM('Dados IBI'!C{er}:N{er})").number_format = N0
partes = [f"SUM('Dados IBI'!C{R[('c',ln)]}:N{R[('c',ln)]})" for g, n, ln, d in VARS]
partes.append(f"SUM('Dados IBI'!C{R[('mat','Fiat')]}:N{R[('mat','CNH')]})")
rc.cell(vrow, 6, "=" + "+".join(partes)).number_format = N0
for k in (rows["ns"], rows["ebit"], vrow):
    rc.cell(k, 7, f'=IF(ABS(E{k}-F{k})<0.01,"OK","DIVERGE")')
rc.cell(rows["x"], 7, f'=IF(ABS(E{rows["x"]})<0.01,"OK","DIVERGE")')
r += 2
rc.cell(r, 1, "Meta ROS minimo Fiat").font = B
rc.cell(r, 2, f"={ROS_CELL}").number_format = PC
rc.cell(r + 1, 1, "ROS Fiat").font = B
rc.cell(r + 1, 2, f"=B{rows['ros']}").number_format = PC
rc.cell(r + 2, 1, "Status").font = B
rc.cell(r + 2, 2, f'=IF(B{r+1}>=B{r}-0.00001,"ATENDE","ABAIXO DA META")')
rc.cell(r + 3, 1, "Realocacao adicional necessaria (R$ mil)").font = B
rc.cell(r + 3, 2, f"=B{rows['x']}").number_format = N0
rc.column_dimensions["A"].width = 38
for col in "BCDEFG":
    rc.column_dimensions[col].width = 16
rc.freeze_panes = "B5"

# ---- EBIT mensal por cliente
em = wb.create_sheet("EBIT Mensal", 1)
head(em, 1, ["EBIT (R$ mil)"] + MESES + ["Total"])
for k, c in enumerate(CLI):
    em.cell(2 + k, 1, c)
    for i in range(13):
        em.cell(2 + k, 2 + i, f"='Calculo Mensal'!{L(3+i)}{BLK[(c,'ebit')]}").number_format = N0
em.cell(5, 1, "Total IBI").font = B
for i in range(13):
    x = em.cell(5, 2 + i, f"=SUM({L(2+i)}2:{L(2+i)}4)")
    x.number_format = N0
    x.font = B
em.column_dimensions["A"].width = 16
# ---- Comparativo de drivers (valores na geracao; mesma logica da aba Calculo Mensal, sem realocacao)
def simula(drv):
    res = {}
    for c in CLI:
        ns = ebit = 0.0
        for m in range(12):
            vols = {k: (ws.cell(VOL[k], 5 + m).value or 0) for k in CLI}
            fats = {k: (ws.cell(FAT[k], 5 + m).value or 0) / 1000 for k in CLI}
            mats = {k: (ws.cell(MAT[k], 5 + m).value or 0) / 1000 for k in CLI}
            sv = vols[c] / sum(vols.values()) if sum(vols.values()) else 0
            sf = fats[c] / sum(fats.values()) if sum(fats.values()) else 0
            sh = sv if drv == "Volume" else sf
            custo = mats[c]
            for g, n, ln, d in CUSTOS:
                tot = ws.cell(ln, 5 + m).value or 0
                if ln == 22:
                    custo += tot if c == "Fiat" else 0
                    continue
                esp = {l2: {k: SPEC_CELL_VAL[(l2, k)] for k in CLI} for (l2, _) in SPEC_CELL}
                if ln in esp:
                    custo += -esp[ln][c] / 12 + (tot + sum(esp[ln].values()) / 12) * sh
                else:
                    custo += tot * sh
            ns += fats[c]
            ebit += fats[c] + custo
        res[c] = (ns, ebit)
    return res


SPEC_CELL_VAL = {(ln, c): d[c] for ln, _, d in SPECS for c in CLI}
cp = wb.create_sheet("Comparativo Drivers")
cp["A1"] = "Comparativo de drivers de rateio (todas as linhas rateadas pelo mesmo driver; ja com custos especificos de CNH/Iveco) - valores na data da geracao"
cp["A1"].font = Font(bold=True, size=12)
head(cp, 3, ["Driver", "Fiat EBIT", "Fiat ROS", "Iveco EBIT", "Iveco ROS", "CNH EBIT", "CNH ROS", "Fiat >= meta?"])
for k, drv in enumerate(("Volume", "Faturamento")):
    res = simula(drv)
    rr = 4 + k
    cp.cell(rr, 1, drv)
    for j, c in enumerate(CLI):
        ns, eb = res[c]
        cp.cell(rr, 2 + 2 * j, eb).number_format = N0
        cp.cell(rr, 3 + 2 * j, eb / ns if ns else 0).number_format = PC
    cp.cell(rr, 8, f'=IF(C{rr}>={ROS_CELL},"SIM","NAO")')
    print(drv, {c: (round(res[c][1], 1), round(res[c][1] / res[c][0] * 100, 1)) for c in CLI})
cp.column_dimensions["A"].width = 16
for col in "BCDEFGH":
    cp.column_dimensions[col].width = 14

tmp = os.path.join(tempfile.mkdtemp(), "rent_ibi_sem_links.xlsx")
wb.save(tmp)

# ---- liga "Dados IBI" a MENS FITTED via Excel (instancia propria, sem tocar nos arquivos abertos da usuaria)
import win32com.client
xl = win32com.client.DispatchEx("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
xl.AskToUpdateLinks = False
try:
    wsrc = xl.Workbooks.Open(src, 0, True)  # UpdateLinks=0, ReadOnly
    wo = xl.Workbooks.Open(tmp)
    sh = wo.Worksheets("Dados IBI")
    ref = f"'[{os.path.basename(src)}]IBI'!"
    for r_dd, r_ibi, div in LINKS:
        for m in range(12):
            f = f"={ref}{L(5 + m)}{r_ibi}" + (f"/{div}" if div != 1 else "")
            sh.Cells(r_dd, 3 + m).Formula = f
    xl.CalculateFull()
    if os.path.exists(out):
        raise SystemExit("saida ja existe, nao sobrescrevo: " + out)
    wo.SaveAs(out, 51)
    wo.Close(False)
    wsrc.Close(False)
finally:
    xl.Quit()
print("ok", out)
