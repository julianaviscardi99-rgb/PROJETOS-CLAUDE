"""PPT MP27 Fitted Units - v7 = v5 (21/10 09:20, com as tabelas coladas por ela) atualizada com MENS v4, Rentabilidade v3 e Base Book de 21/10.

Uso: python atualizar_ppt_mp27_v7.py <v5.pptx> <basebook.xlsx> <mens_v4.xlsx> <rentab_cliente_ibi_v3.xlsx> <saida_v7.pptx>

Edicao cirurgica: nao muda estrutura nem layout; as tabelas coladas (imagens) dos slides 4, 7, 9, 10 e 11 nao sao tocadas;
slide 14 mantem a estrutura (so valores/textos). Reposiciona as barras dos bridges dos slides 2, 3, 4 e 7, troca textos
(mantendo a formatacao dos runs), celulas das tabelas nativas (5, 6, 8) e dados da pizza.
"""
import sys

import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Emu

DECK, BB, MENS, CLI, OUT = sys.argv[1:6]
MINUS = '−'


# ---------------------------------------------------------------- formatos (iguais aos da rev2)
def num(x, d=1):
    s = f"{abs(x):,.{d}f}".replace(',', 'X').replace('.', ',').replace('X', '.')
    return s


def f1(x, d=1):
    s = num(x, d)
    return (MINUS + s) if round(x, d) < 0 or (x < 0 and s.strip('0,.') == '') else s


def sg(x, d=1):
    if round(abs(x), d) == 0:
        return num(0, d)
    return (MINUS if x < 0 else '+') + num(x, d)


def pc(x, d=1, signed=False):
    s = num(x * 100, d) + '%'
    return ((MINUS if x < 0 else '+') + s) if signed else s


def par(x):
    return f"({num(x)})" if round(x, 1) < 0 else num(x)


# ---------------------------------------------------------------- fontes
bb = openpyxl.load_workbook(BB, data_only=True)


def bridge(sheet, hr, dr):
    ws = bb[sheet]
    items, end = [], None
    for col in range(3, 20):
        h, v = ws.cell(hr, col).value, ws.cell(dr, col).value
        if h is None:
            continue
        h = str(h).strip()
        if h == "MP'27" or h.startswith('Delta') or h.startswith('%'):
            continue
        if h == 'EBIT':
            end = float(v)
            break
        items.append((h.replace('\n', ' '), float(v)))
    assert abs(items[0][1] + sum(v for _, v in items[1:]) - end) < 0.003
    return {'start': items[0][1], 'deltas': items[1:], 'end': end, 'd': dict(items[1:])}


BRG = {'BR': bridge('Tabela Variantes s targ', 2, 4), 'IBI': bridge('Tabela Variantes s targ', 6, 8)}
PLS = {'BR': 'Month_FY_LC - MP27', 'IBI': 'IBIRITE', 'SJP': 'SJP', 'GOI': 'GOIANA', 'RES': 'RESENDE'}
ROWS = {'pcs': 42, 'ns': 44, 'nspc': 45, 'ebit': 55, 'ros': 58}


def pl(u, k, i=4):          # i: 2 = MP26, 3 = R9, 4 = MP27
    return float(bb[PLS[u]].cell(ROWS[k], 39 + i).value or 0)


for u in BRG:
    assert abs(BRG[u]['end'] - pl(u, 'ebit')) < 0.002 and abs(BRG[u]['start'] - pl(u, 'ebit', 3)) < 0.002, u
UN = ['IBI', 'SJP', 'GOI', 'RES']
E27 = {u: pl(u, 'ebit') for u in UN}
assert abs(sum(E27.values()) - pl('BR', 'ebit')) < 0.002

mens = openpyxl.load_workbook(MENS, data_only=True)
jl = mens['Impacto JLR']
JLR = {}
for r in range(1, 60):
    k = jl.cell(r, 3).value
    if k in ('Fiat', 'Iveco', 'CNH') and k not in JLR:
        JLR[k] = [jl.cell(r, c).value for c in range(4, 14)]
FIAT_PR9, FIAT_PR27, FIAT_PRECO = JLR['Fiat'][2], JLR['Fiat'][3], JLR['Fiat'][9] / 1000
FIAT_PCT = FIAT_PR27 / FIAT_PR9 - 1
mt = mens['TOTAL']
trow = {mt.cell(r, 3).value: r for r in range(1, 60) if mt.cell(r, 3).value}
tv = lambda k: float(mt.cell(trow[k], 17).value)
SENS = (tv('NET SALES') + tv('Direct Materials') + tv('Transportation') + tv('Other Variable')) / 1000 / 100

wc = openpyxl.load_workbook(CLI, data_only=True)['Rentabilidade por Cliente']
crow = lambda r: [float(wc.cell(r, c).value or 0) for c in (2, 3, 4, 5)]
C_PCS, C_NS, C_NSPC, C_CM, C_FIX, C_REAL, C_EBIT, C_ROS = (crow(5), crow(6), crow(7), crow(14), crow(16), crow(23),
                                                           crow(24), crow(25))
assert str(wc['A23'].value).startswith('Realocacao') and str(wc['A24'].value) == 'EBIT'
assert abs(C_EBIT[3] / 1000 - pl('IBI', 'ebit')) < 0.002
ROS_META = float(openpyxl.load_workbook(CLI, data_only=True)['Premissas']['B23'].value)
C_FIXR = [a + b for a, b in zip(C_FIX, C_REAL)]      # fixo + realocacao para a meta da Fiat (tabela fecha no EBIT)
REAL_FIAT = C_REAL[0] / 1000

T3_MO, T3_DESP = 2.5, 2.0                             # premissa da Juliana (19/10)
ANTT = -BRG['IBI']['d']['ANTT']
COBERT = FIAT_PRECO / (T3_MO + T3_DESP + ANTT)
RISCO_MO = 0.95 + 0.46 + 0.36
RISCOS = FIAT_PRECO + RISCO_MO
OPORT = ANTT + max(0.0, T3_MO + T3_DESP + ANTT - FIAT_PRECO)

br9, br27, brmp26 = pl('BR', 'ebit', 3), pl('BR', 'ebit'), pl('BR', 'ebit', 2)
ns9, ns27 = pl('BR', 'ns', 3), pl('BR', 'ns')
ros27, ros9 = pl('BR', 'ros'), pl('BR', 'ros', 3)
ib9, ib27, ibmp26 = pl('IBI', 'ebit', 3), pl('IBI', 'ebit'), pl('IBI', 'ebit', 2)
r9_4u = sum(pl(u, 'ebit', 3) for u in UN)
sh_eb = {u: E27[u] / br27 for u in UN}
sh_ns = {u: pl(u, 'ns') / ns27 for u in UN}
cl_ns = [C_NS[i] / C_NS[3] for i in range(3)]
cl_eb = [C_EBIT[i] / C_EBIT[3] for i in range(3)]
sem_fiat_br = br27 - FIAT_PRECO
ros_sem_fiat = sem_fiat_br / (ns27 - FIAT_PRECO)

# ---------------------------------------------------------------- utilitarios de edicao
prs = Presentation(DECK)
LOG = []


def sl(n):
    return prs.slides[n - 1]


def find(slide, name, idx=0):
    out = []

    def walk(shs):
        for x in shs:
            if x.name == name:
                out.append(x)
            if x.shape_type == 6:
                walk(x.shapes)
    walk(slide.shapes)
    return out[idx]


def rep(n, name, old, new, idx=0):
    """Troca 'old' por 'new' no texto da forma mantendo a formatacao do run onde o trecho comeca."""
    if old == new:
        return
    sh = find(sl(n), name, idx)
    for p in sh.text_frame.paragraphs:
        runs = p.runs
        full = ''.join(r.text for r in runs)
        i = full.find(old)
        if i < 0:
            continue
        pos, spans = 0, []
        for r in runs:
            spans.append((pos, pos + len(r.text)))
            pos += len(r.text)
        j = i + len(old)
        a = next(k for k, (s, e) in enumerate(spans) if s <= i < e)
        b = next(k for k, (s, e) in enumerate(spans) if s < j <= e)
        head = runs[a].text[:i - spans[a][0]]
        tail = runs[b].text[j - spans[b][0]:]
        runs[a].text = head + new + tail
        for k in range(a + 1, b + 1):
            runs[k].text = ''
        LOG.append((n, name, old, new))
        return
    raise KeyError((n, name, old))


def cell(n, name, r, c, new):
    tf = find(sl(n), name).table.cell(r, c).text_frame
    p = tf.paragraphs[0]
    old = ''.join(x.text for x in p.runs)
    if old != new:
        p.runs[0].text = new
        for x in p.runs[1:]:
            x.text = ''
        LOG.append((n, f'{name}[{r},{c}]', old, new))


# ---------------------------------------------------------------- bridge (mesmo algoritmo da rev2, sem redesenhar)
X0, X1, YT, YB, Y_BR = 300000, 8450000, 4950000, 6420000, 4480000
LBL_D, LBL_T = 150000, 210000


def upd_bridge(n, spec):
    s = sl(n)
    cols = [('Fcst R9', 'tot', spec['start'])] + [(k, 'd', v) for k, v in spec['deltas']] + [('MP’27', 'tot', spec['end'])]
    bars, cum = [], 0.0
    for _, kind, v in cols:
        if kind == 'tot':
            lo, hi, cum = min(0, v), max(0, v), v
        else:
            lo, hi = min(cum, cum + v), max(cum, cum + v)
            cum += v
        bars.append((lo, hi, cum))
    vmax = max(max(b[1] for b in bars), 0)
    vmin = min(min(b[0] for b in bars), 0)
    sc = (YB - (200000 if vmin < 0 else 0) - YT) / (vmax - vmin)
    y = lambda v: int(YT + (vmax - v) * sc)
    grp = find(s, 'Bridge MP27 (Base Book)')
    con = [x for x in grp.shapes if x.name == 'Conector']
    for i, (lab, kind, v) in enumerate(cols):
        lo, hi, lvl = bars[i]
        b = find(s, f'Barra {lab}')
        b.top, b.height = Emu(y(hi)), Emu(max(y(lo) - y(hi), 12700))
        t = find(s, f'Valor {lab}')
        if kind == 'tot':
            t.top = Emu((y(0) + 25000) if v < 0 else (y(hi) - 25000 - LBL_T))
            txt = f1(v)
        else:
            t.top = Emu(y(hi) - 25000 - LBL_D)
            txt = sg(v)
        old = t.text_frame.text
        if old != txt:
            t.text_frame.paragraphs[0].runs[0].text = txt
            LOG.append((n, f'Valor {lab}', old, txt))
        if i < len(cols) - 1:
            con[i].top = Emu(y(lvl))
    lt = lambda i: (y(0) - 20000) if cols[i][2] < 0 else (y(bars[i][1]) - 25000 - LBL_T - 30000)
    find(s, 'Colchete esq').height = Emu(lt(0) - Y_BR)
    find(s, 'Colchete dir').height = Emu(lt(len(cols) - 1) - Y_BR)
    var = find(s, 'Variacao')
    old = var.text_frame.text
    new = pc((spec['end'] - spec['start']) / spec['start'], signed=True)
    var.text_frame.paragraphs[0].runs[0].text = new
    LOG.append((n, 'Variacao', old, new))
    notes = s.notes_slide.notes_text_frame
    lines = notes.text.split('\n')
    vals = {lab.replace('’', "'"): v for lab, _, v in cols}
    out = []
    for ln in lines:
        k = ln.split('\t')[0]
        out.append(f"{k}\t{vals[k]:.3f}" if '\t' in ln and k in vals else ln.replace('de 19/10', 'de 21/10').replace('de 20/10', 'de 21/10'))
    notes.text = '\n'.join(out)



# ================================================================ valores novos
d = BRG['BR']['d']
RES_B = bridge('Tabela Variantes CM s targ', 13, 15)['d']
r4_27 = sum(E27.values())
r4_9 = r9_4u
tgt = 0.13
RISCO_TOT = FIAT_PRECO + RISCO_MO
sh_ns_br = {u: pl(u, 'ns') / ns27 for u in UN}
ibsj = sh_eb['IBI'] + sh_eb['SJP']

# ================================================================ slide 2 - bridge consolidado (estrutura mantida)
rep(2, 'Text 10', '83,2', f1(ns27))
rep(2, 'Text 11', '+8,6%', pc(ns27 / ns9 - 1, signed=True))
rep(2, 'Text 15', '−10,2%', pc(r4_27 / r4_9 - 1, signed=True))
rep(2, 'Text 18', '13,8%', pc(ros27))
rep(2, 'Text 19', '−3,2 p.p.', f"{f1((ros27 - ros9) * 100)} p.p.")
D4 = {u: E27[u] - pl(u, 'ebit', 3) for u in UN}
SC2, YB2 = 205740, 5093208           # escala da v5: 205.740 EMU por R$ mi, base em y = 5.093.208
y2 = lambda v: int(YB2 - v * SC2)
s2 = sl(2)
byid = {x.name: x for x in s2.shapes}
lvl = r4_9
byid['Shape 22'].top, byid['Shape 22'].height = Emu(y2(lvl)), Emu(YB2 - y2(lvl))
bars_ = {'IBI': 'Shape 24', 'SJP': 'Shape 27', 'GOI': 'Shape 30', 'RES': 'Shape 33'}
labs_ = {'IBI': 'Text 26', 'SJP': 'Text 29', 'GOI': 'Text 32', 'RES': 'Text 35'}
conn_after = {'IBI': 'Shape 28', 'SJP': 'Shape 31', 'GOI': 'Shape 34', 'RES': 'Shape 36'}
byid['Shape 25'].top = Emu(y2(lvl))
for u in UN:
    dv = D4[u]
    new = lvl + dv
    lo, hi = min(lvl, new), max(lvl, new)
    b = byid[bars_[u]]
    b.top, b.height = Emu(y2(hi)), Emu(max(y2(lo) - y2(hi), 12700))
    byid[labs_[u]].top = Emu(y2(hi) - 256032)
    old = byid[labs_[u]].text_frame.text
    byid[labs_[u]].text_frame.paragraphs[0].runs[0].text = sg(dv)
    LOG.append((2, 'bridge ' + u, old, sg(dv)))
    lvl = new
    byid[conn_after[u]].top = Emu(y2(lvl))
fin = byid['Shape 37']
fin.top, fin.height = Emu(y2(lvl)), Emu(YB2 - y2(lvl))
byid['Text 38'].top = Emu(y2(lvl) - 256032)
rep(2, 'Text 57', '6,8', num(E27['IBI']))
rep(2, 'Text 58', '−0,6', sg(D4['IBI']))
rep(2, 'Text 59', 'Reajuste de preço amplia a cobertura do 3º turno, ANTT e Jaguar', 'Preço Fiat e fixos compensam MO, ANTT e Jaguar')
rep(2, 'Text 75', '0,2', num(E27['RES']))
rep(2, 'Text 76', '+0,3', sg(D4['RES']))
rep(2, 'Text 79', '86%', pc(ibsj, 0))
s2.notes_slide.notes_text_frame.text = (
    "Fonte: Base Book Fitted Units MP27 (21/10), MENS FITTED MP27 v4. Escopo: Resende, Ibirité, Goiana e São José dos Pinhais.\n"
    f"R09 = {f1(r4_9)}. MP’27 = {f1(E27['IBI'])} + {f1(E27['SJP'])} + {f1(E27['GOI'])} + {f1(E27['RES'])} = {f1(r4_27)}. "
    f"Variação no perímetro = {sg(r4_27 - r4_9)} ({pc(r4_27 / r4_9 - 1, signed=True)}). ROS = {f1(r4_27)} / {f1(ns27)} = {pc(ros27)}.\n"
    "O R09 Brasil de 13,0 inclui 0,2 de outras receitas da SOR encerrada.")

# ================================================================ slide 3 - leitura economica
s3 = sl(3)
b3 = {x.name: x for x in s3.shapes}
SC3, YB3 = 105156, 4297680
h27 = int(ros27 * 100 * SC3)
b3['Shape 11'].top, b3['Shape 11'].height = Emu(YB3 - h27), Emu(h27)
b3['Text 12'].top = Emu(YB3 - h27 - 310896)
rep(3, 'Text 12', '13,8%', pc(ros27))
rep(3, 'Text 19', '−3,2 p.p. vs R09', f"{f1((ros27 - ros9) * 100)} p.p. vs R09")
rep(3, 'Text 19', '+0,8 p.p. vs target', f"{sg((ros27 - tgt) * 100)} p.p. vs target")
rep(3, 'Text 29', '83,2', f1(ns27))
rep(3, 'Text 30', '+8,6%', pc(ns27 / ns9 - 1, signed=True))
rep(3, 'Text 37', '13,8%', pc(ros27))
rep(3, 'Text 47', 'ROS de 13,8%', f"ROS de {pc(ros27)}")
s3.notes_slide.notes_text_frame.text = (
    "Mensagem sugerida para apresentação:\n“O resultado do MP’27 é inferior ao R09, mas permanece aderente ao modelo econômico do negócio. "
    f"O R09 refletia ROS de {pc(ros9, 0)}, acima do target aproximado de 13% dos contratos open book. No MP’27, mesmo com pressões de mão de obra, "
    f"terceiro turno e ANTT, o ROS é de {pc(ros27)} e o EBIT de R$ {f1(br27)} mi.”")

# ================================================================ slide 4 - Brasil (tabela colada intacta)
upd_bridge(4, BRG['BR'])
rep(4, 'Comentarios', 'ROS 13,7%', f"ROS {pc(ros27)}")
rep(4, 'Comentarios', '−1,5 vs Fcst R9 (−11,9%) · +1,6 vs MP\'26',
    f"{sg(br27 - br9)} vs Fcst R9 ({pc(br27 / br9 - 1, signed=True)}) · {sg(br27 - brmp26)} vs MP'26")
rep(4, 'Comentarios', 'Receita +9,2%', f"Receita {pc(ns27 / ns9 - 1, signed=True)}")
rep(4, 'Comentarios', 'Price/Mix +6,1', f"Price/Mix {sg(d['Price/ Mix'])}")
rep(4, 'Comentarios', 'Fixos −2,5: serviços do 3º turno IBI', f"Fixos {sg(d['Cost Fix.'])}: base fixa de IBI revisada")
rep(4, 'Comentarios', 'EBIT de ~5,3', f"EBIT de ~{f1(sem_fiat_br)}")

# ================================================================ slide 5 - mix do EBIT
rep(5, 'Mensagem principal', '84%', pc(ibsj, 0))
ch = [x for x in sl(5).shapes if x.has_chart][1].chart
cd = CategoryChartData()
cd.categories = UN
cd.add_series('EBIT', [E27[u] for u in UN])
ch.replace_data(cd)
NS_C = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
for lbl in ch._chartSpace.findall('.//c:dLbl', NS_C):
    t = lbl.find('.//a:t', NS_C)
    if t is not None:
        u = UN[int(lbl.find('c:idx', NS_C).get('val'))]
        LOG.append((5, f'Pizza MP27 rotulo {u}', t.text, pc(sh_eb[u], 0)))
        t.text = pc(sh_eb[u], 0)
rep(5, 'Variacao EBIT', '−11,9%', pc(br27 / br9 - 1, signed=True))
T5 = 'Table 69'
for r, u in enumerate(UN, start=1):
    cell(5, T5, r, 2, pc(sh_ns_br[u], 0))
    cell(5, T5, r, 5, num(E27[u]) if E27[u] >= 0 else par(E27[u]))
    cell(5, T5, r, 6, pc(sh_eb[u], 0))
    cell(5, T5, r, 7, pc(pl(u, 'ros')))
    cell(5, T5, r, 8, par(E27[u] - pl(u, 'ebit', 3)))
cell(5, T5, 6, 5, num(br27))
cell(5, T5, 6, 7, pc(ros27))
cell(5, T5, 6, 8, par(br27 - br9))
rep(5, 'Nota de rodape', '(20/10)', '(21/10)')
rep(5, 'Comentarios', ': 57% do Net Sales', f": {pc(sh_ns_br['IBI'], 0)} do Net Sales")
rep(5, 'Comentarios', 'Fiat = 48% do EBIT', f"Fiat = {pc(C_EBIT[0] / 1000 / br27, 0)} do EBIT")

# ================================================================ slide 6 - resumo executivo
rep(6, 'Mensagem principal', 'ROS 13,7%', f"ROS {pc(ros27)}")
rep(6, 'KPI 2 valor', '83,7', f1(ns27))
rep(6, 'KPI 2 delta', '+9,2%', pc(ns27 / ns9 - 1, signed=True))
rep(6, 'KPI 2 delta', 'NS/pç 16,13', f"NS/pç {f1(pl('BR', 'nspc'), 2)}")
rep(6, 'KPI 3 delta', '−11,9%', pc(br27 / br9 - 1, signed=True))
rep(6, 'KPI 3 delta', "+16,7% vs MP'26", f"{pc(br27 / brmp26 - 1, signed=True)} vs MP'26")
rep(6, 'KPI 4 valor', '13,7%', pc(ros27))
rep(6, 'KPI 4 delta', '−3,3 p.p.', f"{f1((ros27 - ros9) * 100)} p.p.")
T6 = 'Table 74'
for r, u in enumerate(UN, start=1):
    cell(6, T6, r, 2, f1(pl(u, 'ns')))
    cell(6, T6, r, 4, num(E27[u]) if E27[u] >= 0 else par(E27[u]))
    cell(6, T6, r, 5, par(E27[u] - pl(u, 'ebit', 3)))
    cell(6, T6, r, 6, pc(pl(u, 'ros')))
    cell(6, T6, r, 7, pc(sh_eb[u], 0))
cell(6, T6, 1, 8, 'Preço Fiat e fixos cobrem MO e ANTT')
cell(6, T6, 5, 2, f1(ns27))
cell(6, T6, 5, 4, num(br27))
cell(6, T6, 5, 5, par(br27 - br9))
cell(6, T6, 5, 6, pc(ros27))
rep(6, 'Nota de rodape', '(20/10)', '(21/10)')
rep(6, 'Nota de rodape', 'v3', 'v4')
rep(6, 'Bloco 1 texto', 'Preço: +6,1', f"Preço: {sg(d['Price/ Mix'])}")
rep(6, 'Bloco 2 texto', '3º turno IBI: serviços +1,4 mi', f"Variáveis {f1(d['Cost Var.'])} · fixos RES {f1(RES_B['Cost Fix.'])}")
rep(6, 'Bloco 3 texto', 'Fiat (6,1 mi)', f"Fiat ({f1(FIAT_PRECO)} mi)")

# ================================================================ slide 7 - IBI (tabela colada intacta)
upd_bridge(7, BRG['IBI'])
rep(7, 'Comentarios', 'ROS 13,9%', f"ROS {pc(pl('IBI', 'ros'))}")
rep(7, 'Comentarios', '−0,8 vs Fcst R9 (−11,3%) · +0,9 vs MP\'26',
    f"{sg(ib27 - ib9)} vs Fcst R9 ({pc(ib27 / ib9 - 1, signed=True)}) · {sg(ib27 - ibmp26)} vs MP'26")
rep(7, 'Comentarios', 'Preço adicional Fiat +16,2%/pç', f"Preço Fiat {pc(FIAT_PCT, signed=True)}/pç")
rep(7, 'Comentarios', '+6,1 mi: cobre 102% do 3º turno e ANTT', f"{sg(FIAT_PRECO)} mi vs R9; fixos {sg(BRG['IBI']['d']['Cost Fix.'])} (base revisada)")
rep(7, 'Comentarios', 'EBIT de IBI ~0,4', f"EBIT de IBI ~{f1(ib27 - FIAT_PRECO)}")

# ================================================================ slide 8 - rentabilidade por cliente
rep(8, 'Mensagem principal', '93% da receita e 84% do EBIT', f"{pc(cl_ns[0], 0)} da receita e {pc(cl_eb[0], 0)} do EBIT")
rep(8, 'Mensagem principal', 'ROS de 12,5%', f"ROS de {pc(C_ROS[0])}")
for k, c in enumerate(['Fiat', 'Iveco', 'CNH']):
    old_ros = ''.join(r.text for r in find(sl(8), f'ROS {c}').text_frame.paragraphs[0].runs)[4:]
    rep(8, f'ROS {c}', old_ros, pc(C_ROS[k]))
    for lab, v in (('% da receita IBI', cl_ns[k]), ('% do EBIT IBI', cl_eb[k])):
        t = find(sl(8), f'Valor {c} {lab}')
        rep(8, f'Valor {c} {lab}', t.text_frame.text, pc(v, 0))
        find(sl(8), f'Barra {c} {lab}').width = Emu(max(int(find(sl(8), f'Trilho {c} {lab}').width * v), 25400))
T8 = 'Table 95'
for c in range(4):
    cell(8, T8, 2, c + 1, f1(C_NS[c] / 1000))
    cell(8, T8, 3, c + 1, f1(C_NSPC[c], 2))
    cell(8, T8, 4, c + 1, f1(C_CM[c] / 1000))
    cell(8, T8, 5, c + 1, par(C_FIXR[c] / 1000))
    cell(8, T8, 6, c + 1, f1(C_EBIT[c] / 1000))
    cell(8, T8, 7, c + 1, pc(C_ROS[c]))
rep(8, 'Comentarios', '93% da receita e 84% do EBIT', f"{pc(cl_ns[0], 0)} da receita e {pc(cl_eb[0], 0)} do EBIT")
rep(8, 'Comentarios', 'ROS 12,5%, acima', f"ROS {pc(C_ROS[0])}, acima")
rep(8, 'Comentarios', 'ROS 31,9%', f"ROS {pc(C_ROS[1])}")
rep(8, 'Comentarios', 'ROS 30,6%, com', f"ROS {pc(C_ROS[2])}, com")
rep(8, 'Comentarios', 'Iveco e CNH: 7% da receita, 16% do EBIT',
    f"Iveco e CNH: {pc(cl_ns[1] + cl_ns[2], 0)} da receita, {pc(cl_eb[1] + cl_eb[2], 0)} do EBIT")

# ================================================================ slide 13 - new business (CNH)
rep(13, 'Card NB 3 texto', 'EBIT R$ 0,6 mi · ROS 31%', f"EBIT R$ {f1(C_EBIT[2] / 1000)} mi · ROS {pc(C_ROS[2], 0)}")

# ================================================================ slide 14 - riscos e oportunidades (estrutura mantida)
rep(14, 'Mensagem principal', 'de −7,9 mi', f"de {f1(-RISCO_TOT)} mi")
rep(14, 'Mensagem principal', 'concentrados no reajuste Fiat', 'concentrados no reajuste Fiat e na MO')
rep(14, 'Riscos item 1 texto', '+16,2%/pç', f"{pc(FIAT_PCT, signed=True)}/pç")
rep(14, 'Riscos item 1 valor', '−6,1', f1(-FIAT_PRECO))
rep(14, 'Total Riscos valor', '−7,9', f1(-RISCO_TOT))
rep(14, 'Oportunidades item 2 texto', '3º turno já coberto pelo preço', '3º turno: custo fora da base fixa de IBI')
rep(14, 'Oportunidades item 2 texto', 'Preço Fiat cobre 102% de 3º turno + ANTT', 'Other Fixed −4,8 mi vs MENS anterior')
rep(14, 'Riscos item 5 valor', '−0,6 / 1%', f"{f1(-SENS)} / 1%")
rep(14, 'Oportunidades item 3 valor', '+0,6 / 1%', f"{sg(SENS)} / 1%")
rep(14, 'Nota de rodape', 'v3', 'v4')

prs.save(OUT)
for n, nm, o, w in LOG:
    print(f"[{n}] {nm}: {o!r} -> {w!r}")
print(f"\nEBIT BR {br27:.3f} | IBI {ib27:.3f} | reajuste Fiat {FIAT_PRECO:.3f} ({FIAT_PCT:.4f}) | sens {SENS:.3f} | riscos {RISCO_TOT:.2f}")
