"""PPT MP27 Fitted Units - atualiza valores e comentarios com o novo EBIT de Ibirite (pedido de 20/10).

Uso: python atualizar_ppt_mp27_ebit_ibi.py <deck_v4.pptx> <basebook.xlsx> <mens.xlsx> <rentab_cliente_ibi.xlsx> <saida.pptx>

Parte da v4 salva pela Juliana e NAO muda layout: so troca textos (mantendo a formatacao dos runs),
redimensiona as barras dos bridges (mesma escala/algoritmo do draw_bridge da rev2), as barras de
participacao da rentabilidade por cliente e os dados das pizzas. Slides 2, 3 e 4 nao sao tocados.
"""
import sys

import openpyxl
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Emu

DECK, BB, MENS, CLI, OUT = sys.argv[1:6]
MINUS = '−'


# ---------------------------------------------------------------- formatos (iguais aos da rev2)
def num(x, d=1):
    s = f"{abs(x):,.{d}f}".replace(',', 'X').replace('.', ',').replace('X', '.')
    return s


def f1(x, d=1):
    s = num(x, d)
    return (MINUS + s) if round(x, d) < 0 or (x < 0 and s.strip('0,.') == '') else s


def sg(x, d=1):
    return (MINUS if x < 0 else '+') + num(x, d)


def pc(x, d=1, signed=False):
    s = num(x * 100, d) + '%'
    return ((MINUS if x < 0 else '+') + s) if signed else s


def par(x):
    return f"({num(x)})" if round(x, 1) < 0 else num(x)


# ---------------------------------------------------------------- fontes
bb = openpyxl.load_workbook(BB, data_only=True)


def bridge(sheet, hr, dr):
    ws = bb[sheet]
    items, end = [], None
    for col in range(3, 20):
        h, v = ws.cell(hr, col).value, ws.cell(dr, col).value
        if h is None:
            continue
        h = str(h).strip()
        if h == "MP'27" or h.startswith('Delta') or h.startswith('%'):
            continue
        if h == 'EBIT':
            end = float(v)
            break
        items.append((h.replace('\n', ' '), float(v)))
    assert abs(items[0][1] + sum(v for _, v in items[1:]) - end) < 0.003
    return {'start': items[0][1], 'deltas': items[1:], 'end': end, 'd': dict(items[1:])}


BRG = {'BR': bridge('Tabela Variantes s targ', 2, 4), 'IBI': bridge('Tabela Variantes s targ', 6, 8)}
PLS = {'BR': 'Month_FY_LC - MP27', 'IBI': 'IBIRITE', 'SJP': 'SJP', 'GOI': 'GOIANA', 'RES': 'RESENDE'}
ROWS = {'pcs': 42, 'ns': 44, 'nspc': 45, 'ebit': 55, 'ros': 58}


def pl(u, k, i=4):          # i: 2 = MP26, 3 = R9, 4 = MP27
    return float(bb[PLS[u]].cell(ROWS[k], 39 + i).value or 0)


for u in BRG:
    assert abs(BRG[u]['end'] - pl(u, 'ebit')) < 0.002 and abs(BRG[u]['start'] - pl(u, 'ebit', 3)) < 0.002, u
UN = ['IBI', 'SJP', 'GOI', 'RES']
E27 = {u: pl(u, 'ebit') for u in UN}
assert abs(sum(E27.values()) - pl('BR', 'ebit')) < 0.002

mens = openpyxl.load_workbook(MENS, data_only=True)
jl = mens['Impacto JLR']
JLR = {}
for r in range(1, 60):
    k = jl.cell(r, 3).value
    if k in ('Fiat', 'Iveco', 'CNH') and k not in JLR:
        JLR[k] = [jl.cell(r, c).value for c in range(4, 14)]
FIAT_PR9, FIAT_PR27, FIAT_PRECO = JLR['Fiat'][2], JLR['Fiat'][3], JLR['Fiat'][9] / 1000
FIAT_PCT = FIAT_PR27 / FIAT_PR9 - 1
mt = mens['TOTAL']
trow = {mt.cell(r, 3).value: r for r in range(1, 60) if mt.cell(r, 3).value}
tv = lambda k: float(mt.cell(trow[k], 17).value)
SENS = (tv('NET SALES') + tv('Direct Materials') + tv('Transportation') + tv('Other Variable')) / 1000 / 100

wc = openpyxl.load_workbook(CLI, data_only=True)['Rentabilidade por Cliente']
crow = lambda r: [float(wc.cell(r, c).value or 0) for c in (2, 3, 4, 5)]
C_PCS, C_NS, C_NSPC, C_CM, C_FIX, C_REAL, C_EBIT, C_ROS = (crow(5), crow(6), crow(7), crow(14), crow(16), crow(23),
                                                           crow(24), crow(25))
assert str(wc['A23'].value).startswith('Realocacao') and str(wc['A24'].value) == 'EBIT'
assert abs(C_EBIT[3] / 1000 - pl('IBI', 'ebit')) < 0.002
ROS_META = float(openpyxl.load_workbook(CLI, data_only=True)['Premissas']['B23'].value)
C_FIXR = [a + b for a, b in zip(C_FIX, C_REAL)]      # fixo + realocacao para a meta da Fiat (tabela fecha no EBIT)
REAL_FIAT = C_REAL[0] / 1000

T3_MO, T3_DESP = 2.5, 2.0                             # premissa da Juliana (19/10)
ANTT = -BRG['IBI']['d']['ANTT']
COBERT = FIAT_PRECO / (T3_MO + T3_DESP + ANTT)
RISCO_MO = 0.95 + 0.46 + 0.36
RISCOS = FIAT_PRECO + RISCO_MO
OPORT = ANTT + max(0.0, T3_MO + T3_DESP + ANTT - FIAT_PRECO)

br9, br27, brmp26 = pl('BR', 'ebit', 3), pl('BR', 'ebit'), pl('BR', 'ebit', 2)
ns9, ns27 = pl('BR', 'ns', 3), pl('BR', 'ns')
ros27, ros9 = pl('BR', 'ros'), pl('BR', 'ros', 3)
ib9, ib27, ibmp26 = pl('IBI', 'ebit', 3), pl('IBI', 'ebit'), pl('IBI', 'ebit', 2)
r9_4u = sum(pl(u, 'ebit', 3) for u in UN)
sh_eb = {u: E27[u] / br27 for u in UN}
sh_ns = {u: pl(u, 'ns') / ns27 for u in UN}
cl_ns = [C_NS[i] / C_NS[3] for i in range(3)]
cl_eb = [C_EBIT[i] / C_EBIT[3] for i in range(3)]
sem_fiat_br = br27 - FIAT_PRECO
ros_sem_fiat = sem_fiat_br / (ns27 - FIAT_PRECO)

# ---------------------------------------------------------------- utilitarios de edicao
prs = Presentation(DECK)
LOG = []


def sl(n):
    return prs.slides[n - 1]


def find(slide, name, idx=0):
    out = []

    def walk(shs):
        for x in shs:
            if x.name == name:
                out.append(x)
            if x.shape_type == 6:
                walk(x.shapes)
    walk(slide.shapes)
    return out[idx]


def rep(n, name, old, new, idx=0):
    """Troca 'old' por 'new' no texto da forma mantendo a formatacao do run onde o trecho comeca."""
    if old == new:
        return
    sh = find(sl(n), name, idx)
    for p in sh.text_frame.paragraphs:
        runs = p.runs
        full = ''.join(r.text for r in runs)
        i = full.find(old)
        if i < 0:
            continue
        pos, spans = 0, []
        for r in runs:
            spans.append((pos, pos + len(r.text)))
            pos += len(r.text)
        j = i + len(old)
        a = next(k for k, (s, e) in enumerate(spans) if s <= i < e)
        b = next(k for k, (s, e) in enumerate(spans) if s < j <= e)
        head = runs[a].text[:i - spans[a][0]]
        tail = runs[b].text[j - spans[b][0]:]
        runs[a].text = head + new + tail
        for k in range(a + 1, b + 1):
            runs[k].text = ''
        LOG.append((n, name, old, new))
        return
    raise KeyError((n, name, old))


def cell(n, name, r, c, new):
    tf = find(sl(n), name).table.cell(r, c).text_frame
    p = tf.paragraphs[0]
    old = ''.join(x.text for x in p.runs)
    if old != new:
        p.runs[0].text = new
        for x in p.runs[1:]:
            x.text = ''
        LOG.append((n, f'{name}[{r},{c}]', old, new))


# ---------------------------------------------------------------- bridge (mesmo algoritmo da rev2, sem redesenhar)
X0, X1, YT, YB, Y_BR = 300000, 8450000, 4950000, 6420000, 4480000
LBL_D, LBL_T = 150000, 210000


def upd_bridge(n, spec):
    s = sl(n)
    cols = [('Fcst R9', 'tot', spec['start'])] + [(k, 'd', v) for k, v in spec['deltas']] + [('MP’27', 'tot', spec['end'])]
    bars, cum = [], 0.0
    for _, kind, v in cols:
        if kind == 'tot':
            lo, hi, cum = min(0, v), max(0, v), v
        else:
            lo, hi = min(cum, cum + v), max(cum, cum + v)
            cum += v
        bars.append((lo, hi, cum))
    vmax = max(max(b[1] for b in bars), 0)
    vmin = min(min(b[0] for b in bars), 0)
    sc = (YB - (200000 if vmin < 0 else 0) - YT) / (vmax - vmin)
    y = lambda v: int(YT + (vmax - v) * sc)
    grp = find(s, 'Bridge MP27 (Base Book)')
    con = [x for x in grp.shapes if x.name == 'Conector']
    for i, (lab, kind, v) in enumerate(cols):
        lo, hi, lvl = bars[i]
        b = find(s, f'Barra {lab}')
        b.top, b.height = Emu(y(hi)), Emu(max(y(lo) - y(hi), 12700))
        t = find(s, f'Valor {lab}')
        if kind == 'tot':
            t.top = Emu((y(0) + 25000) if v < 0 else (y(hi) - 25000 - LBL_T))
            txt = f1(v)
        else:
            t.top = Emu(y(hi) - 25000 - LBL_D)
            txt = sg(v)
        old = t.text_frame.text
        if old != txt:
            t.text_frame.paragraphs[0].runs[0].text = txt
            LOG.append((n, f'Valor {lab}', old, txt))
        if i < len(cols) - 1:
            con[i].top = Emu(y(lvl))
    lt = lambda i: (y(0) - 20000) if cols[i][2] < 0 else (y(bars[i][1]) - 25000 - LBL_T - 30000)
    find(s, 'Colchete esq').height = Emu(lt(0) - Y_BR)
    find(s, 'Colchete dir').height = Emu(lt(len(cols) - 1) - Y_BR)
    var = find(s, 'Variacao')
    old = var.text_frame.text
    new = pc((spec['end'] - spec['start']) / spec['start'], signed=True)
    var.text_frame.paragraphs[0].runs[0].text = new
    LOG.append((n, 'Variacao', old, new))
    notes = s.notes_slide.notes_text_frame
    lines = notes.text.split('\n')
    vals = {lab.replace('’', "'"): v for lab, _, v in cols}
    out = []
    for ln in lines:
        k = ln.split('\t')[0]
        out.append(f"{k}\t{vals[k]:.3f}" if '\t' in ln and k in vals else ln.replace('de 19/10', 'de 20/10'))
    notes.text = '\n'.join(out)


# ================================================================ slide 5 - Brasil
upd_bridge(5, BRG['BR'])
d = BRG['BR']['d']
rep(5, 'Comentarios', 'EBIT 11,0  |  ROS 13,2%', f"EBIT {f1(br27)}  |  ROS {pc(ros27)}")
rep(5, 'Comentarios', '−2,0 vs Fcst R9 (−15,7%) · +1,1 vs MP\'26',
    f"{sg(br27 - br9)} vs Fcst R9 ({pc(br27 / br9 - 1, signed=True)}) · {sg(br27 - brmp26)} vs MP'26")
rep(5, 'Comentarios', 'Receita +8,6%', f"Receita {pc(ns27 / ns9 - 1, signed=True)}")
rep(5, 'Comentarios', 'Price/Mix +5,6', f"Price/Mix {sg(d['Price/ Mix'])}")
rep(5, 'Comentarios', 'EBIT de ~5,3', f"EBIT de ~{f1(sem_fiat_br)}")

# ================================================================ slide 6 - mix do EBIT
rep(6, 'Mensagem principal', 'concentram 83%', f"concentram {pc(sh_eb['IBI'] + sh_eb['SJP'], 0)}")
ch = [x for x in sl(6).shapes if x.has_chart][1].chart
cd = CategoryChartData()
cd.categories = UN
cd.add_series('EBIT', [E27[u] for u in UN])
ch.replace_data(cd)
NS_C = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
for lbl in ch._chartSpace.findall('.//c:dLbl', NS_C):     # rotulos da pizza sao texto fixo
    t = lbl.find('.//a:t', NS_C)
    if t is not None:
        u = UN[int(lbl.find('c:idx', NS_C).get('val'))]
        LOG.append((6, f'Pizza MP27 rotulo {u}', t.text, pc(sh_eb[u], 0)))
        t.text = pc(sh_eb[u], 0)
rep(6, 'Centro valor', '11,0', f1(br27), idx=1)
rep(6, 'Variacao EBIT', '−15,7%', pc(br27 / br9 - 1, signed=True))
rep(6, 'Variacao absoluta', 'Δ −2,0', f"Δ {f1(br27 - br9)}")
T6 = 'Table 69'
for r, u in enumerate(UN, start=1):
    cell(6, T6, r, 2, pc(sh_ns[u], 0))
    cell(6, T6, r, 5, num(E27[u]) if E27[u] >= 0 else par(E27[u]))
    cell(6, T6, r, 6, pc(sh_eb[u], 0))
    cell(6, T6, r, 7, pc(pl(u, 'ros')))
    cell(6, T6, r, 8, par(E27[u] - pl(u, 'ebit', 3)))
cell(6, T6, 6, 5, num(br27))
cell(6, T6, 6, 7, pc(ros27))
cell(6, T6, 6, 8, par(br27 - br9))
rep(6, 'Nota de rodape', '(19/10)', '(20/10)')
rep(6, 'Nota de rodape', "MP'27: 11,0", f"MP'27: {f1(br27)}")
rep(6, 'Comentarios', 'IBI: 55% do EBIT', f"IBI: {pc(sh_eb['IBI'], 0)} do EBIT")
rep(6, 'Comentarios', ': 56% do Net Sales', f": {pc(sh_ns['IBI'], 0)} do Net Sales")
rep(6, 'Comentarios', 'SJP: 28%', f"SJP: {pc(sh_eb['SJP'], 0)}")
rep(6, 'Comentarios', 'GOI: 13%', f"GOI: {pc(sh_eb['GOI'], 0)}")
rep(6, 'Comentarios', 'RES: 4%', f"RES: {pc(sh_eb['RES'], 0)}")
rep(6, 'Comentarios', 'Fiat = 43% do EBIT', f"Fiat = {pc(C_EBIT[0] / 1000 / br27, 0)} do EBIT")

# ================================================================ slide 7 - resumo executivo
rep(7, 'Mensagem principal', 'EBIT de R$ 11,0 mi (ROS 13,2%)', f"EBIT de R$ {f1(br27)} mi (ROS {pc(ros27)})")
rep(7, 'KPI 2 valor', '83,2', f1(ns27))
rep(7, 'KPI 2 delta', '+8,6%', pc(ns27 / ns9 - 1, signed=True))
rep(7, 'KPI 2 delta', 'NS/pç 16,03', f"NS/pç {f1(pl('BR', 'nspc'), 2)}")
rep(7, 'KPI 3 valor', '11,0', f1(br27))
rep(7, 'KPI 3 delta', '−15,7%', pc(br27 / br9 - 1, signed=True))
rep(7, 'KPI 3 delta', "+11,7% vs MP'26", f"{pc(br27 / brmp26 - 1, signed=True)} vs MP'26")
rep(7, 'KPI 4 valor', '13,2%', pc(ros27))
rep(7, 'KPI 4 delta', '−3,8 p.p.', f"{f1((ros27 - ros9) * 100)} p.p.")
T7 = 'Table 74'
for r, u in enumerate(UN, start=1):
    cell(7, T7, r, 2, f1(pl(u, 'ns')))
    cell(7, T7, r, 4, num(E27[u]) if E27[u] >= 0 else par(E27[u]))
    cell(7, T7, r, 5, par(E27[u] - pl(u, 'ebit', 3)))
    cell(7, T7, r, 6, pc(pl(u, 'ros')))
    cell(7, T7, r, 7, pc(sh_eb[u], 0))
cell(7, T7, 5, 2, f1(ns27))
cell(7, T7, 5, 4, num(br27))
cell(7, T7, 5, 5, par(br27 - br9))
cell(7, T7, 5, 6, pc(ros27))
rep(7, 'Nota de rodape', '(19/10)', '(20/10)')
rep(7, 'Bloco 1 texto', ': +5,6 (', f": {sg(d['Price/ Mix'])} (")
rep(7, 'Bloco 3 texto', 'Fiat (5,7 mi)', f"Fiat ({f1(FIAT_PRECO)} mi)")

# ================================================================ slide 8 - IBI
upd_bridge(8, BRG['IBI'])
rep(8, 'Comentarios', 'EBIT 6,1  |  ROS 13,0%', f"EBIT {f1(ib27)}  |  ROS {pc(pl('IBI', 'ros'))}")
rep(8, 'Comentarios', '−1,3 vs Fcst R9 (−17,9%) · +0,5 vs MP\'26',
    f"{sg(ib27 - ib9)} vs Fcst R9 ({pc(ib27 / ib9 - 1, signed=True)}) · {sg(ib27 - ibmp26)} vs MP'26")
rep(8, 'Comentarios', 'Fiat +14,9%/', f"Fiat {pc(FIAT_PCT, signed=True)}/")
rep(8, 'Comentarios', '+5,7 mi: cobre 94% do', f"{sg(FIAT_PRECO)} mi: cobre {pc(COBERT, 0)} do")
rep(8, 'Comentarios', 'EBIT de IBI ~0,4', f"EBIT de IBI ~{f1(ib27 - FIAT_PRECO)}")

# ================================================================ slide 9 - rentabilidade por cliente
rep(9, 'Mensagem principal', 'traz 93% da receita e 78% do EBIT', f"traz {pc(cl_ns[0], 0)} da receita e {pc(cl_eb[0], 0)} do EBIT")
rep(9, 'Mensagem principal', 'ROS de 10,9% (meta 10,5%)', f"ROS de {pc(C_ROS[0])} (meta {pc(ROS_META)})")
for k, c in enumerate(['Fiat', 'Iveco', 'CNH']):
    old_ros = ''.join(r.text for r in find(sl(9), f'ROS {c}').text_frame.paragraphs[0].runs)[4:]
    rep(9, f'ROS {c}', old_ros, pc(C_ROS[k]))
    for lab, v in (('% da receita IBI', cl_ns[k]), ('% do EBIT IBI', cl_eb[k])):
        t = find(sl(9), f'Valor {c} {lab}')
        rep(9, f'Valor {c} {lab}', t.text_frame.text, pc(v, 0))
        find(sl(9), f'Barra {c} {lab}').width = Emu(max(int(find(sl(9), f'Trilho {c} {lab}').width * v), 25400))
T9 = 'Table 95'
for c in range(4):
    cell(9, T9, 2, c + 1, f1(C_NS[c] / 1000))
    cell(9, T9, 3, c + 1, f1(C_NSPC[c], 2))
    cell(9, T9, 4, c + 1, f1(C_CM[c] / 1000))
    cell(9, T9, 5, c + 1, par(C_FIXR[c] / 1000))
    cell(9, T9, 6, c + 1, f1(C_EBIT[c] / 1000))
    cell(9, T9, 7, c + 1, pc(C_ROS[c]))
rep(9, 'Comentarios', '93% da receita e 78% do EBIT', f"{pc(cl_ns[0], 0)} da receita e {pc(cl_eb[0], 0)} do EBIT")
rep(9, 'Comentarios', 'ROS 10,9%, acima', f"ROS {pc(C_ROS[0])}, acima")
rep(9, 'Comentarios', 'ROS 41,6%', f"ROS {pc(C_ROS[1])}")
rep(9, 'Comentarios', 'ROS 38,4%, com', f"ROS {pc(C_ROS[2])}, com")
rep(9, 'Comentarios', 'Iveco e CNH: 7% da receita, 22% do EBIT',
    f"Iveco e CNH: {pc(cl_ns[1] + cl_ns[2], 0)} da receita, {pc(cl_eb[1] + cl_eb[2], 0)} do EBIT; "
    f"já após realocar {f1(REAL_FIAT)} mi para a Fiat (no Fixed Cost)")

# ================================================================ slide 14 - new business (CNH)
rep(14, 'Card NB 3 texto', 'EBIT R$ 0,7 mi · ROS 38%', f"EBIT R$ {f1(C_EBIT[2] / 1000)} mi · ROS {pc(C_ROS[2], 0)}")

# ================================================================ slide 15 - riscos e oportunidades
rep(15, 'Mensagem principal', 'de −7,4 mi', f"de {f1(-RISCOS)} mi")
rep(15, 'Mensagem principal', 'de +1,8 mi', f"de {sg(OPORT)} mi")
rep(15, 'Riscos item 1 texto', 'Preço +14,9%/pç', f"Preço {pc(FIAT_PCT, signed=True)}/pç")
rep(15, 'Riscos item 1 valor', '−5,7', f1(-FIAT_PRECO))
rep(15, 'Total Riscos valor', '−7,4', f1(-RISCOS))
gap = T3_MO + T3_DESP + ANTT - FIAT_PRECO
if gap <= 0:
    rep(15, 'Oportunidades item 2 texto', 'Completar o repasse do 3º turno', '3º turno já coberto pelo preço')
    rep(15, 'Oportunidades item 2 texto', 'cobre 94% de', f"cobre {pc(COBERT, 0)} de")
    rep(15, 'Oportunidades item 2 valor', '+0,34', '–')
else:
    rep(15, 'Oportunidades item 2 texto', 'cobre 94% de', f"cobre {pc(COBERT, 0)} de")
    rep(15, 'Oportunidades item 2 valor', '+0,34', sg(gap, 2))
rep(15, 'Total Oportunidades valor', '+1,8', sg(OPORT))
rep(15, 'Riscos item 5 valor', '−0,6 / 1%', f"{f1(-SENS)} / 1%")
rep(15, 'Oportunidades item 3 valor', '+0,6 / 1%', f"{sg(SENS)} / 1%")

# ================================================================ slide 16 - perguntas-chave
rep(16, 'Card 1 pergunta', 'sobe 8,6%', f"sobe {pc(ns27 / ns9 - 1)}")
rep(16, 'Card 1 resposta', 'somam +7,1', f"somam {sg(d['Volume'] + d['Price/ Mix'])}")
rep(16, 'Card 2 resposta', 'Price/Mix +5,6. O reajuste Fiat (+14,9%) vale 5,7 mi; sem ele, EBIT de ~5,3 (ROS ~7%)',
    f"Price/Mix {sg(d['Price/ Mix'])}. O reajuste Fiat ({pc(FIAT_PCT, signed=True)}) vale {f1(FIAT_PRECO)} mi; "
    f"sem ele, EBIT de ~{f1(sem_fiat_br)} (ROS ~{pc(ros_sem_fiat, 0)})")
rep(16, 'Card 3 resposta', 'R$ 0,6 mi de EBIT (6% do EBIT)', f"R$ {f1(SENS)} mi de EBIT ({pc(SENS / br27, 0)} do EBIT)")
rep(16, 'Card 5 resposta', 'ROS 10,9%, perto da meta', f"ROS {pc(C_ROS[0])}, acima da meta")

prs.save(OUT)
for n, nm, o, w in LOG:
    print(f"[{n}] {nm}: {o!r} -> {w!r}")
print(f"\nEBIT BR {br27:.3f} | IBI {ib27:.3f} | Fiat {FIAT_PR9}->{FIAT_PR27} ({FIAT_PCT:.4f}) = {FIAT_PRECO:.3f} | "
      f"cobertura {COBERT:.3f} | sens {SENS:.3f} | riscos {RISCOS:.2f} | oport {OPORT:.2f} | 4 unid R9 {r9_4u:.3f}")
